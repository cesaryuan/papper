"""Verify configured page geometry in actual native DOCX builds."""

import os
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from native_support import ROOT, papper_command


@pytest.mark.parametrize("mathtype", [False, True])
def test_docx_page_margins_and_equation_tabs_follow_text_width(tmp_path: Path, mathtype: bool) -> None:
    """Apply page margins while keeping numbered equations inside the resulting text area."""
    source = tmp_path / "paper.md"
    source.write_text("---\ntitle: Geometry\n---\n\n$$ x+y $$ {#eq:sum}\n\nwhere $x$ is a value.\n", encoding="utf-8")
    (tmp_path / "style.yml").write_text(
        "docxPageMargins:\n  top: 2.54cm\n  bottom: 2.54cm\n  left: 3.17cm\n  right: 3.17cm\n",
        encoding="utf-8")
    reference = ROOT / "pandoc/manuscript-template/reference-doc.docx"
    original = reference.read_bytes()
    output = tmp_path / "paper.docx"
    command = [*papper_command(), "build", "docx", str(source), "-o", str(output)]
    if not mathtype:
        command.append("--no-mathtype")
    result = subprocess.run(command, cwd=tmp_path,
                            env={**os.environ, "PAPPER_HOME": str(tmp_path / "state"),
                                 "PAPPER_RESOURCE_ROOT": str(ROOT)},
                            capture_output=True, text=True, encoding="utf-8", timeout=90)
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    for section in document.sections:
        assert section.top_margin.cm == pytest.approx(2.54, abs=0.001)
        assert section.bottom_margin.cm == pytest.approx(2.54, abs=0.001)
        assert section.left_margin.cm == pytest.approx(3.17, abs=0.001)
        assert section.right_margin.cm == pytest.approx(3.17, abs=0.001)
    if mathtype:
        section = document.sections[0]
        text_width = section.page_width.twips - section.left_margin.twips - section.right_margin.twips
        with ZipFile(output) as archive:
            xml = etree.fromstring(archive.read("word/document.xml"))
        paragraphs = xml.xpath(".//w:p[w:r/w:tab][.//o:OLEObject]", namespaces={
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "o": "urn:schemas-microsoft-com:office:office"})
        assert len(paragraphs) == 1
        equation = next(paragraph for paragraph in document.paragraphs if paragraph.style.name == "Para Equation")
        assert equation.style.name == "Para Equation"
        # Multiple pPr blocks can discard the style when Word opens the file.
        assert len(paragraphs[0].findall(qn("w:pPr"))) == 1
        assert paragraphs[0].find(f"{qn('w:pPr')}/{qn('w:tabs')}") is None
        assert len(paragraphs[0].findall(f"{qn('w:r')}/{qn('w:tab')}")) == 2
        tabs = equation.style.element.find(f"{qn('w:pPr')}/{qn('w:tabs')}")
        assert tabs is not None
        positions = {tab.get(qn("w:val")): int(tab.get(qn("w:pos"))) for tab in tabs}
        assert positions["center"] == round(text_width / 2)
        assert positions["right"] == text_width
        assert next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("where")).style.name == "Para Where"
    assert reference.read_bytes() == original
