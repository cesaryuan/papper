"""Verify border shorthand, native line styles, inheritance and atomic build failures.

These tests inspect actual HTML/DOCX output; visual snapshots additionally protect
connected paragraph boxes and the boundary between inherited and local borders.
"""

from pathlib import Path
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from lxml import etree, html
import pytest

from native_support import ROOT
from test_paragraph_format import build


# This public vocabulary comes from Word's paragraph line styles, excluding page art.
LINE_STYLES = (
    "single", "thick", "double", "dotted", "dashed", "dotDash", "dotDotDash", "triple",
    "thinThickSmallGap", "thickThinSmallGap", "thinThickThinSmallGap",
    "thinThickMediumGap", "thickThinMediumGap", "thinThickThinMediumGap",
    "thinThickLargeGap", "thickThinLargeGap", "thinThickThinLargeGap", "wave",
    "doubleWave", "dashSmallGap", "dashDotStroked", "threeDEmboss", "threeDEngrave", "outset", "inset",
)


@pytest.mark.parametrize("target", ["html", "docx"])
def test_all_word_line_styles_survive_build(tmp_path: Path, rust_executable: Path, target: str) -> None:
    """Catch rejected native styles, unit errors, or silently lost lines in either output backend."""
    source = tmp_path / "lines.md"
    source.write_text("\n\n".join(
        f'::: {{borders.bottom="1.5pt {style} #005588"}}\nLine {style}.\n:::'
        for style in LINE_STYLES
    ), encoding="utf-8")
    output = tmp_path / f"lines.{target}"
    result = build(rust_executable, source, output)
    assert result.returncode == 0, result.stdout + result.stderr
    if target == "docx":
        for paragraph, style in zip(Document(output).paragraphs, LINE_STYLES, strict=True):
            border = paragraph._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom"))
            assert border.get(qn("w:val")) == style
            assert border.get(qn("w:sz")) == "12"
            assert border.get(qn("w:color")) == "005588"
    else:
        document = html.fromstring(output.read_text(encoding="utf-8"))
        paragraphs = document.xpath("//main//p")
        assert len(paragraphs) == len(LINE_STYLES)
        for paragraph in paragraphs:
            assert "border-bottom: 1.5pt " in paragraph.get("style", "")
            assert "#005588" in paragraph.get("style", "")


@pytest.mark.parametrize("target,postprocess", [("html", True), ("docx", True), ("docx", False)])
def test_named_local_and_reset_borders(tmp_path: Path, rust_executable: Path, target: str, postprocess: bool) -> None:
    """Prevent shorthand, nested overrides, RGB colors and explicit resets from diverging."""
    source = ROOT / "tests/snapshot_cases/paragraph_format/paragraph_format.md"
    output = tmp_path / f"borders.{target}"
    result = build(rust_executable, source, output, postprocess=postprocess)
    assert result.returncode == 0, result.stdout + result.stderr
    if target == "html":
        document = html.fromstring(output.read_text(encoding="utf-8"))
        box = document.xpath("//div[@id='local-box']/p")
        assert "border-top: 0.75pt solid #808080" in box[0].get("style")
        assert "border-bottom: 0.5pt dotted #005588" in box[0].get("style")
        assert "border-top: none" in box[1].get("style")
        assert "border-bottom: 0.75pt solid #808080" in box[1].get("style")
        assert "padding-left: 4pt" in box[1].get("style")
        reset = document.xpath("//p[contains(normalize-space(.), 'Local reset removes')]")[0]
        for side in ("top", "right", "bottom", "left"):
            assert f"border-{side}: none" in reset.get("style")
        assert "border-left: 2pt solid #005588" in document.xpath("//div[@id='left-only']/p")[0].get("style")
        assert "data-papper-local-borders" not in output.read_text(encoding="utf-8")
        return
    document = Document(output)
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    named = document.styles["Block Text"]._element.pPr.find(qn("w:pBdr"))
    assert [child.tag for child in named] == [qn("w:" + side) for side in ("top", "left", "bottom", "right", "between")]
    assert named.find(qn("w:top")).get(qn("w:space")) == "4"
    wave = paragraphs["Local wave bottom overrides the named style only on this paragraph."]._p.pPr.find(qn("w:pBdr"))
    assert wave.find(qn("w:bottom")).get(qn("w:val")) == "wave"
    assert wave.find(qn("w:bottom")).get(qn("w:space")) == "4"
    reset = paragraphs["Local reset removes every inherited border."]._p.pPr.find(qn("w:pBdr"))
    assert all(child.get(qn("w:val")) in {"none", "nil"} for child in reset)
    left = paragraphs["A left-side rule using the CSS-like shorthand."]._p.pPr.find(qn("w:pBdr"))
    assert left.find(qn("w:left")).get(qn("w:color")) == "005588"
    with ZipFile(output) as archive:
        assert b"PMT_PARAGRAPH_FORMAT:" not in archive.read("word/document.xml")


@pytest.mark.parametrize("target", ["html", "docx"])
@pytest.mark.parametrize("attributes", [
    'borders="0pt solid #000000"', 'borders="0.3pt solid #000000"',
    'borders="1pt apples #000000"', 'borders="solid dashed"',
    'borders="none #000000"', 'borders="1pt solid #GGGGGG"',
    'borders.space="32pt"', 'borders.space="1.5pt"',
    'borders.between.space="4pt"', 'borders.botom="solid"',
])
def test_invalid_borders_preserve_existing_output(tmp_path: Path, rust_executable: Path, target: str, attributes: str) -> None:
    """Reject ambiguous/unsupported inputs without damaging an already published manuscript."""
    source = tmp_path / "invalid.md"
    source.write_text(f"::: {{{attributes}}}\nParagraph.\n:::\n", encoding="utf-8")
    output = tmp_path / f"output.{target}"
    output.write_bytes(b"previous output")
    result = build(rust_executable, source, output)
    assert result.returncode != 0
    assert "border" in (result.stdout + result.stderr).lower()
    assert output.read_bytes() == b"previous output"


def test_border_distance_inherits_configured_parent_independent_of_order(tmp_path: Path, rust_executable: Path) -> None:
    """A distance-only child must inherit its parent's configured line, even if the child is first."""
    reference = Document(ROOT / "pandoc/manuscript-template/reference-doc.docx")
    parent = reference.styles.add_style("Border Parent", WD_STYLE_TYPE.PARAGRAPH)
    child = reference.styles.add_style("Border Child", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    reference_path = tmp_path / "reference.docx"
    reference.save(reference_path)
    source = tmp_path / "inherit.md"
    source.write_text('''---
papperSettings:
  docxStyle:
    Border Child:
      borders: {space: 6pt}
    Border Parent:
      borders: "1pt wave #005588"
---
::: {custom-style="Border Child"}
Inherited border.
:::
''', encoding="utf-8")
    output = tmp_path / "inherit.docx"
    result = subprocess.run([str(rust_executable), "build", "docx", str(source), "--reference-doc", str(reference_path),
                             "--no-mathtype", "-o", str(output)], capture_output=True, text=True, encoding="utf-8", timeout=90)
    assert result.returncode == 0, result.stdout + result.stderr
    with ZipFile(output) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
        border = styles.xpath("w:style[w:name/@w:val='Border Child']/w:pPr/w:pBdr/w:bottom", namespaces=styles.nsmap)[0]
        assert border.get(qn("w:val")) == "wave"
        assert border.get(qn("w:sz")) == "8"
        assert border.get(qn("w:space")) == "6"


@pytest.mark.parametrize("target,postprocess", [("html", True), ("docx", True), ("docx", False)])
def test_shorthand_defaults_priority_and_story_borders(tmp_path: Path, rust_executable: Path, target: str, postprocess: bool) -> None:
    """Protect compact CSS syntax, distance-only overrides, native page stories and note scopes."""
    source = tmp_path / "compact.md"
    source.write_text('''---
papperSettings:
  docxStyle:
    Quote:
      borders: {all: "1pt wave #005588", space: 5pt}
  docxHeader:
    content: |
      ::: {borders.bottom="0.75pt dashed #005588" borders.space="3pt"}
      Bordered header.
      :::
---
::: {borders="2pt double #123456" borders.all="solid" borders.bottom="#abc dotted 1pt"}
CSS order and short hex.
:::

::: {custom-style="Quote" borders.space="7pt"}
Only distance changes.
:::

::: {borders="solid"}
Body note reference.[^note]
:::

[^note]: Ordinary note.

    ::: {borders.bottom="double"}
    Bordered note.
    :::
''', encoding="utf-8")
    output = tmp_path / f"compact.{target}"
    result = build(rust_executable, source, output, postprocess=postprocess)
    assert result.returncode == 0, result.stdout + result.stderr
    if target == "html":
        document = html.fromstring(output.read_text(encoding="utf-8"))
        css = document.xpath("//p[contains(normalize-space(.), 'CSS order')]")[0].get("style")
        assert "border-top: 0.5pt solid #000000" in css
        assert "border-bottom: 1pt dotted #AABBCC" in css
        distance = document.xpath("//p[contains(normalize-space(.), 'Only distance')]")[0].get("style")
        assert "border-bottom: 1pt solid #005588" in distance
        assert "padding-bottom: 7pt" in distance
        note = document.xpath("//p[contains(normalize-space(.), 'Bordered note.')]")[0]
        assert "border-bottom: 0.5pt double #000000" in note.get("style")
        ordinary = document.xpath("//p[contains(normalize-space(.), 'Ordinary note.')]")[0]
        assert "border" not in ordinary.get("style", "")
        return
    document = Document(output)
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    css = paragraphs["CSS order and short hex."]._p.pPr.find(qn("w:pBdr"))
    assert css.find(qn("w:top")).get(qn("w:sz")) == "4"
    assert css.find(qn("w:top")).get(qn("w:val")) == "single"
    assert css.find(qn("w:bottom")).get(qn("w:color")) == "AABBCC"
    distance = paragraphs["Only distance changes."]._p.pPr.find(qn("w:pBdr"))
    assert distance.find(qn("w:bottom")).get(qn("w:val")) == "wave"
    assert distance.find(qn("w:bottom")).get(qn("w:space")) == "7"
    header = next(paragraph for paragraph in document.sections[0].header.paragraphs if paragraph.text)
    assert header._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")).get(qn("w:val")) == "dashed"
    with ZipFile(output) as archive:
        notes = etree.fromstring(archive.read("word/footnotes.xml"))
        paragraphs = {"".join(node.itertext()): node for node in notes.iter(qn("w:p"))}
        assert paragraphs["Bordered note."].find('.//' + qn("w:bottom")).get(qn("w:val")) == "double"
        ordinary = next(node for text, node in paragraphs.items() if "Ordinary note." in text)
        assert ordinary.find('.//' + qn("w:pBdr")) is None


def test_first_paragraph_style_can_clear_body_border(tmp_path: Path, rust_executable: Path) -> None:
    """A semantic first-paragraph override must outrank the body border in rendered HTML."""
    source = tmp_path / "first.md"
    source.write_text('''---
papperSettings:
  docxStyle:
    Body Text:
      borders: "1pt solid #005588"
    First Paragraph:
      borders: none
---
# Heading

First paragraph.

Second paragraph.
''', encoding="utf-8")
    output = tmp_path / "first.html"
    result = build(rust_executable, source, output)
    assert result.returncode == 0, result.stdout + result.stderr
    document = html.fromstring(output.read_text(encoding="utf-8"))
    first, second = document.xpath("//main/p")
    assert "border-bottom: none" in first.get("style")
    assert "border-bottom: 1pt solid #005588" in second.get("style")
