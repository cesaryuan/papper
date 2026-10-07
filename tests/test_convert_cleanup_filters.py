"""Verify Convert caption pairing and TOC cleanup with real Pandoc output.

Markdown fixtures snapshot each filter independently, including nested content
and near misses. A generated Word document exercises the public Convert chain
to protect caption text, bookmark references, extracted images and dimensions.
"""

from __future__ import annotations

import json
import os
import subprocess
from pathlib import Path

import pytest

from native_support import ROOT, native_pandoc_executable
from snapshot_utils import assert_snapshot


@pytest.mark.parametrize(("case", "filter_name"), [
    ("figure_captions", "detect_figure"),
    ("toc_cleanup", "remove_toc_anchors"),
    ("inline_images", "extract_inline_images"),
    ("table_figures", "detect_figure"),
    ("indented_figures", "detect_figure"),
    ("table_captions", "detect_table"),
    ("figure_prefixes", "detect_figure"),
    ("table_prefixes", "detect_table"),
])
def test_convert_cleanup_filter_snapshot(
    case: str, filter_name: str, tmp_path: Path, snapshot_update: bool,
) -> None:
    """Protect actual imported syntax and content preservation for each new filter."""
    pandoc = native_pandoc_executable()
    assert pandoc is not None, "Convert filters require the retained Pandoc engine"
    fixture = ROOT / "tests/snapshot_cases_convert" / case
    result = subprocess.run(
        [str(pandoc), str(fixture / "input.md"), "--from=markdown", "--to=markdown", "--wrap=none",
         "--lua-filter", str(ROOT / "pandoc/filters/convert" / f"{filter_name}.lua")],
        cwd=tmp_path, capture_output=True, text=True, encoding="utf-8", timeout=30,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert_snapshot(result.stdout, fixture / "snapshots-content" / f"{case}.md", update=snapshot_update)


def test_inline_image_filter_creates_figure_ast(tmp_path: Path) -> None:
    """Verify the extracted image is a Figure block before Markdown writing flattens it."""
    pandoc = native_pandoc_executable()
    assert pandoc is not None, "Convert filters require the retained Pandoc engine"
    fixture = ROOT / "tests/snapshot_cases_convert/inline_images"
    result = subprocess.run(
        [str(pandoc), str(fixture / "input.md"), "--from=markdown", "--to=json",
         "--lua-filter", str(ROOT / "pandoc/filters/convert/extract_inline_images.lua")],
        cwd=tmp_path, capture_output=True, text=True, encoding="utf-8", timeout=30,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    blocks = json.loads(result.stdout)["blocks"]
    assert blocks[0]["t"] == "Para"
    assert "".join(inline["c"] for inline in blocks[0]["c"]) == "对桥隧基本信息进行预处理，建立区域"
    assert blocks[1]["t"] == "Figure"
    image = blocks[1]["c"][2][0]["c"][0]
    assert image["t"] == "Image"
    assert image["c"][2][0] == "media/image16.png"
    assert dict(image["c"][0][2]) == {"width": "5.78125in", "height": "3.120138888888889in"}


def test_paired_image_crossrefs_snapshot(tmp_path: Path, snapshot_update: bool) -> None:
    """Protect bracketed figure/equation references and unknown bookmark links."""
    pandoc = native_pandoc_executable()
    assert pandoc is not None, "Convert filters require the retained Pandoc engine"
    fixture = ROOT / "tests/snapshot_cases_convert/paired_crossrefs"
    result = subprocess.run(
        [str(pandoc), str(fixture / "input.md"), "--from=markdown", "--to=markdown", "--wrap=none",
         "--lua-filter", str(ROOT / "pandoc/filters/convert/detect_figure.lua"),
         "--lua-filter", str(ROOT / "pandoc/filters/convert/crossrefs.lua")],
        cwd=tmp_path, capture_output=True, text=True, encoding="utf-8", timeout=30,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "[@fig:_Ref202795969]" in result.stdout
    assert "[@fig:_RefSingle]" in result.stdout
    assert "[@eq:_RefEquation]" in result.stdout
    assert '[First bookmark](#_RefFirst)' in result.stdout
    assert '[Inline picture](#_RefInline)' in result.stdout
    assert_snapshot(
        result.stdout, fixture / "snapshots-content/paired_crossrefs.md", update=snapshot_update,
    )


@pytest.mark.parametrize("image_layout", ["standalone", "inline", "table", "indented"])
def test_convert_pairs_word_captions_and_removes_nested_toc_links(
    tmp_path: Path, rust_executable: Path, snapshot_update: bool, image_layout: str,
) -> None:
    """Pair real Word captions, clean TOC links and resolve inbound figure references."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    document = Document()
    image_path = str(ROOT / "template/examples/images/single-figure-example.png")
    if image_layout == "inline":
        paragraph = document.add_paragraph("对桥隧基本信息进行预")
        paragraph.add_run().add_picture(image_path, width=Inches(3))
        paragraph.add_run("处理，建立区域")
        caption = document.add_paragraph()
    elif image_layout == "table":
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).paragraphs[0].add_run().add_picture(image_path, width=Inches(1))
        caption = table.cell(1, 0).paragraphs[0]
    else:
        document.add_picture(image_path, width=Inches(1))
        if image_layout == "indented":
            # Word left indentation imports as a BlockQuote, outside its caption.
            document.paragraphs[-1].paragraph_format.left_indent = Inches(1)
        caption = document.add_paragraph()
    # Nested Word bookmarks reproduce the user's nested Span pattern in Pandoc.
    for bookmark_id, name in [(1, "_Ref181174213"), (2, "_Toc241697898")]:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id))
        start.set(qn("w:name"), name)
        caption._p.append(start)
    for bookmark_id in (2, 1):
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id))
        caption._p.append(end)
    caption.add_run("图1‑11 区域桥隧网络脆弱节点识别结果")
    toc = document.add_paragraph()
    # Markdown readers reject nested links; real nested Word links are needed
    # here to verify that removing both wrappers preserves the page number.
    outer = OxmlElement("w:hyperlink")
    outer.set(qn("w:anchor"), "_Toc241697943")
    text_run = document.add_paragraph("图2‑35 各时段核密度估计 ")._p[0]
    document._body._body.remove(text_run.getparent())
    outer.append(text_run)
    inner = OxmlElement("w:hyperlink")
    inner.set(qn("w:anchor"), "_Toc241697943")
    page_run = document.add_paragraph("78")._p[0]
    document._body._body.remove(page_run.getparent())
    inner.append(page_run)
    outer.append(inner)
    toc._p.append(outer)
    reference = document.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "_Ref181174213")
    run = document.add_paragraph("Figure reference")._p[0]
    document._body._body.remove(run.getparent())
    link.append(run)
    reference._p.append(link)
    source = tmp_path / "captions.docx"
    document.save(source)
    output = tmp_path / "converted"
    result = subprocess.run(
        [str(rust_executable), "convert", str(source), "-o", str(output)], cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "home")},
        capture_output=True, text=True, encoding="utf-8", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    markdown = (output / "captions.md").read_text(encoding="utf-8")
    assert "_Toc" not in markdown
    assert "![图1‑11 区域桥隧网络脆弱节点识别结果](media/" in markdown
    assert "#fig:_Ref181174213" in markdown
    if image_layout == "inline":
        assert "对桥隧基本信息进行预处理，建立区域\n\n![" in markdown
        assert 'width="3.0in"' in markdown
    else:
        assert 'width="1.0in"' in markdown
    assert "图2‑35 各时段核密度估计 78" in markdown
    assert "[@fig:_Ref181174213]" in markdown
    assert len(list((output / "media").glob("*.png"))) == 1
    snapshot_name = "word_inline_cleanup.md" if image_layout == "inline" else "word_cleanup.md"
    assert_snapshot(
        markdown, ROOT / "tests/snapshot_cases_convert/word_cleanup/snapshots-content" / snapshot_name,
        update=snapshot_update,
    )


def test_convert_pairs_word_table_caption(
    tmp_path: Path, rust_executable: Path, snapshot_update: bool,
) -> None:
    """Protect table titles, merged cells and bookmarks through the public DOCX importer."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("An introduction that remains outside the caption and table.")
    caption = document.add_paragraph()
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "1")
    start.set(qn("w:name"), "_Ref202795830")
    caption._p.append(start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "1")
    caption._p.append(end)
    caption.add_run("表2‑1 武汉市")
    caption.add_run("车辆类型组成").bold = True
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "车型组成"
    table.cell(0, 2).text = "平均客货比/%"
    table.cell(1, 0).text = "两轴车"
    table.cell(1, 1).text = "83.51"
    table.cell(1, 2).merge(table.cell(2, 2)).text = "8：1"
    table.cell(2, 0).text = "三轴车"
    table.cell(2, 1).text = "2.17"
    document.add_paragraph("Text following the table remains a separate paragraph.")
    # Pandoc drops unreferenced hidden Word bookmarks; an actual hyperlink
    # keeps the _Ref anchor available so its preservation can be verified.
    reference = document.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "_Ref202795830")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Table reference"
    run.append(text)
    link.append(run)
    reference._p.append(link)
    source = tmp_path / "table-captions.docx"
    document.save(source)
    output = tmp_path / "converted"
    result = subprocess.run(
        [str(rust_executable), "convert", str(source), "-o", str(output)], cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "home")},
        capture_output=True, text=True, encoding="utf-8", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    markdown = (output / "table-captions.md").read_text(encoding="utf-8")
    assert markdown.count("表2‑1 武汉市**车辆类型组成**") == 1
    assert ": []{#_Ref202795830 .anchor}表2‑1 武汉市**车辆类型组成**" in markdown
    assert "[Table reference](#_Ref202795830)" in markdown
    assert_snapshot(
        markdown, ROOT / "tests/snapshot_cases_convert/table_captions/snapshots-content/word_table_captions.md",
        update=snapshot_update,
    )


@pytest.mark.parametrize("field_layout", ["spanning", "nested_mathtype"])
def test_convert_preserves_numbered_paragraphs_inside_word_fields(
    tmp_path: Path, rust_executable: Path, field_layout: str,
) -> None:
    """Keep heading/list text in place instead of leaking field contents into later blocks."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()

    def field_run(paragraph, kind: str, instruction: str = "") -> None:
        """Append the Word field delimiters used by spanning and nested MathType fields."""
        run = paragraph.add_run()._r
        if instruction:
            text = OxmlElement("w:instrText")
            text.text = instruction
            run.append(text)
        else:
            marker = OxmlElement("w:fldChar")
            marker.set(qn("w:fldCharType"), kind)
            run.append(marker)

    intro = document.add_paragraph("Before the field. ")
    field_run(intro, "begin")
    field_run(intro, "", " GOTOBUTTON ZEqnNum1 ")
    if field_layout == "nested_mathtype":
        # MathType nests a REF result inside its GOTOBUTTON instruction. Pandoc
        # can retain that field state after its delimiters, as in convert-test.docx.
        field_run(intro, "begin")
        field_run(intro, "", " REF ZEqnNum1 ")
    field_run(intro, "separate")
    intro.add_run("Equation reference")
    if field_layout == "nested_mathtype":
        field_run(intro, "end")
        field_run(intro, "end")

    heading = document.add_heading("Multi-head attention", level=4)
    numbering = heading._p.get_or_add_pPr().get_or_add_numPr()
    numbering.get_or_add_ilvl().val = 0
    numbering.get_or_add_numId().val = 1
    document.add_paragraph("The following steps retain their own paragraphs.")
    steps = ["Draw a sample.", "Compute responses.", "Collect data.", "Update parameters.", "Start the next stage."]
    for text in steps:
        paragraph = document.add_paragraph(text)
        numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_ilvl().val = 0
        numbering.get_or_add_numId().val = 7
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "The table contains only this text."
    tail = document.add_paragraph("After the steps.")
    if field_layout == "spanning":
        field_run(tail, "end")
    source = tmp_path / "numbered-fields.docx"
    document.save(source)
    output = tmp_path / "converted"
    result = subprocess.run(
        [str(rust_executable), "convert", str(source), "-o", str(output)], cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "home")},
        capture_output=True, text=True, encoding="utf-8", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    markdown = (output / "numbered-fields.md").read_text(encoding="utf-8")
    assert "#### Multi-head attention\n" in markdown
    for index, text in enumerate(steps, start=1):
        assert f"{index}.  {text}\n" in markdown
        assert markdown.count(text) == 1
    assert markdown.index(steps[-1]) < markdown.index("The table contains only this text.")
    assert "The following steps retain their own paragraphs." in markdown
