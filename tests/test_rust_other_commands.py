"""Compare remaining native CLI commands using isolated projects and real output.

Product operations run through the Rust executable; cleanup contracts create and remove their own
temporary projects, and conversion checks real MathType OLE and extracted media.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest

from test_build_snapshots import ROOT, copy_case
from native_support import native_pandoc_executable
from test_rust_docx_contract import rust_mathtype_converter


def _environment(home: Path) -> dict[str, str]:
    """Confine all native mutable user state to the test's temporary root."""
    return {**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(home),
            "UV_CACHE_DIR": str(ROOT / ".uv-cache")}


def _native(executable: Path, arguments: list[str], project: Path, home: Path) -> subprocess.CompletedProcess[str]:
    """Execute an actual native command with captured diagnostics and bounded runtime."""
    return subprocess.run([str(executable), *arguments], cwd=project, env=_environment(home),
                          capture_output=True, text=True, encoding="utf-8", timeout=90)


def _tree(directory: Path) -> dict[str, bytes]:
    """Compare complete initialized/exported files rather than template path wiring."""
    return {path.relative_to(directory).as_posix(): path.read_bytes()
            for path in directory.rglob("*") if path.is_file()}


def _state(home: Path, project: Path) -> Path:
    """Locate existing project state using the public canonical-path identity contract."""
    key = hashlib.sha256(os.path.normcase(str(project.resolve())).encode()).hexdigest()[:20]
    return home / "projects" / key


@pytest.mark.parametrize("language", [None, "zh-cn"])
def test_native_init_preserves_localized_templates_and_user_edits(
    tmp_path: Path, rust_executable: Path, language: str | None,
) -> None:
    """Match real localized templates and merge missing guidance without replacing edited files."""
    native = tmp_path / "native"
    native.mkdir()
    arguments = ["init", *( ["--lang", language] if language else [] )]
    actual = _native(rust_executable, arguments, native, tmp_path / "home")
    assert actual.returncode == 0, actual.stdout + actual.stderr
    manuscript = "manuscript-cn.md" if language else "manuscript.md"
    reply = "reply_to_reviewers-cn.md" if language else "reply_to_reviewers.md"
    assert (native / "manuscript.md").read_bytes() == (ROOT / "template" / manuscript).read_bytes()
    assert (native / "reply_to_reviewers.md").read_bytes() == (ROOT / "template" / reply).read_bytes()
    assert (native / "style.yml").read_bytes() == (ROOT / "template/style-project.yml").read_bytes()
    edited = native / ".agents/word-manuscript-fix/SKILL.md"
    edited.write_text("local project guidance\n", encoding="utf-8")
    (native / "manuscript.md").write_text("user manuscript\n", encoding="utf-8")
    missing = native / ".agents/manuscript-review/SKILL.md"
    missing.unlink()
    merged = _native(rust_executable, [*arguments, "--merge"], native, tmp_path / "home")
    assert merged.returncode == 0, merged.stdout + merged.stderr
    assert edited.read_text(encoding="utf-8") == "local project guidance\n"
    assert (native / "manuscript.md").read_text(encoding="utf-8") == "user manuscript\n"
    assert missing.read_bytes() == (ROOT / "template/.agents/manuscript-review/SKILL.md").read_bytes()
    rejected = _native(rust_executable, arguments, native, tmp_path / "home")
    assert rejected.returncode != 0
    assert (native / "manuscript.md").read_text(encoding="utf-8") == "user manuscript\n"


def test_native_clean_removes_project_caches_and_preserves_other_projects(
    tmp_path: Path, rust_executable: Path,
) -> None:
    """Remove output and all project caches while preserving manuscripts and another project's data."""
    home = tmp_path / "home"
    first = tmp_path / "first"
    second = tmp_path / "second"
    first.mkdir()
    second.mkdir()
    (first / "manuscript.md").write_text("user content", encoding="utf-8")
    (first / "output/nested").mkdir(parents=True)
    (first / "output/nested/generated.docx").write_bytes(b"generated")
    (first / ".pandoc-cache").mkdir()
    (first / ".pandoc-cache/legacy.bin").write_bytes(b"legacy cache")
    for project in (first, second):
        state = _state(home, project)
        (state / "cache").mkdir(parents=True)
        (state / "cache/result.bin").write_bytes(project.name.encode())
    work = _state(home, first) / "work"
    (work / "rust-v1").mkdir(parents=True)
    (work / "rust-v1/input.md").write_text("transient", encoding="utf-8")
    (work / "legacy.tmp").write_bytes(b"transient legacy work")
    cleaned = _native(rust_executable, ["clean"], first, home)
    assert cleaned.returncode == 0, cleaned.stdout + cleaned.stderr
    assert not (first / "output").exists()
    assert not work.exists()
    assert not _state(home, first).exists()
    assert not (first / ".pandoc-cache").exists()
    assert (_state(home, second) / "cache/result.bin").read_bytes() == b"second"
    assert (first / "manuscript.md").read_text(encoding="utf-8") == "user content"
    repeated = _native(rust_executable, ["clean"], first, home)
    assert repeated.returncode == 0, repeated.stdout + repeated.stderr
    removed = _native(rust_executable, ["distclean"], first, home)
    assert removed.returncode != 0
    assert "unrecognized subcommand" in removed.stderr


@pytest.mark.parametrize("target", [".", ".."])
def test_native_clean_rejects_project_or_ancestor_without_deleting_user_data(
    tmp_path: Path, rust_executable: Path, target: str,
) -> None:
    """Reject destructive custom outputs before any existing project file is removed."""
    project = tmp_path / "isolated-project"
    project.mkdir()
    manuscript = project / "manuscript.md"
    manuscript.write_bytes(b"must remain intact")
    result = _native(rust_executable, ["clean", "--output-dir", target], project, tmp_path / "home")
    assert result.returncode != 0
    assert "Refusing to clean unsafe" in result.stderr
    assert manuscript.read_bytes() == b"must remain intact"


@pytest.mark.parametrize("target", ["json", "latex"])
def test_native_json_and_latex_preserve_crossrefs_and_relative_resources(
    tmp_path: Path, rust_executable: Path, target: str,
) -> None:
    """Publish resolved references and portable figure resources from an isolated manuscript."""
    native = tmp_path / "native"
    source = ROOT / "tests/snapshot_cases/crossrefs"
    copy_case(source, native)
    extension = "json" if target == "json" else "tex"
    arguments = ["build", target, "-m", "crossrefs.md", "-o", f"output/result.{extension}"]
    actual = _native(rust_executable, arguments, native, tmp_path / "home")
    assert actual.returncode == 0, actual.stdout + actual.stderr
    native_bytes = (native / f"output/result.{extension}").read_bytes()
    if target == "json":
        ast = json.loads(native_bytes)
        assert ast["blocks"] and ast["meta"]
        assert "Unknown reference" not in json.dumps(ast)
        links = []

        def collect_links(value) -> None:
            """Read crossref destinations from the public Pandoc JSON document."""
            if isinstance(value, dict):
                if value.get("t") == "Link":
                    links.append(value["c"][2][0])
                for child in value.values():
                    collect_links(child)
            elif isinstance(value, list):
                for child in value:
                    collect_links(child)

        collect_links(ast["blocks"])
        assert {"#fig:trend", "#tbl:values", "#eq:circle", "#sec:overview"} <= set(links)
    else:
        text = native_bytes.decode("utf-8")
        assert r"\begin{document}" in text and r"\end{document}" in text
        assert "Unknown reference" not in text
        assert any(data == (source / "figure.svg").read_bytes()
                   for data in _tree(native / "output/latex").values())


def test_native_convert_recovers_mathtype_tex_and_extracts_media_without_overwriting_user_files(
    tmp_path: Path, rust_executable: Path, rust_mathtype_converter: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Round-trip real editable MathType OLE and embedded media through the product convert command."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Inches

    project = tmp_path / "project"
    project.mkdir()
    marked = project / "marked.docx"
    document = Document()
    paragraph = document.add_paragraph("A real inline formula: ")
    marker = paragraph.add_run("MTLATEX:inline:x^2+y^2")
    marker.font.hidden = True
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x^2+y^2"
    run.append(text)
    math.append(run)
    paragraph._p.append(math)
    # A real bundled PNG tests media extraction alongside an editable OLE object.
    image = ROOT / "template/examples/images/single-figure-example.png"
    document.add_picture(str(image), width=Inches(1))
    document.save(marked)
    metadata = project / "metadata.json"
    metadata.write_text(json.dumps({"pmt_settings": {"values": {"mathtypeConversionMethod": "rust"}, "provided": [], "pandoc_metadata": {}, "reply": None},
                                    "pandoc_metadata": {}, "has_yaml_header": False}), encoding="utf-8")
    source = project / "paper.docx"
    # Exercise the Lua map's Unicode temporary path: Windows ANSI getenv formerly
    # corrupted that path despite successful native equation decoding.
    state = tmp_path / "中文 状态"
    temporary = tmp_path / "中文 临时"
    temporary.mkdir()
    for variable in ("TMP", "TEMP", "TMPDIR"):
        monkeypatch.setenv(variable, str(temporary))
    converted = subprocess.run([str(rust_mathtype_converter), str(marked), str(source), str(metadata), str(project)],
                               env=_environment(state), capture_output=True, text=True, encoding="utf-8", timeout=90)
    assert converted.returncode == 0, converted.stdout + converted.stderr
    assert json.loads(converted.stdout)["converted"] == 1
    with ZipFile(source) as archive:
        assert archive.read("word/embeddings/mathtype_formula_1.bin").startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
    destination = project / "native-converted"
    destination.mkdir()
    (destination / "user-notes.txt").write_bytes(b"existing user notes")
    result = _native(rust_executable, ["convert", str(source), "-o", str(destination)], project, state)
    assert result.returncode == 0, result.stdout + result.stderr
    markdown = (destination / "paper.md").read_text(encoding="utf-8")
    assert "$x^2+y^2$" in markdown
    assert "MTLATEX:" not in markdown
    assert "](media/" in markdown
    assert (destination / "user-notes.txt").read_bytes() == b"existing user notes"
    media = [path for path in (destination / "media").iterdir() if path.is_file()]
    assert any(path.read_bytes() == image.read_bytes() for path in media)
    # The standalone Lua path must decode in one native batch without a Python
    # interpreter. With no predecoded map, Pandoc invokes PAPPER_EXECUTABLE.
    pandoc = native_pandoc_executable()
    assert pandoc is not None
    standalone = project / "standalone"
    standalone.mkdir()
    filter_environment = _environment(tmp_path / "standalone-home")
    filter_environment.pop("MATHTYPE_LATEX_MAP", None)
    filter_environment["PAPPER_EXECUTABLE"] = str(rust_executable)
    filter_environment["PATH"] = str(pandoc.parent) + os.pathsep + os.environ.get("SystemRoot", r"C:\Windows") + r"\System32"
    result = subprocess.run([str(pandoc), str(source), "--from=docx", "--to=markdown", "--wrap=none",
                             "--lua-filter", str(ROOT / "pandoc/filters/convert/mtef_parser.lua"),
                             "--extract-media=.", "--output", "paper.md"], cwd=standalone, env=filter_environment,
                            capture_output=True, text=True, encoding="utf-8", timeout=45)
    assert result.returncode == 0, result.stdout + result.stderr
    assert "$x^2+y^2$" in (standalone / "paper.md").read_text(encoding="utf-8")

    # Mock only the external decoder result to reproduce an empty equation.
    # A standalone preview is visited by both display and inline recovery; its
    # bytes must survive extraction and its path must be reported exactly once.
    import hashlib

    with ZipFile(source) as archive:
        previews = [name for name in archive.namelist() if name.endswith(".wmf")]
        assert len(previews) == 1
        preview = previews[0]
        preview_bytes = archive.read(preview)
        empty_map = {
            hashlib.sha1(archive.read(name)).hexdigest(): r"\[\]"
            for name in archive.namelist() if name.endswith((".wmf", ".bin"))
        }
    mapping = project / "empty-equations.json"
    mapping.write_text(json.dumps(empty_map), encoding="utf-8")
    filter_environment["MATHTYPE_LATEX_MAP"] = str(mapping)
    result = subprocess.run(
        [str(pandoc), str(source), "--from=docx", "--to=markdown", "--wrap=none",
         "--lua-filter", str(ROOT / "pandoc/filters/convert/mtef_parser.lua"),
         "--extract-media=.", "--output", "empty.md"], cwd=standalone,
        env=filter_environment, capture_output=True, text=True, encoding="utf-8", timeout=45,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    relative_preview = preview.removeprefix("word/")
    assert relative_preview in (standalone / "empty.md").read_text(encoding="utf-8")
    assert (standalone / relative_preview).read_bytes() == preview_bytes
    warnings = [line for line in result.stderr.splitlines() if "keeping preview image" in line]
    assert len(warnings) == 1
    assert relative_preview in warnings[0] and "empty equation" in warnings[0]


def test_native_build_uses_environment_defaults_and_explicit_cli_overrides(
    tmp_path: Path, rust_executable: Path,
) -> None:
    """Honor runtime source/style/output defaults while explicit source/style/output flags win."""
    project = tmp_path / "project"
    project.mkdir()
    for filename, title in [("environment.md", "Environment source"), ("explicit.md", "Explicit source")]:
        (project / filename).write_text(f"---\ntitle: {title}\n---\n\nA paragraph.\n", encoding="utf-8")
    for filename, subtitle in [("environment.yml", "Environment style"), ("explicit.yml", "Explicit style")]:
        (project / filename).write_text(f"pandocMetadata:\n  subtitle: {subtitle}\n", encoding="utf-8")
    environment = {**_environment(tmp_path / "home"), "PMT_MANUSCRIPT_FILE": "environment.md",
                   "PMT_STYLE_FILE": "environment.yml", "PMT_PROJECT_NAME": "environment-project",
                   "PMT_JSON_DIR": "custom-json"}
    defaults = subprocess.run([str(rust_executable), "build", "json"], cwd=project, env=environment,
                              capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert defaults.returncode == 0, defaults.stdout + defaults.stderr
    default_output = project / "custom-json/environment-project.json"
    ast = json.loads(default_output.read_bytes())

    def meta_text(field: str) -> str:
        """Read visible metadata text from the actual Pandoc AST."""
        return " ".join(item["c"] for item in ast["meta"][field]["c"] if item["t"] == "Str")

    assert meta_text("title") == "Environment source"
    assert meta_text("subtitle") == "Environment style"
    original = default_output.read_bytes()
    explicit = subprocess.run([str(rust_executable), "build", "json", "-m", "explicit.md",
                               "--style-file", "explicit.yml", "-o", "chosen.json"], cwd=project, env=environment,
                              capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert explicit.returncode == 0, explicit.stdout + explicit.stderr
    ast = json.loads((project / "chosen.json").read_bytes())
    assert meta_text("title") == "Explicit source"
    assert meta_text("subtitle") == "Explicit style"
    assert default_output.read_bytes() == original
    assert not (project / "output").exists()


def test_native_mathtype_explicit_work_directory_retains_reviewed_parts(
    tmp_path: Path, rust_executable: Path,
) -> None:
    """Keep requested intermediate formulas inspectable while custom output names preserve source identity."""
    from io import BytesIO
    import olefile

    project = tmp_path / "project"
    project.mkdir()
    (project / "paper.md").write_text("---\ntitle: Formula inspection\nmathtype: true\nmathtypeConversionMethod: rust\n---\n\nAn inline $x_1$ formula.\n", encoding="utf-8")
    debug = tmp_path / "reviewed-equations"
    environment = {**_environment(tmp_path / "home"), "PMT_MATHTYPE_WORK_DIR": str(debug),
                   "PMT_PROJECT_NAME": "unused-environment-name"}
    output = project / "renamed-output.docx"
    result = subprocess.run([str(rust_executable), "build", "docx", "-m", "paper.md", "-o", str(output)],
                            cwd=project, env=environment, capture_output=True, text=True, encoding="utf-8", timeout=45)
    assert result.returncode == 0, result.stdout + result.stderr
    with ZipFile(output) as archive:
        ole = archive.read("word/embeddings/mathtype_formula_1.bin")
        preview = archive.read("word/media/mathtype_formula_1.wmf")
        assert b"MTLATEX:" not in archive.read("word/document.xml")
    parts = debug / "paper"
    assert (parts / "eq_001.tex").read_text(encoding="utf-8") == "$x_1$"
    assert (parts / "eq_001.ole.bin").read_bytes() == ole
    assert (parts / "eq_001.wmf").read_bytes() == preview
    with olefile.OleFileIO(BytesIO(ole)) as compound:
        native_stream = compound.openstream("Equation Native").read()
    assert (parts / "eq_001.mtef.bin").read_bytes() == native_stream[28:]
    assert json.loads((parts / "eq_001.json").read_bytes())["mathtype"]["baseline_from_bottom_pt"] >= 0
    with ZipFile(debug / "paper.marked.docx") as archive:
        assert b"MTLATEX:inline:x_1" in archive.read("word/document.xml")
    assert not (debug / "renamed-output").exists()
