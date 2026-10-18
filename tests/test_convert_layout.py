"""Exercise Word layout import through public Convert and DOCX rebuilds.

Existing Convert snapshots contain no page stories or explicit pagination. These
cases use a checked-in layout DOCX and targeted generated inputs to protect
layout boundaries, dynamic fields, story resources, and inheritance
by inspecting actual Markdown and reopened DOCX output.
"""

from __future__ import annotations

import os
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest
import yaml
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from native_support import ROOT
from snapshot_utils import assert_snapshot


def add_field(paragraph, instruction: str, result: str, *, simple: bool = False) -> None:
    """Author real Word fields, including instructions split across separate runs."""
    if simple:
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = result
        run.append(text)
        field.append(run)
        paragraph._p.append(field)
        return
    for kind in ("begin", "instruction", "separate", "result", "end"):
        run = paragraph.add_run()
        run.bold = True
        if kind == "instruction":
            for chunk in (instruction[:3], instruction[3:]):
                text = OxmlElement("w:instrText")
                text.set(qn("xml:space"), "preserve")
                text.text = chunk
                paragraph.add_run()._r.append(text)
        elif kind == "result":
            run.text = result
        else:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), kind)
            run._r.append(field)


def run_papper(executable: Path, directory: Path, *arguments: str) -> subprocess.CompletedProcess[str]:
    """Run the real CLI with isolated tool state and surface conversion/build failures."""
    result = subprocess.run(
        [str(executable), *arguments], cwd=directory,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(directory / "home")},
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    return result


def test_convert_layout_snapshot_and_rebuild(
    tmp_path: Path, rust_executable: Path, snapshot_update: bool,
) -> None:
    """Protect section order, first/even inheritance, dynamic fields, and story media on rebuild."""
    source = ROOT / "tests/snapshot_cases_convert/page_layout/layout.docx"
    original = source.read_bytes()
    with ZipFile(source) as archive:
        source_logo = archive.read("word/media/image1.png")
    output = tmp_path / "converted"
    (output / "media").mkdir(parents=True)
    existing = output / "media/image1.png"
    existing.write_bytes(b"previous imported image")
    run_papper(rust_executable, tmp_path, "convert", str(source), "-o", str(output))
    markdown = (output / "layout.md").read_text(encoding="utf-8")
    assert source.read_bytes() == original
    assert existing.read_bytes() == b"previous imported image"
    assert "PAPPERLAYOUT" not in markdown and "papper-convert-layout" not in markdown
    assert markdown.count("Layout specimen") == 1
    assert_snapshot(
        markdown, ROOT / "tests/snapshot_cases_convert/page_layout/snapshots-content/layout.md",
        update=snapshot_update,
    )
    rebuilt = output / "rebuilt.docx"
    run_papper(rust_executable, output, "build", "docx", "layout.md", "--no-mathtype", "-o", str(rebuilt))
    reopened = Document(rebuilt)
    assert len(reopened.sections) == 5, "Initial settings must not create an extra empty section"
    assert [section.start_type for section in reopened.sections[1:]] == [
        WD_SECTION_START.CONTINUOUS, WD_SECTION_START.ODD_PAGE,
        WD_SECTION_START.EVEN_PAGE, WD_SECTION_START.NEW_PAGE,
    ]
    assert [section.different_first_page_header_footer for section in reopened.sections] == [True, False, True, False, False]
    assert reopened.settings.odd_and_even_pages_header_footer
    first, second, third, fourth, fifth = reopened.sections
    assert first.header.paragraphs[0].text == 'Header "quoted" & italic Project'
    assert first.header.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert first.header.tables[0].cell(1, 1).text == "Draft"
    assert any(paragraph.text == "Chapter: Chapter 4" for paragraph in first.header.paragraphs)
    assert [node.text for node in first.header._element.iter(qn("w:instrText"))] == ["STYLEREF 1"]
    assert list(first.header._element.iter(qn("m:oMath")))
    assert first.footer.paragraphs[0].text == "Page I of 12"
    assert first.first_page_header.paragraphs[0].text == third.first_page_header.paragraphs[0].text == "Opening header"
    assert first.first_page_footer.paragraphs[0].text == ""
    assert second.footer.paragraphs[0].text == third.footer.paragraphs[0].text == ""
    assert second.even_page_footer.paragraphs[0].text == "Even footer"
    assert third.header.paragraphs[0].text == "New normal header"
    assert fourth.header.paragraphs[0].text == fourth.even_page_header.paragraphs[0].text == ""
    assert fifth.footer.paragraphs[0].text == "Final - a"
    assert second._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) == "3"
    assert third._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) is None
    assert third.header_distance.pt == 40 and third.footer_distance.pt == 24
    assert sum(node.get(qn("w:type")) == "page" for node in reopened.element.iter(qn("w:br"))) == 2
    assert [field.text for field in first.footer._element.iter(qn("w:instrText"))] == [
        r" PAGE \* ROMAN ", "NUMPAGES",
    ]
    assert [field.text for field in reopened.element.iter(qn("w:instrText"))] == [
        " PAGE ", r" PAGE \* roman ",
    ]
    with ZipFile(rebuilt) as archive:
        assert any(archive.read(name) == source_logo
                   for name in archive.namelist() if name.startswith("word/media/"))
        assert "Opening header" not in reopened.element.xml
        assert any(b"https://example.org/paper?draft=1&amp;view=full" in archive.read(name)
                   for name in archive.namelist() if name.startswith("word/_rels/header"))


@pytest.mark.parametrize("format_name", ["upperRoman", "decimal"])
def test_convert_inherited_page_field_format(
    tmp_path: Path, rust_executable: Path, format_name: str,
) -> None:
    """Keep inherited PAGE fields dynamic when the following section changes its numbering format."""
    document = Document()
    first = document.sections[0]
    add_field(first.footer.paragraphs[0], " PAGE ", "1", simple=True)
    # Inactive first/even stories must not enable layout modes during import.
    first.first_page_header.paragraphs[0].text = "Dormant first"
    first.even_page_header.paragraphs[0].text = "Dormant even"
    document.add_paragraph("Initial")
    second = document.add_section(WD_SECTION_START.NEW_PAGE)
    number = OxmlElement("w:pgNumType")
    number.set(qn("w:fmt"), format_name)
    second._sectPr.append(number)
    document.add_paragraph("Following")
    source = tmp_path / "inherited.docx"
    document.save(source)
    output = tmp_path / "converted"
    run_papper(rust_executable, tmp_path, "convert", str(source), "-o", str(output))
    markdown = (output / "inherited.md").read_text(encoding="utf-8")
    assert "Dormant" not in markdown
    rebuilt = output / "rebuilt.docx"
    run_papper(rust_executable, output, "build", "docx", "inherited.md", "--no-mathtype", "-o", str(rebuilt))
    reopened = Document(rebuilt)
    assert len(reopened.sections) == 2
    assert not reopened.settings.odd_and_even_pages_header_footer
    assert not any(section.different_first_page_header_footer for section in reopened.sections)
    assert reopened.sections[0].footer.paragraphs[0].text == "1"
    assert reopened.sections[1].footer.paragraphs[0].text == ("I" if format_name == "upperRoman" else "1")


@pytest.mark.parametrize("note_kind", ["footnote", "endnote"])
def test_convert_layout_in_tables_and_notes_with_pretty_xml(
    tmp_path: Path, rust_executable: Path, note_kind: str,
) -> None:
    """Retain nested fields/breaks and authored marker-like text in pretty-printed OOXML."""
    document = Document()
    document.add_paragraph("Literal PAPPERLAYOUT1END")
    if note_kind == "endnote":
        number = OxmlElement("w:pgNumType")
        number.set(qn("w:fmt"), "lowerRoman")
        document.sections[0]._sectPr.append(number)
    paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph.add_run("Cell before")
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run("Cell after ")
    add_field(paragraph, " PAGE ", "9", simple=True)
    note_reference = OxmlElement(f"w:{note_kind}Reference")
    note_reference.set(qn("w:id"), "1")
    document.add_paragraph("Note").add_run()._r.append(note_reference)
    source = tmp_path / "nested.docx"
    document.save(source)
    with ZipFile(source) as archive:
        parts = {entry.filename: archive.read(entry) for entry in archive.infolist()}
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    footnotes = etree.fromstring(f'<w:{note_kind}s xmlns:w="{w}"><w:{note_kind} w:id="1"><w:p/></w:{note_kind}></w:{note_kind}s>')
    note_paragraph = footnotes[0][0]
    for field_type in ("begin", "separate", "end"):
        run = etree.SubElement(note_paragraph, qn("w:r"))
        etree.SubElement(run, qn("w:fldChar"), {qn("w:fldCharType"): field_type})
        if field_type == "begin":
            etree.SubElement(run, qn("w:instrText")).text = " PAGE "
        if field_type == "separate":
            etree.SubElement(run, qn("w:t")).text = "9"
    parts[f"word/{note_kind}s.xml"] = etree.tostring(footnotes)
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    etree.SubElement(rels, f"{{{rels.nsmap[None]}}}Relationship", {
        "Id": "rIdLayoutNote", "Target": f"{note_kind}s.xml",
        "Type": f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{note_kind}s",
    })
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
    types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(types, f"{{{types.nsmap[None]}}}Override", {
        "PartName": f"/word/{note_kind}s.xml",
        "ContentType": f"application/vnd.openxmlformats-officedocument.wordprocessingml.{note_kind}s+xml",
    })
    parts["[Content_Types].xml"] = etree.tostring(types)
    parts["word/document.xml"] = etree.tostring(etree.fromstring(parts["word/document.xml"]), pretty_print=True)
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    original = source.read_bytes()
    output = tmp_path / "converted"
    run_papper(rust_executable, tmp_path, "convert", str(source), "-o", str(output))
    markdown = (output / "nested.md").read_text(encoding="utf-8")
    assert "Literal PAPPERLAYOUT1END" in markdown
    assert markdown.count('t="pageNum"') == 2
    assert markdown.count('t="pageBreak"') == 1
    assert "Cell before" in markdown and "Cell after" in markdown and "[^1]:" in markdown
    assert source.read_bytes() == original
    rebuilt = output / "rebuilt.docx"
    run_papper(rust_executable, output, "build", "docx", "nested.md", "--no-mathtype", "-o", str(rebuilt))
    reopened = Document(rebuilt)
    assert len(reopened.sections) == 1
    assert reopened.sections[0].footer.paragraphs[0].text == ""
    expected_instruction = r" PAGE \* roman " if note_kind == "endnote" else " PAGE "
    assert [node.text for node in reopened.element.iter(qn("w:instrText"))] == [expected_instruction]
    with ZipFile(rebuilt) as archive:
        assert [node.text for node in etree.fromstring(archive.read("word/footnotes.xml")).iter(qn("w:instrText"))] == [expected_instruction]


def test_convert_layout_retains_comprehensive_import_contract(
    tmp_path: Path, rust_executable: Path,
) -> None:
    """Protect equation recovery, captions, crossrefs, notes, and media when layout forces a reread."""
    case = ROOT / "tests/snapshot_cases_convert/comprehensive"
    with ZipFile(case / "comprehensive.docx") as archive:
        parts = {entry.filename: archive.read(entry) for entry in archive.infolist()}
    root = etree.fromstring(parts["word/document.xml"])
    body = root.find(qn("w:body"))
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    run.append(page_break)
    paragraph.append(run)
    body.insert(len(body) - 1, paragraph)
    parts["word/document.xml"] = etree.tostring(root)
    source = tmp_path / "comprehensive.docx"
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    output = tmp_path / "converted"
    run_papper(rust_executable, tmp_path, "convert", str(source), "-o", str(output))
    markdown = (output / "comprehensive.md").read_text(encoding="utf-8")
    expected = (case / "snapshots-content/comprehensive.md").read_text(encoding="utf-8")
    header, content = markdown[4:].split("\n---\n", 1)
    original_metadata, original_content = {}, expected
    if expected.startswith("---\n"):
        original_header, original_content = expected[4:].split("\n---\n", 1)
        original_metadata = yaml.safe_load(original_header)
    metadata = yaml.safe_load(header)
    assert metadata.pop("papperSettings")["docxFooter"]["content"] == ""
    assert metadata == original_metadata
    assert content.count('[]{t="pageBreak"}') == 1
    assert content.replace('\n\n[]{t="pageBreak"}\n\n', "\n\n").lstrip("\n") == original_content.lstrip("\n")
    expected_media = {path.relative_to(case / "snapshots-content").as_posix()
                      for path in (case / "snapshots-content/media").rglob("*") if path.is_file()}
    assert {path.relative_to(output).as_posix() for path in (output / "media").rglob("*")
            if path.is_file()} == expected_media
