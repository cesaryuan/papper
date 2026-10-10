"""Verify explicit style selection through the public build command."""

import json
import shutil
import os
import subprocess
from pathlib import Path

import pytest

from native_support import ROOT, native_pandoc_executable, papper_command


@pytest.fixture
def style_project(tmp_path: Path, monkeypatch) -> Path:
    """Create conflicting default styles to expose accidental discovery or merging."""
    source = tmp_path / "chapters"
    source.mkdir()
    (source / "paper.md").write_text(
        "---\nsubtitle: Manuscript\n---\nBody\n", encoding="utf-8"
    )
    (source / "style.yml").write_text(
        "pandocMetadata:\n  title: Local\n  localOnly: true\n", encoding="utf-8"
    )
    (tmp_path / "style.yml").write_text(
        "pandocMetadata:\n  title: Working\n  workingOnly: true\n", encoding="utf-8"
    )
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("PAPPER_HOME", str(tmp_path / "state"))
    monkeypatch.setenv("PAPPER_RESOURCE_ROOT", str(ROOT))
    return tmp_path


def run_build(arguments: list[str]) -> subprocess.CompletedProcess[str]:
    """Run the current Rust CLI and capture user-visible selection errors."""
    return subprocess.run([*papper_command(), *arguments], env=os.environ,
                          capture_output=True, text=True, encoding="utf-8", timeout=45)


@pytest.mark.skipif(
    native_pandoc_executable() is None and
    (shutil.which("pandoc") is None or shutil.which("pandoc-crossref") is None),
    reason="JSON builds require the native engine or standalone Pandoc and crossref",
)
@pytest.mark.parametrize("path_kind", ["relative", "absolute", "explicit_default"])
def test_build_uses_only_explicit_style_for_one_invocation(
    style_project: Path, path_kind: str
) -> None:
    """Select the requested style in output, preserve YAML priority, and restore discovery."""
    selected = style_project / "styles with spaces" / "journal.yml"
    selected.parent.mkdir()
    selected.write_text(
        "pandocMetadata:\n  title: Selected\n  subtitle: Style\n", encoding="utf-8"
    )
    if path_kind == "explicit_default":
        selected = style_project / "style.yml"
        selected.write_text(
            "pandocMetadata:\n  title: Selected\n  subtitle: Style\n", encoding="utf-8"
        )
    style_arg = str(selected if path_kind == "absolute" else selected.relative_to(style_project))
    output = style_project / "selected.json"
    arguments = ["build", "json", "chapters/paper.md"]

    result = run_build([*arguments, "--style-file", style_arg, "-o", str(output)])
    assert result.returncode == 0, result.stdout + result.stderr
    metadata = json.loads(output.read_text(encoding="utf-8"))["meta"]
    assert metadata["title"]["c"] == [{"t": "Str", "c": "Selected"}]
    assert metadata["subtitle"]["c"] == [{"t": "Str", "c": "Manuscript"}]
    assert "localOnly" not in metadata
    assert "workingOnly" not in metadata

    result = run_build([*arguments, "-o", str(output)])
    assert result.returncode == 0, result.stdout + result.stderr
    metadata = json.loads(output.read_text(encoding="utf-8"))["meta"]
    assert metadata["title"]["c"] == [{"t": "Str", "c": "Local"}]
    assert metadata["localOnly"]["c"] is True
    if path_kind != "explicit_default":
        assert metadata["workingOnly"]["c"] is True


@pytest.mark.parametrize(
    ("style_arg", "message"),
    [("missing.yml", "Style file not found"), ("chapters", "Style path is not a file")],
)
def test_build_rejects_invalid_explicit_style(
    style_project: Path, style_arg: str, message: str
) -> None:
    """Report invalid explicit paths instead of silently falling back to default styles."""
    output = style_project / "result.json"

    result = run_build(
        ["build", "json", "chapters/paper.md", "--style-file", style_arg, "-o", str(output)]
    )
    assert result.returncode != 0
    assert message in result.stderr
    assert not output.exists()


@pytest.mark.parametrize("postprocess", ["true", "false"])
@pytest.mark.parametrize("margins", ["null", "{left: 2cm, right: 2cm}"])
def test_docx_styles_are_inherited_from_exported_reference(
    tmp_path: Path, postprocess: str, margins: str,
) -> None:
    """Export styles without building, then inherit them without changing the source reference."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    reference = tmp_path / "journal.docx"
    template = Document(ROOT / "pandoc/manuscript-template/reference-doc.docx")
    template.sections[0].left_margin = Cm(1.23)
    body = template.styles.add_style("Journal Body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = template.styles["Body Text"]
    body.font.size = Pt(9)
    body.font.italic = True
    emphasis = template.styles.add_style("Journal Emphasis", WD_STYLE_TYPE.CHARACTER)
    emphasis.font.color.rgb = RGBColor.from_string("112233")
    template.add_paragraph("Reference content only")
    template.save(reference)
    original = reference.read_bytes()

    source = tmp_path / "paper.md"
    source.write_text(
        '::: {custom-style="Journal Body"}\n'
        'Journal text with [styled text]{custom-style="Journal Emphasis"}.\n:::\n',
        encoding="utf-8",
    )
    (tmp_path / "style.yml").write_text(
        f"docxPageMargins: {margins}\n"
        "docxStyle:\n"
        "  Journal Body:\n"
        "    fontFamily: Times New Roman\n"
        "    fontSize: 17pt\n"
        "    firstLineIndentChars: 2\n"
        "    paragraphSpacing: {before: 7pt, after: 9pt}\n"
        "  Journal Emphasis:\n"
        "    fontSize: 14pt\n"
        "    fontColor: '#0055AA'\n",
        encoding="utf-8",
    )
    output = tmp_path / "paper.docx"
    output.write_bytes(b"Previous manuscript output")
    exported = tmp_path / "exported-reference.docx"
    result = subprocess.run(
        [*papper_command(), "build", "docx", str(source), "--no-mathtype",
         "--reference-doc", str(reference), "--export-reference-doc", str(exported),
         "-o", str(output)],
        cwd=tmp_path,
        env={**os.environ, "PAPPER_HOME": str(tmp_path / "state"),
             "PAPPER_RESOURCE_ROOT": str(ROOT), "PMT_ENABLE_DOCX_POSTPROCESS": postprocess},
        capture_output=True, text=True, encoding="utf-8", timeout=45,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert output.read_bytes() == b"Previous manuscript output"
    assert reference.read_bytes() == original
    result = subprocess.run(
        [*papper_command(), "build", "docx", str(source), "--no-mathtype",
         "--reference-doc", str(exported), "-o", str(output)],
        cwd=tmp_path,
        env={**os.environ, "PAPPER_HOME": str(tmp_path / "state"),
             "PAPPER_RESOURCE_ROOT": str(ROOT), "PMT_ENABLE_DOCX_POSTPROCESS": postprocess},
        capture_output=True, text=True, encoding="utf-8", timeout=45,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    for path in [exported, output]:
        document = Document(path)
        style = document.styles["Journal Body"]
        assert style.font.name == "Times New Roman"
        assert style.font.size.pt == 17
        assert style.font.italic is True
        assert style.base_style.name == "Body Text"
        assert style.paragraph_format.space_before.pt == 7
        assert style.paragraph_format.space_after.pt == 9
        assert style.element.find(f"{qn('w:pPr')}/{qn('w:ind')}").get(qn("w:firstLineChars")) == "200"
        assert document.styles["Journal Emphasis"].font.size.pt == 14
        assert str(document.styles["Journal Emphasis"].font.color.rgb) == "0055AA"
        assert document.sections[0].left_margin.twips == (
            template.sections[0].left_margin.twips if margins == "null" else Cm(2).twips
        )
    manuscript = Document(output)
    paragraph = next(p for p in manuscript.paragraphs if p.text.startswith("Journal text"))
    assert paragraph.style.name == "Journal Body"
    assert next(run for run in paragraph.runs if run.text == "styled text").style.name == "Journal Emphasis"
    assert not any(p.text == "Reference content only" for p in manuscript.paragraphs)


@pytest.mark.parametrize("mode", [pytest.param(None, id="default"), "window", "content", "fixed", "none"])
def test_table_autofit_defaults_are_shared_by_html_and_docx(
    tmp_path: Path, rust_executable: Path, mode: str | None,
) -> None:
    """Verify implicit/explicit table defaults and overrides in output, preserving layout tables."""
    from lxml import etree, html
    from zipfile import ZipFile

    image = (ROOT / "template/examples/images/subfigure-a-example.png").as_posix()
    source = tmp_path / "tables.md"
    source.write_text(f"""---
subfigGrid: true
tableAutofit: content
---
| Default | Value |
|---|---|
| Sample | 1 |

: Default table.

| Content | Value |
|---|---|
| Sample | 2 |

: Content table. {{autofit="content"}}

| Fixed | Value |
|---|---|
| Sample | 3 |

: Fixed table. {{autofit="fixed"}}

| Window | Value |
|---|---|
| Sample | 4 |

: Window table. {{autofit="window"}}

<div id="fig:panels">
![Left.]({image}){{#fig:left width=49%}}
![Right.]({image}){{#fig:right width=49%}}

Panels.
</div>

$$
x=1
$$ {{#eq:example}}
""", encoding="utf-8")
    style = tmp_path / "style.yml"
    style.write_text(f"tableAutofit: {mode}\n" if mode else "", encoding="utf-8")
    mode = mode or "none"
    environment = {**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT)}
    for target in ["html", "docx"]:
        output = tmp_path / f"tables.{target}"
        command = [str(rust_executable), "build", target, "-m", str(source),
                   "--style-file", str(style), "-o", str(output)]
        if target == "docx":
            command.append("--no-mathtype")
        result = subprocess.run(command, cwd=tmp_path, env=environment,
                                capture_output=True, text=True, encoding="utf-8", timeout=45)
        assert result.returncode == 0, result.stdout + result.stderr
        if target == "html":
            document = html.fromstring(output.read_text(encoding="utf-8"))
            for heading, expected in [("Default", mode if mode != "none" else None),
                                      ("Content", "content"), ("Fixed", "fixed"), ("Window", "window")]:
                table = document.xpath(f"//table[.//th='{heading}']")[0]
                assert table.get("data-autofit") == expected
            layout = document.xpath("//figure[contains(@class, 'subfigures')]/table")[0]
            assert layout.get("data-autofit") is None
            assert 'table[data-autofit="window"]' in output.read_text(encoding="utf-8")
        else:
            with ZipFile(output) as archive:
                document = etree.fromstring(archive.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            # Pandoc emits a full-width table here; none/fixed preserve it rather than forcing auto width.
            for heading, kind, width in [("Default", "auto" if mode == "content" else "pct", "0" if mode == "content" else "5000"),
                                         ("Content", "auto", "0"), ("Window", "pct", "5000")]:
                table = document.xpath(f"//w:tbl[.//w:t='{heading}']", namespaces=ns)[0]
                preferred = table.find("w:tblPr/w:tblW", ns)
                assert preferred.get(f"{{{ns['w']}}}type") == kind
                assert preferred.get(f"{{{ns['w']}}}w") == width
                if heading == "Default" and mode != "none":
                    assert table.find("w:tblPr/w:tblLayout", ns).get(f"{{{ns['w']}}}type") == ("fixed" if mode == "fixed" else "autofit")
            fixed = document.xpath("//w:tbl[.//w:t='Fixed']", namespaces=ns)[0]
            assert fixed.find("w:tblPr/w:tblLayout", ns).get(f"{{{ns['w']}}}type") == "fixed"
            layout = document.xpath("//w:tbl[w:tblPr/w:tblStyle/@w:val='TableSubfigure']", namespaces=ns)[0]
            assert layout.find("w:tblPr/w:tblW", ns).get(f"{{{ns['w']}}}w") == "4900"
            assert layout.find("w:tblPr/w:tblLayout", ns).get(f"{{{ns['w']}}}type") == "fixed"
            # Defaults must not resize crossref's three-column equation layout.
            equation = document.xpath("//w:tbl[.//m:oMath]", namespaces={**ns, "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"})[0]
            assert equation.find("w:tblPr/w:tblLayout", ns).get(f"{{{ns['w']}}}type") == "fixed"
            # Walking metadata would put autofit on eqnBlockTemplate and change this generated width.
            assert equation.find("w:tblPr/w:tblW", ns).get(f"{{{ns['w']}}}w") == "4931"
            assert not document.xpath("//w:t[contains(., 'PMT_TABLE_METADATA:')]", namespaces=ns)
