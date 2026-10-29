"""Exercise exported crop pixels and editable DOCX crops through the public CLI.

Existing Convert fixtures have no image crop metadata. Generated shared-image
inputs catch accidental changes to other occurrences, while pixel comparisons
and reopened DOCX parts verify the actual result rather than filter wiring.
"""

from __future__ import annotations

import os
import re
import struct
import subprocess
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from PIL import Image, ImageChops

from native_support import ROOT
from snapshot_utils import assert_snapshot


def run_cli(executable: Path, directory: Path, *arguments: str, postprocess: bool = True) -> subprocess.CompletedProcess[str]:
    """Run real conversion/builds in isolated state and report actionable failures."""
    result = subprocess.run(
        [str(executable), *arguments], cwd=directory,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(directory / "home"),
             "PMT_ENABLE_DOCX_POSTPROCESS": str(postprocess).lower()},
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    return result


def make_source(directory: Path, format_name: str) -> tuple[Path, bytes]:
    """Author distinct crops of one embedded picture in the body, table, header and note."""
    image = Image.new("RGB", (80, 60))
    image.putdata([(x * 3, y * 4, (x + y) % 256) for y in range(60) for x in range(80)])
    stream = BytesIO()
    image.save(stream, format=format_name)
    original = stream.getvalue()
    document = Document()
    for title, paragraph, crop in [
        ('Detail\'s "quoted" & <original>', document.add_paragraph("First "), (.25, .2, .125, .1)),
        ("Uncropped", document.add_paragraph("Plain "), (0, 0, 0, 0)),
        ("Half", document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0], (0, 0, .5, 0)),
        ("Padded", document.add_paragraph("Padding "), (-.25, 0, .25, 0)),
        ("Header", document.sections[0].header.paragraphs[0], (0, 0, 0, .25)),
    ]:
        picture = paragraph.add_run().add_picture(BytesIO(original), width=Inches(2), height=Inches(1))
        picture._inline.docPr.set("title", title)
        picture._inline.docPr.set("descr", "Original description")
        if any(crop):
            rectangle = OxmlElement("a:srcRect")
            for side, value in zip(("l", "t", "r", "b"), crop):
                if value:
                    rectangle.set(side, str(round(value * 100_000)))
            picture._inline.find('.//' + qn("pic:blipFill")).insert(1, rectangle)
    source = directory / "images.docx"
    document.save(source)
    # A real note uses a separate relationship part, so this also catches imports
    # that incorrectly resolve every image against document.xml.rels.
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    body = etree.fromstring(parts["word/document.xml"])
    picture_run = body.find('.//' + qn("w:drawing")).getparent()
    note_picture = etree.fromstring(etree.tostring(picture_run))
    note_picture.find('.//' + qn("wp:docPr")).set("title", "Note")
    crop = note_picture.find('.//' + qn("a:srcRect"))
    crop.attrib.clear()
    crop.set("l", "50000")
    reference = etree.SubElement(body.find(qn("w:body"))[0], qn("w:r"))
    etree.SubElement(reference, qn("w:footnoteReference"), {qn("w:id"): "1"})
    source_xml = etree.tostring(body)
    # Legal single-quoted XML attributes must survive crop transport even when
    # the original title itself contains an apostrophe and double quotes.
    title_match = re.search(rb'title="([^"]*)"', source_xml)
    replacement = b"title='" + title_match[1].replace(b"'", b"&apos;") + b"'"
    parts["word/document.xml"] = source_xml[:title_match.start()] + replacement + source_xml[title_match.end():]
    notes = etree.Element(qn("w:footnotes"), nsmap=body.nsmap)
    note = etree.SubElement(notes, qn("w:footnote"), {qn("w:id"): "1"})
    etree.SubElement(note, qn("w:p")).append(note_picture)
    parts["word/footnotes.xml"] = etree.tostring(notes)
    parts["word/_rels/footnotes.xml.rels"] = parts["word/_rels/document.xml.rels"]
    relationships = etree.fromstring(parts["word/_rels/document.xml.rels"])
    etree.SubElement(relationships, "{" + relationships.nsmap[None] + "}Relationship", {
        "Id": "rIdNotes", "Target": "footnotes.xml",
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
    })
    parts["word/_rels/document.xml.rels"] = etree.tostring(relationships)
    types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(types, "{" + types.nsmap[None] + "}Override", {
        "PartName": "/word/footnotes.xml",
        "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
    })
    parts["[Content_Types].xml"] = etree.tostring(types)
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return source, original


@pytest.mark.parametrize("format_name", ["PNG", "JPEG"])
def test_convert_bakes_independent_crop_pixels(tmp_path: Path, rust_executable: Path, format_name: str) -> None:
    """Export visible pixels, including padding, without changing shared originals or existing files."""
    source, original = make_source(tmp_path, format_name)
    untouched = source.read_bytes()
    output = tmp_path / "converted"
    output.mkdir()
    (output / "media").mkdir()
    preserved = output / "media/image1.png"
    preserved.write_bytes(b"user edited image")
    run_cli(rust_executable, tmp_path, "convert", str(source), "-o", str(output), "--image-crop", "apply", "--offline")
    assert source.read_bytes() == untouched
    assert preserved.read_bytes() == b"user edited image"
    decoded = Image.open(BytesIO(original)).convert("RGBA")
    expected = [decoded.crop(box) for box in [(20, 12, 70, 54), (40, 0, 80, 60), (0, 0, 40, 60), (0, 0, 80, 45)]]
    padded = Image.new("RGBA", (80, 60))
    padded.paste(decoded.crop((0, 0, 60, 60)), (20, 0))
    expected.append(padded)
    media = [path.read_bytes() for path in (output / "media").iterdir() if path != preserved]
    assert original in media, "An uncropped occurrence must retain the exact original file"
    actual = [Image.open(BytesIO(data)).convert("RGBA") for data in media if data != original]
    assert len(actual) == len(expected)
    # Independent JPEG decoders round color conversion slightly differently.
    # PNG stays exact; JPEG allows two levels while still catching incorrect crops.
    tolerance = 2 if format_name == "JPEG" else 0
    for expected_image in expected:
        match = next((image for image in actual if image.size == expected_image.size
                      and max(high for _, high in ImageChops.difference(image, expected_image).getextrema()) <= tolerance), None)
        assert match is not None, f"Missing visible crop {expected_image.size}"
        actual.remove(match)
    markdown = (output / "images.md").read_text(encoding="utf-8")
    assert 'crop-right=' not in markdown and 'PAPPER_IMAGE_CROP:' not in markdown
    for link in re.findall(r"\]\((media/[^ )]+)", markdown):
        assert (output / link).is_file()


@pytest.mark.parametrize(("postprocess", "crop_arguments"), [
    (True, ()), (False, ("--image-crop", "attributes")),
])
def test_crop_attributes_roundtrip_and_snapshot(
    tmp_path: Path, rust_executable: Path, snapshot_update: bool, postprocess: bool, crop_arguments: tuple[str, ...],
) -> None:
    """Keep originals, crops, titles and sizes with default/explicit attribute mode across a real round trip."""
    source, original = make_source(tmp_path, "PNG")
    untouched = source.read_bytes()
    output = tmp_path / "converted"
    run_cli(rust_executable, tmp_path, "convert", str(source), "-o", str(output), *crop_arguments, "--offline")
    markdown = (output / "images.md").read_text(encoding="utf-8")
    assert source.read_bytes() == untouched
    assert [path.read_bytes() for path in (output / "media").iterdir()] == [original]
    assert_snapshot(markdown, ROOT / "tests/snapshot_cases_convert/image_crops/snapshots-content/attributes.md", update=snapshot_update)
    rebuilt = output / "rebuilt.docx"
    run_cli(rust_executable, output, "build", "docx", "images.md", "--no-mathtype", "-o", str(rebuilt), postprocess=postprocess)
    actual = {}
    with ZipFile(rebuilt) as archive:
        assert original in [archive.read(name) for name in archive.namelist() if name.startswith("word/media/")]
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            data = archive.read(name)
            assert b"PAPPER_IMAGE_CROP:" not in data
            for drawing in etree.fromstring(data).iter(qn("w:drawing")):
                properties = drawing.find('.//' + qn("wp:docPr"))
                rectangle = drawing.find('.//' + qn("a:srcRect"))
                actual[properties.get("title", "")] = dict(rectangle.attrib) if rectangle is not None else {}
                extent = drawing.find('.//' + qn("wp:extent"))
                assert int(extent.get("cx")) == pytest.approx(Inches(2), abs=50)
                assert int(extent.get("cy")) == pytest.approx(Inches(1), abs=50)
    assert actual == {
        'Detail\'s "quoted" & <original>': {"l": "25000", "t": "20000", "r": "12500", "b": "10000"},
        "Uncropped": {}, "Half": {"r": "50000"}, "Padded": {"l": "-25000", "r": "25000"},
        "Header": {"b": "25000"}, "Note": {"l": "50000"},
    }


def test_authored_crop_sizes_and_validation(tmp_path: Path, rust_executable: Path) -> None:
    """Infer missing dimensions from the visible crop and reject invalid authored percentages."""
    image = Image.new("RGB", (80, 60), "red")
    image.save(tmp_path / "photo.png")
    markdown = tmp_path / "sizing.md"
    markdown.write_text('''Inline ![Half](photo.png "width only"){width="2in" crop-right="50%"}.

![Half](photo.png "height only"){height="1in" crop-right="50"}

![Half](photo.png "both"){width="2in" height="1in" crop-right="50"}

![Half](photo.png "natural"){crop-right="50"}
''', encoding="utf-8")
    output = tmp_path / "sizing.docx"
    run_cli(rust_executable, tmp_path, "build", "docx", "sizing.md", "--no-mathtype", "-o", str(output))
    dimensions = {}
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    for drawing in root.iter(qn("w:drawing")):
        title = drawing.find('.//' + qn("wp:docPr")).get("title")
        extent = drawing.find('.//' + qn("wp:extent"))
        dimensions[title] = (int(extent.get("cx")) / Inches(1), int(extent.get("cy")) / Inches(1))
    assert dimensions["width only"] == pytest.approx((2, 3), abs=1e-4)
    assert dimensions["height only"] == pytest.approx((2 / 3, 1), abs=1e-4)
    assert dimensions["both"] == pytest.approx((2, 1), abs=1e-4)
    assert dimensions["natural"][0] / dimensions["natural"][1] == pytest.approx(2 / 3, abs=1e-4)
    for invalid in ['crop-right="banana"', 'crop-left="50" crop-right="50"']:
        markdown.write_text(f'![Invalid](photo.png){{{invalid}}}', encoding="utf-8")
        previous = output.read_bytes()
        result = subprocess.run(
            [str(rust_executable), "build", "docx", "sizing.md", "--no-mathtype", "-o", str(output)],
            cwd=tmp_path, env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT)},
            capture_output=True, text=True, encoding="utf-8", timeout=90,
        )
        assert result.returncode != 0
        assert output.read_bytes() == previous, "A crop error must preserve the previous DOCX"


def vector_picture(extension: str, svg_viewport: str = "normal") -> tuple[bytes, list[bytes]]:
    """Author native colored vector shapes with shifted coordinates and content outside the canvas."""
    if extension == "svg":
        viewport = {
            "normal": 'width="80" height="60" viewBox="10 20 80 60" preserveAspectRatio="none"',
            "letterbox": 'width="80" height="60" viewBox="10 20 80 80"',
            "percent": 'width="100%" height="100%" viewBox="10 20 80 60"',
            "physical": 'width="60pt" height="45pt" viewBox="10 20 80 60"',
        }[svg_viewport]
        paths = [
            b'<path d="M -100 -100 H 200 V 200 H -100 Z" fill="cyan"/>',
            b'<path d="M 10 20 H 50 V 50 H 10 Z" fill="red"/>',
            b'<path d="M 50 20 H 90 V 50 H 50 Z" fill="blue"/>',
            b'<path d="M 10 50 H 50 V 80 H 10 Z" fill="green"/>',
            b'<path d="M 50 50 H 90 V 80 H 50 Z" fill="yellow"/>',
        ]
        return (b'<?xml version="1.0" encoding="UTF-8"?>\n'
                b'<!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN" "http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd">\n'
                + f'<svg xmlns="http://www.w3.org/2000/svg" {viewport}>'.encode()
                + b''.join(paths) + b'</svg>'), paths
    # Use 20 logical units per point in WMF, and a 254 DPI reference device
    # in EMF. Both render to the same 80 x 60 image at the test's chosen DPI.
    records = []
    drawings = []
    if extension == "wmf":
        records.extend([
            wmf_record(0x0103, 8), wmf_record(0x020b, -200, 400),
            wmf_record(0x020c, 1200, 1600), wmf_record(0x02fa, 5, 0, 0, 0, 0),
            wmf_record(0x012d, 0),
        ])
    else:
        records.append(emf_record(37, -2147483640))  # Select the stock NULL_PEN.
        records.append(emf_record(75, 0, 5))  # Clear ordinary clipping; a crop must survive this.
    for index, (color, box) in enumerate([
        (0xFFFF00, (-80, -60, 160, 120)), (0x0000FF, (0, 0, 40, 30)),
        (0xFF0000, (40, 0, 80, 30)), (0x008000, (0, 30, 40, 60)),
        (0x00FFFF, (40, 30, 80, 60)),
    ]):
        if extension == "wmf":
            records.extend([wmf_record(0x02fc, 0, color & 0xFFFF, color >> 16, 0), wmf_record(0x012d, index + 1)])
            left, top, right, bottom = [value * 20 + (400 if axis % 2 == 0 else -200)
                                        for axis, value in enumerate(box)]
            drawing = wmf_record(0x041b, bottom, right, top, left)
        else:
            records.extend([emf_record(39, index + 1, 0, color, 0), emf_record(37, index + 1)])
            drawing = emf_record(43, *(value + (100 if axis % 2 == 0 else 200)
                                       for axis, value in enumerate(box)))
        drawings.append(drawing)
        records.append(drawing)
    if extension == "wmf":
        records.append(wmf_record(0))
        placeable = struct.pack('<IHhhhhHI', 0x9AC6CDD7, 0, 0, 0, 1600, 1200, 1440, 0)
        checksum = 0
        for word in struct.unpack('<10H', placeable):
            checksum ^= word
        header = struct.pack('<HHHIHIH', 1, 9, 0x300, 9 + sum(len(record) // 2 for record in records),
                             6, max(len(record) // 2 for record in records), 0)
        return placeable + struct.pack('<H', checksum) + header + b''.join(records), drawings
    records.append(emf_record(14, 0, 0, 20))
    header = bytearray(108)
    struct.pack_into('<II8i', header, 0, 1, 108, 100, 200, 180, 260, 1000, 2000, 1800, 2600)
    struct.pack_into('<4I', header, 40, 0x464D4520, 0x10000, 108 + sum(map(len, records)), 1 + len(records))
    struct.pack_into('<H', header, 56, 6)
    struct.pack_into('<4i', header, 72, 1000, 1000, 100, 100)
    struct.pack_into('<2i', header, 100, 100000, 100000)
    return bytes(header) + b''.join(records), drawings


def wmf_record(kind: int, *arguments: int) -> bytes:
    """Encode an independent WMF fixture record, accepting both signed coordinates and unsigned words."""
    return struct.pack('<IH', 3 + len(arguments), kind) + b''.join(
        struct.pack('<H', value & 0xFFFF) for value in arguments)


def emf_record(kind: int, *arguments: int) -> bytes:
    """Encode an independent EMF fixture record with 32-bit arguments."""
    return struct.pack('<II', kind, 8 + 4 * len(arguments)) + b''.join(
        struct.pack('<I', value & 0xFFFFFFFF) for value in arguments)


def replace_with_vector(source: Path, original: bytes, extension: str, *, svg_fallback: bool) -> None:
    """Embed shared vectors, optionally using Word's SVG extension with a deliberately different PNG fallback."""
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    png_name = next(name for name in parts if name.startswith('word/media/'))
    fallback = parts.pop(png_name)
    vector_name = png_name.rsplit('.', 1)[0] + '.' + extension
    parts[vector_name] = original
    for name, data in list(parts.items()):
        if name.endswith('.rels'):
            parts[name] = data.replace(png_name.rsplit('/', 1)[1].encode(), vector_name.rsplit('/', 1)[1].encode())
    if svg_fallback:
        parts['word/media/fallback.png'] = fallback
        relationships = etree.fromstring(parts['word/_rels/document.xml.rels'])
        etree.SubElement(relationships, '{' + relationships.nsmap[None] + '}Relationship', {
            'Id': 'rIdFallback', 'Target': 'media/fallback.png',
            'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image',
        })
        parts['word/_rels/document.xml.rels'] = etree.tostring(relationships)
        body = etree.fromstring(parts['word/document.xml'])
        blip = body.find('.//' + qn('a:blip'))
        vector_id = blip.get(qn('r:embed'))
        blip.set(qn('r:embed'), 'rIdFallback')
        extension_list = etree.SubElement(blip, qn('a:extLst'))
        svg_extension = etree.SubElement(extension_list, qn('a:ext'), {'uri': '{96DAC541-7B7A-43D3-8B79-37D633B846F1}'})
        etree.SubElement(svg_extension, '{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',
                         {qn('r:embed'): vector_id})
        parts['word/document.xml'] = etree.tostring(body)
    types = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(types, '{' + types.nsmap[None] + '}Default', {
        'Extension': extension, 'ContentType': 'image/svg+xml' if extension == 'svg' else f'image/x-{extension}',
    })
    parts['[Content_Types].xml'] = etree.tostring(types)
    with ZipFile(source, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def render_vector(data: bytes, extension: str, directory: Path) -> Image.Image:
    """Rasterize only in verification, using actual SVG/GDI consumers to check the exported vector canvas."""
    if extension == 'svg':
        executable = 'papper-svg.exe' if os.name == 'nt' else 'papper-svg'
        renderer = Path(os.environ.get('PAPPER_SVG_RENDERER', ROOT / 'target/debug' / executable))
        assert renderer.is_file(), 'Build papper-svg before running vector crop integration tests'
        result = subprocess.run([str(renderer), 'render', '--source', str(directory / 'picture.svg'), '--dpi', '96'],
                                input=data, capture_output=True, timeout=30)
        assert result.returncode == 0, result.stderr.decode(errors='replace')
        image = Image.open(BytesIO(result.stdout)).convert('RGBA')
    elif extension == 'emf':
        # Play the original EMF records directly through GDI. GDI+ translates GDI
        # rectangle clipping with different lower/right-edge pixel conventions.
        image = Image.open(BytesIO(data))
        image.load(dpi=254)
        image = image.convert('RGBA')
    else:
        image = render_metafile(data, extension, directory)
    background = Image.new('RGBA', image.size, 'white')
    background.alpha_composite(image)
    return background


def render_metafile(data: bytes, extension: str, directory: Path) -> Image.Image:
    """Use GDI+ playback, since Pillow's WMF-to-EMF conversion can refit out-of-frame drawing bounds."""
    if extension == 'wmf':
        left, top, right, bottom = struct.unpack_from('<4h', data, 6)
        units = struct.unpack_from('<H', data, 14)[0]
        width, height = round((right - left) * 72 / units), round((bottom - top) * 72 / units)
    else:
        left, top, right, bottom = struct.unpack_from('<4i', data, 24)
        width, height = round((right - left) / 10), round((bottom - top) / 10)
    source = directory / ('verification.' + extension)
    output = directory / 'verification.png'
    source.write_bytes(data)
    script = directory / 'render-vector.ps1'
    script.write_text('''# Render a test-only metafile with Windows GDI+ to verify its visible vector canvas.
# Parameters select the source, PNG verification artifact and physical-size-derived pixel dimensions.
param([string]$Source, [string]$Output, [int]$Width, [int]$Height)
$ErrorActionPreference = 'Stop'
Add-Type -AssemblyName System.Drawing
$image = [System.Drawing.Image]::FromFile($Source)
$bitmap = [System.Drawing.Bitmap]::new($Width, $Height)
$graphics = [System.Drawing.Graphics]::FromImage($bitmap)
try {
    $graphics.Clear([System.Drawing.Color]::White)
    $graphics.DrawImage($image, 0, 0, $Width, $Height)
    $bitmap.Save($Output, [System.Drawing.Imaging.ImageFormat]::Png)
} finally {
    $graphics.Dispose()
    $bitmap.Dispose()
    $image.Dispose()
}
''', encoding='utf-8')
    result = subprocess.run(['powershell.exe', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(script),
                             '-Source', str(source), '-Output', str(output), '-Width', str(width), '-Height', str(height)],
                            capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    with Image.open(output) as image:
        return image.convert('RGBA')


@pytest.mark.parametrize(('extension', 'viewport'), [
    ('svg', 'normal'), ('svg', 'letterbox'), ('svg', 'percent'), ('svg', 'physical'), ('emf', ''), ('wmf', ''),
])
def test_convert_crops_vectors_without_rasterizing(
    tmp_path: Path, rust_executable: Path, extension: str, viewport: str,
) -> None:
    """Keep native vector drawings and exact visible crops, including SVG fallback priority and negative padding."""
    if extension != 'svg' and not hasattr(Image.core, 'drawwmf'):
        pytest.skip('Independent metafile playback requires Windows GDI')
    source, _ = make_source(tmp_path, 'PNG')
    original, drawings = vector_picture(extension, viewport)
    replace_with_vector(source, original, extension, svg_fallback=extension == 'svg')
    untouched = source.read_bytes()
    output = tmp_path / 'converted'
    result = run_cli(rust_executable, tmp_path, 'convert', str(source), '-o', str(output), '--image-crop', 'apply', '--offline')
    assert 'Could not recover image crop' not in result.stderr
    assert source.read_bytes() == untouched
    paths = list((output / 'media').iterdir())
    assert {path.suffix for path in paths} == {'.' + extension}, 'Apply must preserve the native vector format'
    media = [path.read_bytes() for path in paths]
    assert original in media
    decoded = render_vector(original, extension, tmp_path)
    assert decoded.size == (80, 60)
    expected = [decoded.crop(box) for box in [(20, 12, 70, 54), (40, 0, 80, 60), (0, 0, 40, 60), (0, 0, 80, 45)]]
    padded = Image.new('RGBA', (80, 60), 'white')
    padded.paste(decoded.crop((0, 0, 60, 60)), (20, 0))
    expected.append(padded)
    actual = []
    for data in media:
        if data != original:
            assert all(drawing in data for drawing in drawings), 'Original paths/drawing records must remain vector content'
            actual.append(render_vector(data, extension, tmp_path))
    assert len(actual) == len(expected)
    for expected_image in expected:
        match = next((image for image in actual if image.size == expected_image.size
                      and ImageChops.difference(image, expected_image).getbbox(alpha_only=False) is None), None)
        assert match is not None, f'Missing visible vector crop {expected_image.size}'
        actual.remove(match)


def test_unsupported_wmf_mapping_keeps_original(tmp_path: Path, rust_executable: Path) -> None:
    """Warn instead of exporting a distorted image when a WMF explicitly changes its viewport."""
    source, _ = make_source(tmp_path, 'PNG')
    original, _ = vector_picture('wmf')
    original = bytearray(original)
    extra = wmf_record(0x020e, 1200, 1600)
    original[40:40] = extra
    struct.pack_into('<I', original, 28, (len(original) - 22) // 2)
    replace_with_vector(source, bytes(original), 'wmf', svg_fallback=False)
    output = tmp_path / 'converted'
    result = run_cli(rust_executable, tmp_path, 'convert', str(source), '-o', str(output), '--image-crop', 'apply', '--offline')
    assert 'WMF crop does not support viewport scaling' in result.stderr
    assert [path.read_bytes() for path in (output / 'media').iterdir()] == [original]
