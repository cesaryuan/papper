"""Guard imported data, source-local resources, and reply overrides through real output.

These cases cover cross-document and multi-directory behavior absent from ordinary
single-project snapshots. Native commands write isolated files; assertions inspect
their media bytes, visible text, references, and portable resource destinations.
"""

from __future__ import annotations

from io import BytesIO
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
from urllib.parse import unquote
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image
import pytest

from native_support import ROOT
from test_rust_docx_contract import rust_postprocessor


def run_cli(executable: Path, project: Path, *arguments: str) -> subprocess.CompletedProcess[str]:
    """Run current native code with isolated state and actionable failure diagnostics."""
    result = subprocess.run(
        [str(executable), *arguments], cwd=project,
        env={**os.environ, "PAPPER_HOME": str(project / "state"),
             "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_DISABLE_UPDATE_CHECK": "1"},
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    return result


def imported_image(markdown: Path) -> Path:
    """Resolve the first image actually referenced by an imported Markdown file."""
    match = re.search(r"\]\((media/[^)]+)\)", markdown.read_text(encoding="utf-8"))
    assert match is not None
    return markdown.parent / unquote(match[1])


def test_imports_preserve_other_documents_and_edited_media(tmp_path: Path, rust_executable: Path) -> None:
    """Same-named Word pictures must remain correct across imports, edits, and retries."""
    images = {}
    for name, color in [("first", "red"), ("second", "blue")]:
        picture = tmp_path / f"{name}.png"
        Image.new("RGB", (16, 16), color).save(picture)
        images[name] = picture.read_bytes()
        document = Document()
        document.add_picture(str(picture), width=Inches(1))
        document.save(tmp_path / f"{name}.docx")
        run_cli(rust_executable, tmp_path, "convert", f"{name}.docx", "-o", "converted")

    first = imported_image(tmp_path / "converted/first.md")
    second = imported_image(tmp_path / "converted/second.md")
    assert first != second
    assert first.read_bytes() == images["first"]
    assert second.read_bytes() == images["second"]
    run_cli(rust_executable, tmp_path, "convert", "second.docx", "-o", "converted")
    assert imported_image(tmp_path / "converted/second.md") == second

    # Edited hash-named resources need another name, rather than being overwritten on retry.
    second.write_bytes(images["first"])
    run_cli(rust_executable, tmp_path, "convert", "second.docx", "-o", "converted")
    replacement = imported_image(tmp_path / "converted/second.md")
    assert replacement != second
    assert replacement.read_bytes() == images["second"]
    assert second.read_bytes() == first.read_bytes() == images["first"]


@pytest.mark.parametrize("explicit", [False, True])
def test_svg_conversion_obeys_source_and_explicit_resource_roots(
    tmp_path: Path, rust_executable: Path, explicit: bool,
) -> None:
    """Final DOCX pixels must follow the same effective resource order as ordinary images."""
    chapter = tmp_path / "chapter"
    custom = tmp_path / "custom"
    chapter.mkdir()
    custom.mkdir()
    for directory, color in [(tmp_path, "red"), (chapter, "blue"), (custom, "lime")]:
        (directory / "figure.svg").write_text(
            f'<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20">'
            f'<rect width="20" height="20" fill="{color}"/></svg>', encoding="utf-8",
        )
    (tmp_path / "style.yml").write_text("mathtype: false\ndocxConvertSvgToPng: true\n", encoding="utf-8")
    (chapter / "paper.md").write_text("![Figure](figure.svg){width=1in}\n", encoding="utf-8")
    run_cli(rust_executable, tmp_path, "build", "docx", "chapter/paper.md", "-o", "paper.docx",
            *(["--resource-path", "custom"] if explicit else []))
    with ZipFile(tmp_path / "paper.docx") as archive:
        pixels = [Image.open(BytesIO(archive.read(name))).convert("RGB")
                  for name in archive.namelist() if name.startswith("word/media/") and name.endswith(".png")]
    assert pixels
    expected = (0, 255, 0) if explicit else (0, 0, 255)
    assert all(image.getpixel((image.width // 2, image.height // 2)) == expected for image in pixels)


@pytest.mark.parametrize("label", ["_RefConvertTwoCell", "_RefConvertTextLabel"])
def test_imported_bookmarked_equation_resolves_after_rebuild(
    tmp_path: Path, rust_executable: Path, label: str,
) -> None:
    """Flattened Word equation tables must survive a Markdown round trip with working references."""
    source = ROOT / "tests/snapshot_cases_convert/comprehensive/comprehensive.docx"
    run_cli(rust_executable, tmp_path, "convert", str(source), "-o", "converted")
    markdown = (tmp_path / "converted/comprehensive.md").read_text(encoding="utf-8")
    equation = next(line for line in markdown.splitlines() if f"{{#eq:{label}}}" in line)
    (tmp_path / "paper.md").write_text(equation + f"\n\nSee [@eq:{label}].\n", encoding="utf-8")
    result = run_cli(rust_executable, tmp_path, "build", "json", "paper.md", "-o", "paper.json")
    assert "Undefined cross-reference" not in result.stderr
    ast = json.loads((tmp_path / "paper.json").read_bytes())
    links = [inline["c"][2][0] for block in ast["blocks"] if block["t"] == "Para"
             for inline in block["c"] if inline["t"] == "Link"]
    assert f"#eq:{label}" in links


def test_reply_metadata_and_settings_survive_local_style_overlay(tmp_path: Path, rust_executable: Path) -> None:
    """Reply overrides inherit across styles, beat base values, and yield to reply YAML."""
    responses = tmp_path / "responses"
    responses.mkdir()
    (tmp_path / "style.yml").write_text(
        "mathtype: false\npandocMetadata:\n  title: Base title\n"
        "reply:\n  tableAutofit: content\n  pandocMetadata:\n    title: Reply title\n", encoding="utf-8",
    )
    (responses / "style.yml").write_text(
        "tableAutofit: fixed\nreply:\n  pandocMetadata:\n    subtitle: Local reply subtitle\n", encoding="utf-8",
    )
    reply = responses / "response.md"
    reply.write_text("Reviewer response.\n\n| Item | Value |\n|---|---|\n| A | B |\n", encoding="utf-8")
    run_cli(rust_executable, tmp_path, "build-reply", "responses/response.md", "-o", "response.docx")
    paragraphs = [paragraph.text for paragraph in Document(tmp_path / "response.docx").paragraphs]
    assert "Reply title" in paragraphs and "Local reply subtitle" in paragraphs
    assert "Base title" not in paragraphs
    with ZipFile(tmp_path / "response.docx") as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    table = xml.xpath("//w:tbl[.//w:t='Item']", namespaces=namespace)[0]
    assert table.find("w:tblPr/w:tblW", namespace).get(f"{{{namespace['w']}}}type") == "auto"
    reply.write_text("---\ntitle: Header title\n---\n\nReviewer response.\n", encoding="utf-8")
    run_cli(rust_executable, tmp_path, "build-reply", "responses/response.md", "-o", "response.docx")
    paragraphs = [paragraph.text for paragraph in Document(tmp_path / "response.docx").paragraphs]
    assert "Header title" in paragraphs and "Reply title" not in paragraphs


@pytest.mark.parametrize("external", [False, True])
def test_latex_images_are_portable_beside_custom_output(
    tmp_path: Path, rust_executable: Path, external: bool,
) -> None:
    """Image references must resolve from the exported .tex directory, even for absolute inputs."""
    chapter = tmp_path / "chapter"
    chapter.mkdir()
    picture = chapter / "figure.png"
    shutil.copy2(ROOT / "template/examples/images/single-figure-example.png", picture)
    path = picture.as_posix() if external else "figure.png"
    (chapter / "paper.md").write_text(f"![Figure]({path}){{width=1in}}\n", encoding="utf-8")
    run_cli(rust_executable, tmp_path, "build", "latex", "chapter/paper.md", "-o", "submission/paper.tex")
    exported = tmp_path / "submission/paper.tex"
    paths = re.findall(r"\\includegraphics\[[^\]]*\]\{([^}]+)\}", exported.read_text(encoding="utf-8"))
    assert paths
    assert all(not Path(path).is_absolute() for path in paths)
    assert all((exported.parent / path).read_bytes() == picture.read_bytes() for path in paths)


def test_docx_rewrite_preserves_large_binary_and_unknown_parts(
    tmp_path: Path, rust_postprocessor: Path,
) -> None:
    """Streaming publication must preserve opaque embedded data and remain readable as Word."""
    source = tmp_path / "original.docx"
    document = Document()
    document.add_paragraph("Preserve my embedded data.")
    document.save(source)
    opaque = os.urandom(4 * 1024 * 1024)
    with ZipFile(source, "a") as archive:
        archive.writestr("word/embeddings/custom-object.bin", opaque)
        archive.writestr("customXml/review.txt", b"unknown package content")
    output = tmp_path / "rewritten.docx"
    settings = tmp_path / "settings.json"
    settings.write_text(json.dumps({
        "pmt_settings": {"values": {}, "provided": [], "pandoc_metadata": {}, "reply": None},
        "pandoc_metadata": {}, "has_yaml_header": False,
    }), encoding="utf-8")
    result = subprocess.run([str(rust_postprocessor), str(source), str(output), str(settings)],
                            capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    with ZipFile(output) as archive:
        assert archive.read("word/embeddings/custom-object.bin") == opaque
        assert archive.read("customXml/review.txt") == b"unknown package content"
        assert archive.testzip() is None
    assert "Preserve my embedded data." in [paragraph.text for paragraph in Document(output).paragraphs]
