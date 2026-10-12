"""Verify custom header/footer output, dynamic page fields, and configuration precedence.

Crafted references cover first/even footers and authored content that the fixed
manuscript snapshots cannot exercise; output assertions protect data preservation.
"""

import os
import subprocess
import posixpath
from pathlib import Path
from zipfile import ZipFile

import pytest
import yaml
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from lxml import etree

from native_support import ROOT


def process_footer(source: Path, output: Path, executable: Path, config: dict | None, *, header: dict | None = None) -> None:
    """Build configured Markdown with an authored reference so Pandoc renders header/footer content."""
    markdown = output.with_suffix(".md")
    markdown.write_text("---\n" + yaml.safe_dump({"papperSettings": {"docxFooter": config, "docxHeader": header}}, allow_unicode=True)
                        + "---\nBody contents\n", encoding="utf-8")
    result = subprocess.run([
        str(executable), "build", "docx", str(markdown), "--reference-doc", str(source),
        "--no-mathtype", "-o", str(output),
    ], cwd=output.parent, env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT)},
        capture_output=True, text=True, encoding="utf-8", timeout=45)
    assert result.returncode == 0, result.stdout + result.stderr


@pytest.mark.parametrize("alignment", ["left", "center", "right"])
def test_custom_footer_replaces_variants_and_remains_dynamic(
    tmp_path: Path, rust_executable: Path, alignment: str,
) -> None:
    """Replace first/even page stories, create the missing primary, and avoid duplicate fields on reuse."""
    document = Document()
    document.styles.add_style("page number", WD_STYLE_TYPE.CHARACTER)
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.different_first_page_header_footer = True
    # Access only first/even parts: ordinary pages initially have no primary header/footer.
    section.first_page_footer.paragraphs[0].text = "Old first page"
    section.even_page_footer.paragraphs[0].text = "Old even page"
    section.first_page_header.paragraphs[0].text = "Old first header"
    section.even_page_header.paragraphs[0].text = "Old even header"
    document.add_section()
    source = tmp_path / "source.docx"
    document.save(source)
    config = {"content": '审稿 & \\<draft\\> - []{t="pageNum"} - / []{t=\'pageNum\'}',
              "alignment": alignment}
    for index in range(2):
        output = tmp_path / f"output-{index}.docx"
        process_footer(source, output, rust_executable, config, header=config)
        document = Document(output)
        for section in document.sections:
            for footer in (section.footer, section.first_page_footer, section.even_page_footer,
                           section.header, section.first_page_header, section.even_page_header):
                assert len(footer.paragraphs) == 1
                paragraph = footer.paragraphs[0]
                assert paragraph.text == "审稿 & <draft> - 1 - / 1"
                assert paragraph._p.find(qn("w:pPr")).find(qn("w:jc")).get(qn("w:val")) == alignment
                assert [field.text for field in paragraph._p.iter(qn("w:instrText"))] == [" PAGE ", " PAGE "]
                assert [field.get(qn("w:fldCharType")) for field in paragraph._p.iter(qn("w:fldChar"))] == [
                    "begin", "separate", "end", "begin", "separate", "end",
                ]
            assert section.different_first_page_header_footer
        assert document.settings.odd_and_even_pages_header_footer
        source = output


@pytest.mark.parametrize("content", [None, "", "Confidential & internal\\\nSecond line Right"])
def test_footer_preserve_clear_and_plain_text_without_page_style(
    tmp_path: Path, rust_executable: Path, content: str | None,
) -> None:
    """Preserve authored content with null, clear explicitly, and use Markdown hard line breaks."""
    document = Document()
    document.sections[0].footer.paragraphs[0].text = "Authored footer"
    document.sections[0].footer.add_paragraph("Second paragraph")
    document.sections[0].header.paragraphs[0].text = "Authored header"
    document.sections[0].header.add_paragraph("Second header paragraph")
    source = tmp_path / "source.docx"
    document.save(source)
    original = etree.tostring(document.sections[0].footer._element, method="c14n")
    original_header = etree.tostring(document.sections[0].header._element, method="c14n")
    output = tmp_path / "output.docx"
    process_footer(source, output, rust_executable, None if content is None else {"content": content},
                   header=None if content is None else {"content": content})
    header = Document(output).sections[0].header
    footer = Document(output).sections[0].footer
    if content is None:
        assert etree.tostring(header._element, method="c14n") == original_header
        assert etree.tostring(footer._element, method="c14n") == original
    else:
        assert len(header.paragraphs) == 1
        assert header.paragraphs[0].text == content.replace("\\\n", "\n")
        assert not list(header._element.iter(qn("w:instrText")))
        assert header.paragraphs[0].style.name == "Header"
        assert len(footer.paragraphs) == 1
        assert footer.paragraphs[0].text == content.replace("\\\n", "\n")
        assert not list(footer._element.iter(qn("w:instrText")))


@pytest.mark.parametrize(("project", "inline", "expected", "alignment", "header_inline", "header_expected", "header_alignment"), [
    ('docxFooter: {content: \'- []{t="pageNum"} -\', alignment: left}',
     'docxFooter: {alignment: right}', "- 1 -", "right", "docxHeader: {alignment: right}", "Header i", "right"),
    ('docxShowPageNumbers: false',
     'docxFooter: {content: \'- []{t="pageNum"} -\'}', "- 1 -", "center", "docxHeader: {content: ''}", "", "center"),
    ('docxFooter: {content: \'- []{t="pageNum"} -\'}',
     'docxShowPageNumbers: false', "Authored footer", None, "docxHeader: null", "Authored header", None),
    ('docxFooter: {content: \'- []{t="pageNum"} -\'}',
     'docxFooter: null', "Authored footer", None, "docxHeader: {content: Plain}", "Plain", "center"),
    ('docxShowPageNumbers: true\ndocxFooter: null',
     '', "Authored footer", None, "", "Header i", "center"),
])
def test_footer_build_obeys_source_priority(
    tmp_path: Path, rust_executable: Path, project: str, inline: str,
    expected: str, alignment: str | None, header_inline: str, header_expected: str,
    header_alignment: str | None,
) -> None:
    """Verify actual build output when new, legacy, and null values cross source layers."""
    reference = tmp_path / "reference.docx"
    document = Document(ROOT / "pandoc/manuscript-template/reference-doc.docx")
    document.sections[0].footer.paragraphs[0].text = "Authored footer"
    document.sections[0].header.paragraphs[0].text = "Authored header"
    document.save(reference)
    project += '\ndocxHeader: {content: \'Header []{t="pageNum" format="lowerRoman"}\', alignment: center}'
    inline = '\n'.join(value for value in (inline, header_inline) if value)
    (tmp_path / "style.yml").write_text(project + "\n", encoding="utf-8")
    source = tmp_path / "paper.md"
    settings = "\n".join("  " + line for line in inline.splitlines())
    source.write_text(f"---\npapperSettings:\n{settings}\n---\nBody\n" if inline else "Body\n", encoding="utf-8")
    output = tmp_path / "output.docx"
    result = subprocess.run([
        str(rust_executable), "build", "docx", str(source), "--reference-doc", str(reference),
        "--no-mathtype", "-o", str(output),
    ], cwd=tmp_path, env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT)},
        capture_output=True, text=True, encoding="utf-8", timeout=45)
    assert result.returncode == 0, result.stdout + result.stderr
    paragraph = Document(output).sections[0].footer.paragraphs[0]
    assert paragraph.text == expected
    assert len(list(paragraph._p.iter(qn("w:instrText")))) == (1 if expected == "- 1 -" else 0)
    if alignment is not None:
        assert paragraph._p.find(qn("w:pPr")).find(qn("w:jc")).get(qn("w:val")) == alignment

    header = Document(output).sections[0].header.paragraphs[0]
    assert header.text == header_expected
    assert len(list(header._p.iter(qn("w:instrText")))) == (1 if header_expected == "Header i" else 0)
    if header_alignment is not None:
        assert header._p.pPr.find(qn("w:jc")).get(qn("w:val")) == header_alignment


@pytest.mark.parametrize(("format_name", "instruction", "value"), [
    (None, " PAGE ", "1"),
    ("decimal", " PAGE \\* Arabic ", "1"),
    ("upperRoman", " PAGE \\* ROMAN ", "I"),
    ("lowerRoman", " PAGE \\* roman ", "i"),
    ("upperLetter", " PAGE \\* ALPHABETIC ", "A"),
    ("lowerLetter", " PAGE \\* alphabetic ", "a"),
    ("decimalEnclosedDash", " PAGE \\* ArabicDash ", "- 1 -"),
    ("ordinal", " PAGE \\* Ordinal ", "1st"),
    ("cardinalText", " PAGE \\* CardText ", "one"),
    ("ordinalText", " PAGE \\* OrdText ", "first"),
    ("hex", " PAGE \\* Hex ", "1"),
    ("ROMAN", " PAGE \\* ROMAN ", "I"),
])
def test_page_format_keeps_markdown_emphasis_and_dynamic_fields(
    tmp_path: Path, rust_executable: Path, format_name: str | None, instruction: str, value: str,
) -> None:
    """Protect actual page-field instructions and cached results, including enclosing bold text."""
    reference = tmp_path / "reference.docx"
    Document(ROOT / "pandoc/manuscript-template/reference-doc.docx").save(reference)
    attribute = '' if format_name is None else f' format="{format_name}"'
    output = tmp_path / "output.docx"
    config = {"content": f'**Page []{{t="pageNum"{attribute}}}**'}
    process_footer(reference, output, rust_executable, config, header=config)
    paragraph = Document(output).sections[0].footer.paragraphs[0]
    assert paragraph.text == "Page " + value
    assert [node.text for node in paragraph._p.iter(qn("w:instrText"))] == [instruction]
    assert all(run.bold for run in paragraph.runs)
    header = Document(output).sections[0].header.paragraphs[0]
    assert header.text == "Page " + value
    assert [node.text for node in header._p.iter(qn("w:instrText"))] == [instruction]
    assert all(run.bold for run in header.runs)


@pytest.mark.parametrize("postprocess", ["true", "false"])
def test_markdown_footer_preserves_media_links_tables_lists_and_math(
    tmp_path: Path, rust_executable: Path, postprocess: str, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Verify resource relationships, local anchors, and body isolation in each header/footer variant."""
    from PIL import Image

    monkeypatch.setenv("PMT_ENABLE_DOCX_POSTPROCESS", postprocess)
    image = tmp_path / "logo.png"
    Image.new("RGB", (24, 12), "#336699").save(image)
    reference = tmp_path / "reference.docx"
    document = Document(ROOT / "pandoc/manuscript-template/reference-doc.docx")
    document.sections[0].different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True
    document.sections[0].first_page_footer.paragraphs[0].text = "Old first"
    document.sections[0].even_page_footer.paragraphs[0].text = "Old even"
    document.sections[0].footer.paragraphs[0].text = "Old primary"
    document.sections[0].first_page_header.paragraphs[0].text = "Old first header"
    document.sections[0].even_page_header.paragraphs[0].text = "Old even header"
    document.save(reference)
    content = '''# Footer title {#footer-local}

**Bold** *Italic* `Code` $x^2$ [External](https://example.org/footer) [Local](#footer-local)

1. First item
2. Second item

| Logo | Number |
| :--- | ---: |
| ![Logo](logo.png){width=0.3in} | []{t="pageNum" format="lowerRoman"} |
'''
    source = tmp_path / "paper.md"
    source.write_text("---\n" + yaml.safe_dump({"papperSettings": {"docxFooter": {"content": content, "alignment": "right"},
                                                               "docxHeader": {"content": content.replace("Footer title", "Header title").replace("footer-local", "header-local"), "alignment": "left"}}})
                      + "---\nBody only. ![Body logo](logo.png){width=0.3in}\n", encoding="utf-8")
    output = tmp_path / "output.docx"
    result = subprocess.run([
        str(rust_executable), "build", "docx", str(source), "--reference-doc", str(reference),
        "--no-mathtype", "-o", str(output),
    ], cwd=tmp_path, env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT)},
        capture_output=True, text=True, encoding="utf-8", timeout=45)
    assert result.returncode == 0, result.stdout + result.stderr
    reopened = Document(output)
    assert reopened.sections[0].header.paragraphs[0].text == "Header title"
    with ZipFile(output) as archive:
        namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
                      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
        body = etree.fromstring(archive.read("word/document.xml"))
        assert "Header title" not in ''.join(body.itertext())
        assert "PMT_DOCX_HEADER" not in ''.join(body.itertext())
        assert "Footer title" not in ''.join(body.itertext())
        assert "PMT_DOCX_FOOTER" not in ''.join(body.itertext())
        assert len(body.xpath('//a:blip', namespaces=namespaces)) == 1
        drawing_ids = body.xpath('//wp:docPr/@id', namespaces=namespaces)
        bookmark_ids = body.xpath('//w:bookmarkStart/@w:id', namespaces=namespaces)
        bookmark_names = body.xpath('//w:bookmarkStart/@w:name', namespaces=namespaces)
        footers = [name for name in archive.namelist() if name.startswith('word/footer') and name.endswith('.xml')]
        headers = [name for name in archive.namelist() if name.startswith("word/header") and name.endswith(".xml")]
        assert len(footers) == 3
        assert len(headers) == 3
        for name in footers + headers:
            footer = etree.fromstring(archive.read(name))
            text = ''.join(footer.itertext())
            assert "Bold" in text and "Italic" in text and "Code" in text
            assert "First item" in text and "Second item" in text
            assert footer.xpath('//w:b', namespaces=namespaces)
            assert footer.xpath('//w:i', namespaces=namespaces)
            assert footer.xpath('//m:oMath', namespaces=namespaces)
            assert len(footer.xpath('//w:tbl', namespaces=namespaces)) == 1
            assert len(footer.xpath('//w:numPr', namespaces=namespaces)) == 2
            assert footer.xpath('//w:instrText/text()', namespaces=namespaces) == [' PAGE \\* roman ']
            # Paragraph placement applies outside tables; table cells keep their authored alignments.
            root = 'w:ftr' if name in footers else 'w:hdr'
            alignment = 'right' if name in footers else 'left'
            assert all(value == alignment for value in footer.xpath(f'/{root}/w:p/w:pPr/w:jc/@w:val', namespaces=namespaces))
            assert footer.xpath('//w:tc/w:p/w:pPr/w:jc/@w:val', namespaces=namespaces) == ['left', 'right', 'left', 'right']
            relations = etree.fromstring(archive.read('word/_rels/' + Path(name).name + '.rels'))
            by_id = {relation.get('Id'): relation for relation in relations}
            hyperlink = footer.xpath('//w:hyperlink[@r:id]/@r:id', namespaces=namespaces)
            assert len(hyperlink) == 1
            assert by_id[hyperlink[0]].get('Target') == 'https://example.org/footer'
            assert by_id[hyperlink[0]].get('TargetMode') == 'External'
            embedding = footer.xpath('//a:blip/@r:embed', namespaces=namespaces)
            assert len(embedding) == 1
            target = posixpath.normpath('word/' + by_id[embedding[0]].get('Target'))
            assert archive.read(target) == image.read_bytes()
            local_names = footer.xpath('//w:bookmarkStart/@w:name', namespaces=namespaces)
            assert footer.xpath('//w:hyperlink[@w:anchor]/@w:anchor', namespaces=namespaces) == local_names
            drawing_ids.extend(footer.xpath('//wp:docPr/@id', namespaces=namespaces))
            bookmark_ids.extend(footer.xpath('//w:bookmarkStart/@w:id', namespaces=namespaces))
            bookmark_names.extend(local_names)
        assert len(drawing_ids) == len(set(drawing_ids))
        assert len(bookmark_ids) == len(set(bookmark_ids))
        assert len(bookmark_names) == len(set(bookmark_names))


@pytest.mark.parametrize(("content", "message"), [
    ('[]{t="pageNum" format="invalid"}', 'Unsupported docxFooter pageNum format'),
    ('[]{t="pageNum" somekey="upperRoman"}', 'Unknown docxFooter pageNum attribute'),
    ('[text]{t="pageNum"}', 'pageNum spans must be empty'),
    ('Footer^[A note]', 'does not support footnotes'),
    ('See [@paper2026]', 'does not support bibliography citations'),
])
def test_invalid_footer_preserves_existing_output(
    tmp_path: Path, rust_executable: Path, content: str, message: str,
) -> None:
    """Reject unsupported page formats and Word-invalid content before replacing user output."""
    for setting in ("docxFooter", "docxHeader"):
        source = tmp_path / "paper.md"
        source.write_text("---\n" + yaml.safe_dump({"papperSettings": {setting: {"content": content}}})
                          + "---\nBody\n", encoding="utf-8")
        output = tmp_path / "output.docx"
        original = b"Previous successful document"
        output.write_bytes(original)
        result = subprocess.run([str(rust_executable), "build", "docx", str(source), "--no-mathtype", "-o", str(output)],
                                cwd=tmp_path, env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT)},
                                capture_output=True, text=True, encoding="utf-8", timeout=45)
        assert result.returncode != 0
        assert message.replace("docxFooter", setting) in result.stdout + result.stderr
        assert output.read_bytes() == original
