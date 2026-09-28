"""Verify user-visible Chinese DOCX builds and style customization."""

from pathlib import Path
import os
import re
import shutil
import subprocess
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn


PNG_PIXEL = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000b49444154789c636000020000050001a5f645400000000049454e44ae426082"
)


def run_docx_build(project: Path, cli_args: list[str]) -> tuple[str, Path]:
    """Run the public build command in an isolated manuscript project."""
    output = project / "result.docx"
    environment = os.environ.copy()
    environment["LANG"] = "en-US"
    result = subprocess.run(
        [sys.executable, "-m", "pandoc_manuscript.cli", "build", "docx", "paper.md", "--no-mathtype", "-o", str(output), *cli_args],
        cwd=project,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        env=environment,
        timeout=120,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    return result.stdout + result.stderr, output


@pytest.mark.parametrize(
    ("metadata_lang", "style_lang", "cli_args"),
    [
        ("zh-CN", None, []),
        (None, "zh-CN", []),
        ("en-US", None, ["--lang", "zh-cn"]),
        ("en-US", None, ["--lang", "zhcn"]),
    ],
)
def test_chinese_docx_build_numbers_figures_tables_and_formats_headings(
    tmp_path: Path,
    metadata_lang: str | None,
    style_lang: str | None,
    cli_args: list[str],
) -> None:
    """Use manuscript/style metadata or CLI override for the same Chinese DOCX behavior."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    (tmp_path / "figure.png").write_bytes(PNG_PIXEL)
    manuscript = tmp_path / "paper.md"
    original = (
        "---\n"
        + (f"lang: {metadata_lang}\n" if metadata_lang is not None else "")
        + "title: 中文标题\n"
        "references:\n"
        "  - id: csl-sample\n"
        "    type: article-journal\n"
        "    author:\n"
        "      - family: Wang\n"
        "        given: Wei\n"
        "    title: CSL Sample Article\n"
        "    container-title: Journal of Example Research\n"
        "    issued: {date-parts: [[2024]]}\n"
        "    volume: 5\n"
        "    page: 1-10\n"
        "    DOI: 10.1234/pmt-csl-sample\n"
        "---\n\n"
        "# 第一章\n\n# 第二章\n\n# 第三章\n\n## 第一节 {#sec:section}\n\n"
        "参见 [@sec:section] 和 [@csl-sample]。\n\n"
        "![测试图](figure.png){#fig:one}\n\n"
        "| 列一 | 列二 |\n| --- | --- |\n| 值一 | 值二 |\n\n: 测试表 {#tbl:one}\n"
    )
    manuscript.write_text(original, encoding="utf-8")
    if style_lang is not None:
        (tmp_path / "style.yml").write_text(
            f"docxShowLineNumbers: 连续\npandocMetadata:\n  lang: {style_lang}\n",
            encoding="utf-8",
        )

    output_log, output = run_docx_build(tmp_path, cli_args)

    assert "Could not load translations" not in output_log
    assert "has no translation defined" not in output_log
    assert manuscript.read_text(encoding="utf-8") == original
    assert not list(tmp_path.glob(".paper.pmt-no-lang-*.md"))
    doc = Document(output)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "图 3-1" in text
    assert "表 3-1" in text
    assert any(paragraph.text.startswith(("3.1\t", "3.1 ")) for paragraph in doc.paragraphs)
    assert re.search(r"节\s+3\.1", text)
    assert "CSL Sample Article" in text
    assert "10.1234/pmt-csl-sample" not in text
    assert "参考文献" in text
    assert all(section._sectPr.find(qn("w:lnNumType")) is None for section in doc.sections)
    assert doc.styles["Heading 1"].font.size.pt == 15
    assert doc.styles["Heading 2"].font.size.pt == 14
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        assert style.font.bold is False, name
        assert style.element.rPr.bCs.val is False, name
        fonts = style.element.rPr.rFonts
        assert fonts.get(qn("w:eastAsia")) == "黑体", name
        assert fonts.get(qn("w:ascii")) == "Times New Roman", name
        assert fonts.get(qn("w:eastAsiaTheme")) is None, name


def test_chinese_section_reference_uses_dots_without_links(tmp_path: Path) -> None:
    """Keep dotted section references when crossref emits plain text instead of links."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    (tmp_path / "paper.md").write_text(
        "---\nlang: zh-CN\n---\n\n# 第一章\n\n## 第一节 {#sec:section}\n\n参见 [@sec:section]。\n",
        encoding="utf-8",
    )
    (tmp_path / "style.yml").write_text(
        "pandocMetadata:\n  linkReferences: false\n",
        encoding="utf-8",
    )

    _, output = run_docx_build(tmp_path, [])
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert re.search(r"节\s+1\.1", text)


@pytest.mark.parametrize(
    ("font_family", "western", "chinese"),
    [
        ("黑体", "黑体", "黑体"),
        ('{western: "Times New Roman", chinese: "宋体"}', "Times New Roman", "宋体"),
    ],
)
def test_docx_style_font_and_bold_work_without_language_mode(
    tmp_path: Path,
    font_family: str,
    western: str,
    chinese: str,
) -> None:
    """Apply string or script-specific fonts in an ordinary DOCX build."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    (tmp_path / "paper.md").write_text("# Heading\n\nBody.\n", encoding="utf-8")
    (tmp_path / "style.yml").write_text(
        f"docxStyle:\n  标题 1: {{fontFamily: {font_family}, bold: false}}\n",
        encoding="utf-8",
    )

    _, output = run_docx_build(tmp_path, [])

    doc = Document(output)
    style = doc.styles["Heading 1"]
    assert style.font.bold is False
    assert style.element.rPr.rFonts.get(qn("w:ascii")) == western
    assert style.element.rPr.rFonts.get(qn("w:hAnsi")) == western
    assert style.element.rPr.rFonts.get(qn("w:eastAsia")) == chinese
    assert all(section._sectPr.find(qn("w:lnNumType")) is not None for section in doc.sections)


@pytest.mark.parametrize("csl_source", ["manuscript", "style"])
def test_explicit_csl_overrides_chinese_docx_default(tmp_path: Path, csl_source: str) -> None:
    """Keep manuscript and style CSL overrides when Chinese defaults apply."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    csl = Path(__file__).resolve().parents[1] / "pandoc" / "csl" / "elsevier-vancouver.csl"
    if csl_source == "style":
        (tmp_path / "style.yml").write_text(
            f"pandocMetadata:\n  csl: '{csl.as_posix()}'\n",
            encoding="utf-8",
        )
    manuscript_csl = f"csl: '{csl.as_posix()}'\n" if csl_source == "manuscript" else ""
    (tmp_path / "paper.md").write_text(
        "---\n" + manuscript_csl +
        "references:\n"
        "  - id: sample\n"
        "    type: article-journal\n"
        "    title: CSL Override Article\n"
        "    issued: {date-parts: [[2024]]}\n"
        "    DOI: 10.1234/pmt-csl-override\n"
        "---\n\nCite [@sample].\n",
        encoding="utf-8",
    )

    _, output = run_docx_build(tmp_path, ["--lang", "zhcn"])

    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "10.1234/pmt-csl-override" in text


def test_chinese_build_preserves_explicit_crossref_metadata(tmp_path: Path) -> None:
    """Render manuscript cross-reference choices above Chinese language defaults."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    (tmp_path / "figure.png").write_bytes(PNG_PIXEL)
    (tmp_path / "paper.md").write_text(
        "---\n"
        "lang: zh-CN\n"
        "figureTitle: 'Fig. '\n"
        "figPrefix: Fig.\n"
        "---\n\n"
        "参见 @fig:one。\n\n"
        "![测试图](figure.png){#fig:one}\n",
        encoding="utf-8",
    )

    _, output = run_docx_build(tmp_path, [])

    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "Fig. 1 测试图" in text


def test_chinese_docx_preserves_explicit_heading_style(tmp_path: Path) -> None:
    """Apply user DOCX style choices after Chinese heading defaults."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    (tmp_path / "paper.md").write_text("# 标题\n\n正文。\n", encoding="utf-8")
    (tmp_path / "style.yml").write_text(
        "pandocMetadata:\n  lang: zh-CN\n"
        "docxStyle:\n  标题 1: {fontSize: 四号, bold: true}\n",
        encoding="utf-8",
    )

    _, output = run_docx_build(tmp_path, [])

    heading = Document(output).styles["Heading 1"]
    assert heading.font.size.pt == 14
    assert heading.font.bold is True


def test_chinese_docx_can_explicitly_enable_line_numbers(tmp_path: Path) -> None:
    """Allow an explicit line-number setting to override the Chinese default."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    (tmp_path / "paper.md").write_text("# 标题\n\n正文。\n", encoding="utf-8")
    (tmp_path / "style.yml").write_text(
        "docxShowLineNumbers: true\npandocMetadata:\n  lang: zh-CN\n",
        encoding="utf-8",
    )

    _, output = run_docx_build(tmp_path, [])

    doc = Document(output)
    assert all(section._sectPr.find(qn("w:lnNumType")) is not None for section in doc.sections)
