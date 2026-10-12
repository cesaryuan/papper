"""Verify Lua author output and custom note marks without optional DOCX formatting."""

from __future__ import annotations

import shutil
import os
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree, html

from native_support import ROOT, papper_command

PANDOC = shutil.which("pandoc")
pytestmark = pytest.mark.skipif(PANDOC is None, reason="author rendering requires Pandoc")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def render(source: str, target: str, tmp_path: Path, *, standalone: bool = True) -> Path:
    """Render Lua authors with only required note finalization for DOCX output."""
    markdown = tmp_path / "paper.md"
    markdown.write_text(source, encoding="utf-8")
    output = tmp_path / f"paper.{target}"
    command = [PANDOC, str(markdown), "-t", target, "--wrap=none", "-L",
               str(ROOT / "pandoc/filters/shared/author_info.lua"), "-o", str(output)]
    if standalone:
        command.extend(["-s", "--template", str(ROOT / f"pandoc/templates/default.{'xml' if target == 'docx' else 'html'}")])
    environment = {**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "state"),
                   "PMT_ENABLE_DOCX_POSTPROCESS": "false", "PAPPER_DISABLE_UPDATE_CHECK": "1"}
    if target == "docx":
        command = [*papper_command(), "build", "docx", str(markdown), "--no-mathtype", "-o", str(output)]
    result = subprocess.run(command, cwd=tmp_path, env=environment,
                            capture_output=True, text=True, encoding="utf-8", timeout=45)
    assert result.returncode == 0, result.stderr
    return output


@pytest.mark.parametrize("target", ["docx", "html"])
def test_keyed_inline_authors_keep_labels_escape_text_and_precede_abstract(target: str, tmp_path: Path) -> None:
    """Catch sparse-label overwrites, alias loss, bad escaping, and misplaced institutions."""
    source = """---
title: Author paper
abstract: Abstract text
author:
  - name: Alice & Bob
    affiliation: [b, New <Lab>, Existing Lab]
    corresponding: Ask <Alice> & Bob
  - name: Carol
    affiliations: [New <Lab>, b]
affiliation:
  b: Existing Lab
  '10': Tenth Lab
  '2': Second Lab
---

Body text with a note.[^body]

[^body]: Original body note.
"""
    output = render(source, target, tmp_path)
    if target == "html":
        tree = html.fromstring(output.read_text(encoding="utf-8"))
        block = tree.xpath('//div[@class="papper-authors"]')[0]
        assert block.xpath('.//span[@class="author-name"]/text()') == ["Alice & Bob", "Carol"]
        assert block.xpath('.//div[@class="author"]/p/sup/text()') == ["b,a,b", "*", "a,b"]
        institutions = block.xpath('.//div[@class="affiliation"]/p')
        assert [p.text_content() for p in institutions] == ["2 Second Lab", "10 Tenth Lab", "a New <Lab>", "b Existing Lab"]
        assert "*Ask <Alice> & Bob" in block.text_content()
        assert tree.xpath('//header/*')[1] == block
        assert "Original body note." in tree.text_content()
    else:
        with ZipFile(output) as package:
            document = etree.fromstring(package.read("word/document.xml"))
            notes = etree.fromstring(package.read("word/footnotes.xml"))
        paragraphs = document.xpath('//w:body/w:p', namespaces=NS)
        text = ["".join(p.xpath('.//w:t/text()', namespaces=NS)) for p in paragraphs]
        assert text[:6] == ["Author paper", "Alice & Bobb,a,b*, Carola,b", "2 Second Lab", "10 Tenth Lab", "a New <Lab>", "b Existing Lab"]
        assert text.index("Abstract text") > 5
        refs = document.xpath('//w:footnoteReference/@w:id', namespaces=NS)
        assert len(refs) == len(set(refs)) == 2
        actual_notes = {n.get(f"{{{NS['w']}}}id"): "".join(n.xpath('.//w:t/text()', namespaces=NS)).strip()
                        for n in notes.xpath('//w:footnote[not(@w:type)]', namespaces=NS)}
        assert [actual_notes[ref] for ref in refs] == ["* Ask <Alice> & Bob", "Original body note."]
        references = document.xpath('//w:footnoteReference', namespaces=NS)
        assert references[0].get(f"{{{NS['w']}}}customMarkFollows") == "1"
        assert references[0].getparent().xpath('./w:t/text()', namespaces=NS) == ["*"]
        assert references[1].get(f"{{{NS['w']}}}customMarkFollows") is None
        assert notes.xpath('//w:footnote[@w:id=$id]//w:footnoteRef', id=refs[1], namespaces=NS)
        assert not notes.xpath('//w:footnote[@w:id=$id]//w:footnoteRef', id=refs[0], namespaces=NS)
        assert "PMT_CORRESPONDENCE_MARK" not in etree.tostring(notes, encoding="unicode")
        assert paragraphs[1].xpath('./w:pPr/w:pStyle/@w:val', namespaces=NS) == ["Author"]
        assert paragraphs[2].xpath('./w:pPr/w:pStyle/@w:val', namespaces=NS) == ["Affiliation"]


@pytest.mark.parametrize("target", ["docx", "html"])
def test_correspondence_without_affiliation_or_title_is_rendered(target: str, tmp_path: Path) -> None:
    """Catch the old missing correspondence marker when an author had no institution."""
    source = """---
authors:
  - name: Alice
    corresponding: true
    title: Professor
    email: alice@example.test
---

Body text.
"""
    output = render(source, target, tmp_path)
    sentence = "Correspondence to: Dr. Alice (Professor). E-mail: alice@example.test."
    if target == "html":
        tree = html.fromstring(output.read_text(encoding="utf-8"))
        assert sentence in tree.xpath('//div[@class="papper-authors"]')[0].text_content()
        fragment = render(source, target, tmp_path, standalone=False).read_text(encoding="utf-8")
        assert "Alice" not in fragment
        assert "Body text." in fragment
    else:
        with ZipFile(output) as package:
            document = etree.fromstring(package.read("word/document.xml"))
            notes = etree.fromstring(package.read("word/footnotes.xml"))
        assert document.xpath('//w:footnoteReference/@w:id', namespaces=NS)
        assert sentence in "".join(notes.xpath('//w:t/text()', namespaces=NS))


@pytest.mark.parametrize("postprocess", ["true", "false"])
def test_author_style_overrides_work_without_postprocessing(tmp_path: Path, postprocess: str) -> None:
    """Catch author output that still depends on postprocessing or hardcodes its typography."""
    from docx import Document

    source = tmp_path / "paper.md"
    source.write_text("---\ntitle: Paper\nauthors:\n  - name: Alice\n    affiliation: Example Lab\n    corresponding: true\n  - name: Bob\n    corresponding: Contact Bob\n---\n\nBody.[^body]\n\n[^body]: Body note.\n", encoding="utf-8")
    style = tmp_path / "style.yml"
    style.write_text("docxStyle:\n  Author:\n    fontSize: 17pt\n  Affiliation:\n    fontSize: 9pt\n", encoding="utf-8")
    output = tmp_path / "paper.docx"
    result = subprocess.run(
        [*papper_command(), "build", "docx", str(source), "--no-mathtype", "--style-file", str(style), "-o", str(output)],
        cwd=tmp_path, env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "state"),
                         "PMT_ENABLE_DOCX_POSTPROCESS": postprocess, "PAPPER_DISABLE_UPDATE_CHECK": "1"},
        capture_output=True, text=True, encoding="utf-8", timeout=45,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    author = next(p for p in document.paragraphs if p.text == "Alicea*, Bob*")
    affiliation = next(p for p in document.paragraphs if p.text == "a Example Lab")
    assert author.style.font.size.pt == 17
    assert affiliation.style.font.size.pt == 9
    assert all(run.font.size is None for run in author.runs + affiliation.runs)
    with ZipFile(output) as package:
        notes = etree.fromstring(package.read("word/footnotes.xml"))
        xml = etree.fromstring(package.read("word/document.xml"))
    refs = xml.xpath('//w:footnoteReference', namespaces=NS)
    assert [ref.get(f"{{{NS['w']}}}customMarkFollows") for ref in refs] == ["1", "1", None]
    assert len({ref.get(f"{{{NS['w']}}}id") for ref in refs}) == 3
    assert "* Correspondence to: Dr. Alice, Example Lab" in "".join(notes.xpath('//w:t/text()', namespaces=NS))
    assert "* Contact Bob" in "".join(notes.xpath('//w:t/text()', namespaces=NS))
