"""Protect native insertion output, section inheritance, and failed-build data preservation.

The comprehensive snapshot covers serialized output; these contracts exercise
optional postprocessing/native numbering and invalid source recovery separately.
"""

from __future__ import annotations

from datetime import date
import json
import os
import re
from pathlib import Path
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import html
import pytest

from native_support import ROOT


def test_exported_directive_choices_render_through_the_public_cli(tmp_path: Path, rust_executable: Path) -> None:
    """Prevent editor choices drifting from accepted directives and their real Word fields."""
    exported = subprocess.run(
        [str(rust_executable), "config", "describe", "--format", "json"],
        cwd=tmp_path, capture_output=True, text=True, encoding="utf-8", check=True, timeout=15,
    )
    element = json.loads(exported.stdout)["elementAttributes"]["emptySpan"]
    directives = []
    page_fields = 0
    for kind, variant in element["variants"].items():
        for spelling in [kind, *variant.get("aliases", [])]:
            required = ' instruction="PAGE"' if kind == "field" else ""
            directives.append(f'[]{{t="{spelling}"{required}}}')
            page_fields += kind == "pageNum"
        for key, field in variant["properties"].items():
            for value in [*field.get("values", []), *field.get("suggestions", [])]:
                attributes = {element["discriminator"]: kind, key: value}
                if kind == "field":
                    attributes.setdefault("instruction", "PAGE")
                # Alignment is meaningful only alongside an independent story.
                if kind == "sectionBreak" and key in {"headerAlignment", "footerAlignment"}:
                    attributes[key.removesuffix("Alignment")] = "Story"
                authored = " ".join(f'{name}={json.dumps(value, ensure_ascii=False)}' for name, value in attributes.items())
                directives.append(f"[]{{{authored}}}")
                page_fields += kind == "pageNum"
    source = tmp_path / "choices.md"
    source.write_text("---\npapperSettings:\n  mathtype: false\n---\n\n" + "\n\n".join(directives), encoding="utf-8")
    output = tmp_path / "choices.docx"
    result = build_insertions(rust_executable, source, output)
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    instructions = [node.text or "" for node in document.element.iter(qn("w:instrText"))]
    assert sum(" PAGE " in instruction for instruction in instructions) >= page_fields
    assert any(" TOC " in instruction for instruction in instructions)
    assert any(' DATE \\@ "yyyy年M月d日" ' in instruction for instruction in instructions)
    today = date.today()
    assert f"{today.year}年{today.month}月{today.day}日" in "".join(document.element.itertext())


def build_insertions(executable: Path, source: Path, output: Path, *, postprocess: bool = True, target: str = "docx", reference: Path | None = None) -> subprocess.CompletedProcess:
    """Build through the public CLI with request-local state and the selected finalization path."""
    return subprocess.run(
        [str(executable), "build", target, str(source), *(["--no-mathtype"] if target == "docx" else []),
         *(["--reference-doc", str(reference)] if reference is not None else []), "-o", str(output)],
        cwd=output.parent,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PMT_ENABLE_DOCX_POSTPROCESS": str(postprocess).lower()},
        capture_output=True, text=True, encoding="utf-8", timeout=60,
    )


@pytest.mark.parametrize("postprocess,native", [(True, False), (True, True), (False, False)])
def test_insertions_survive_finalization_and_preserve_section_stories(
    tmp_path: Path, rust_executable: Path, postprocess: bool, native: bool,
) -> None:
    """Verify effective stories and numbering, including an unlinked copy followed by a clear."""
    source = tmp_path / "input.md"
    fixture = ROOT / "tests/snapshot_cases/docx_directives/docx_directives.md"
    text = fixture.read_text(encoding="utf-8").replace(
        "../crossrefs/figure.svg", (ROOT / "tests/snapshot_cases/crossrefs/figure.svg").as_posix(),
    ).replace("  mathtype: false", f"  mathtype: false\n  docxNativeCrossref: {str(native).lower()}")
    # A hyperlink in an independent story exercises OPC relationship relocation,
    # which plain-text output snapshots cannot distinguish from a broken target.
    text = text.replace('header="**Methods header**"',
                        'header="**Methods header** [Site](https://example.org) []{t=\'currentDate\'}"')
    source.write_text(text, encoding="utf-8")
    output = tmp_path / "output.docx"
    result = build_insertions(rust_executable, source, output, postprocess=postprocess)
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    first, methods, copied, cleared = document.sections
    assert first.header.paragraphs[0].text == "Global header"
    assert methods.header.paragraphs[0].text == copied.header.paragraphs[0].text
    assert "Methods header" in methods.header.paragraphs[0].text
    assert date.today().isoformat() in methods.header.paragraphs[0].text
    assert not methods.header.is_linked_to_previous
    assert not copied.header.is_linked_to_previous
    assert methods.header.part.partname != copied.header.part.partname
    assert copied.footer.is_linked_to_previous
    assert methods.footer.paragraphs[0].text == copied.footer.paragraphs[0].text == "Section page 1"
    assert cleared.header.paragraphs[0].text == cleared.footer.paragraphs[0].text == ""
    assert methods._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) == "1"
    assert copied._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) is None
    assert [section._sectPr.find(qn("w:type")).get(qn("w:val")) for section in (methods, copied, cleared)] == [
        "nextPage", "continuous", "oddPage",
    ]
    for story in (methods.header, copied.header):
        assert any(relationship.target_ref == "https://example.org" for relationship in story.part.rels.values())
    instructions = [node.text.strip() for node in document.element.iter(qn("w:instrText"))]
    assert instructions[0] == r'TOC \o "1-2" \h \z'
    directories = [paragraph for paragraph in document.paragraphs
                   if any(' \\c' in node.text for node in paragraph._p.iter(qn("w:instrText")))]
    assert len(directories) == 12
    assert all(paragraph.style.name.lower() == "table of figures" for paragraph in directories)
    # Style-based lists must not add SEQ fields or change build-time numbering.
    assert sum(instruction.startswith("SEQ Figure") for instruction in instructions) == (6 if native else 0)
    assert sum(instruction.startswith("SEQ Table") for instruction in instructions) == (4 if native else 0)
    translated = [paragraph for paragraph in document.paragraphs if "Translated" in paragraph.text]
    assert len(translated) == 2
    assert all(not any("SEQ " in node.text for node in paragraph._p.iter(qn("w:instrText")))
               for paragraph in translated)
    # Panel captions stay in layout-table cells and are available only to expanded lists.
    panel_paragraphs = document.element.xpath(
        './/w:tc/w:p[w:pPr/w:pStyle[starts-with(@w:val,"PapperSubfigureCaption")]]',
    )
    assert len(panel_paragraphs) == 4
    assert document.settings.element.find(qn("w:updateFields")).get(qn("w:val")) == "true"
    assert any(value.get(qn("w:type")) == "page" for value in document.element.iter(qn("w:br")))
    with ZipFile(output) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                assert b"PMT_DOCX_" not in archive.read(name)
                assert b"PMT_FOOTER_PAGE:" not in archive.read(name)


@pytest.mark.parametrize("native", [False, True])
def test_caption_lists_follow_custom_labels(tmp_path: Path, rust_executable: Path, native: bool) -> None:
    """Prevent empty native lists when localized caption labels differ from defaults."""
    source = tmp_path / "labels.md"
    figure = (ROOT / "tests/snapshot_cases/crossrefs/figure.svg").as_posix()
    source.write_text(
        f'---\nfigureTitle: "\u00a0图\u3000"\ntableTitle: 表\n'
        f'papperSettings:\n  mathtype: false\n  docxNativeCrossref: {str(native).lower()}\n'
        '  docxStyle:\n    图表目录:\n      fontSize: 11pt\n---\n\n'
        '[]{t="listofFigures"}\n\n[]{t="listofTables"}\n\n'
        f'![Primary]({figure}){{#fig:one}}\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n: Values {{#tbl:one}}\n',
        encoding="utf-8",
    )
    output = tmp_path / "labels.docx"
    result = build_insertions(rust_executable, source, output)
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    instructions = [node.text.strip() for node in document.element.iter(qn("w:instrText"))]
    if native:
        assert r'TOC \h \z \c "图"' in instructions
        assert r'TOC \h \z \c "表"' in instructions
        assert sum(instruction.startswith("SEQ 图 ") for instruction in instructions) == 1
        assert sum(instruction.startswith("SEQ 表 ") for instruction in instructions) == 1
    else:
        assert r'TOC \t "Image Caption,1" \h \z \c' in instructions
        assert r'TOC \t "Table Caption,1" \h \z \c' in instructions
        assert not any(instruction.startswith("SEQ ") for instruction in instructions)
    assert document.styles["table of figures"].font.size.pt == 11


@pytest.mark.parametrize("postprocess,native", [(True, False), (True, True), (False, False)])
def test_caption_lists_select_sources_independently(
    tmp_path: Path, rust_executable: Path, postprocess: bool, native: bool,
) -> None:
    """Keep main and caption-en directories independent of the text's language and numbering mode."""
    fixture = ROOT / "tests/snapshot_cases/docx_directives/docx_directives.md"
    text = fixture.read_text(encoding="utf-8").replace(
        "../crossrefs/figure.svg", (ROOT / "tests/snapshot_cases/crossrefs/figure.svg").as_posix(),
    ).replace("  mathtype: false", f"  mathtype: false\n  docxNativeCrossref: {str(native).lower()}\n"
              '  docxStyle:\n    Image Caption:\n      paragraphSpacing:\n        after: 12pt\n'
              '    Table Caption:\n      paragraphSpacing:\n        before: 12pt\n        after: 2pt')
    source = tmp_path / "sources.md"
    source.write_text(text, encoding="utf-8")
    output = tmp_path / "sources.docx"
    result = build_insertions(rust_executable, source, output, postprocess=postprocess)
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    styles = {style.style_id: style.name for style in document.styles}
    captions = []
    for paragraph in document.element.iter(qn("w:p")):
        style = paragraph.find(f'{qn("w:pPr")}/{qn("w:pStyle")}')
        if style is not None:
            name = styles[style.get(qn("w:val"))]
            if "Caption" in name:
                captions.append((name, "".join(node.text or "" for node in paragraph.iter(qn("w:t"))),
                                 [node.text.strip() for node in paragraph.iter(qn("w:instrText"))]))
    instructions = [node.text.strip() for node in document.element.iter(qn("w:instrText"))]
    directories = [instruction for instruction in instructions if instruction.startswith("TOC ") and "\\c" in instruction]
    selected = []
    for instruction in directories:
        # Resolve Word's actual style or SEQ selectors against the published paragraphs.
        selector = re.search(r'\\t "([^"]+)"', instruction)
        if selector:
            names = selector[1].split(",")[::2]
            selected.append([text for name, text, _ in captions if name in names])
        else:
            label = re.search(r'\\c "([^"]+)"', instruction)
            assert label is not None
            selected.append([text for _, text, codes in captions
                             if any(code.startswith(f"SEQ {label[1]} ") for code in codes)])
    assert [len(entries) for entries in selected] == [6, 10, 4, 6, 6, 2, 10, 10, 2, 4, 4, 2]
    selected = selected[-9:]
    assert selected[0] == selected[1]
    assert selected[3] == selected[4]
    assert selected[6] == selected[7]
    assert any("English only figure" in text for text in selected[0])
    assert any("中文单语图" in text for text in selected[0])
    assert any("English only table" in text for text in selected[6])
    assert any("中文单语表" in text for text in selected[6])
    assert not any("translation" in text.lower() or "Translated" in text
                   for index in (0, 3, 6) for text in selected[index])
    assert all("translation" in text.lower() or "Translated" in text
               for index in (2, 5, 8) for text in selected[index])
    assert any("译注" in text for text in selected[2])
    assert any("English table translation" in text for text in selected[8])
    assert not any(panel in text for text in selected[0]
                   for panel in ("First panel", "Second panel", "English panel", "中文子图"))
    assert any("English panel" in text for text in selected[3])
    assert {name for name, text, _ in captions if "English only figure" in text} == {"Image Caption"}
    assert {name for name, text, _ in captions if "译注" in text} == {"Image Caption EN"}
    assert document.styles["Image Caption EN"].base_style.name == "Image Caption"
    assert document.styles["Table Caption EN"].base_style.name == "Table Caption"
    # Both kinds clear main space-after; tables also clear English space-before,
    # retaining the configured outer gaps and single-caption spacing.
    assert document.styles["Image Caption"].paragraph_format.space_after.pt == 12
    primary = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Image Caption"]
    assert len(primary) == 6
    for paragraph in primary:
        bilingual = "Primary figure" in paragraph.text or "中文双语图" in paragraph.text
        assert paragraph.paragraph_format.space_after == (0 if bilingual else None)
    assert document.styles["Table Caption"].paragraph_format.space_before.pt == 12
    assert document.styles["Table Caption"].paragraph_format.space_after.pt == 2
    primary_tables = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Table Caption"]
    assert len(primary_tables) == 4
    for paragraph in primary_tables:
        bilingual = "Measurements" in paragraph.text or "中文双语表" in paragraph.text
        assert paragraph.paragraph_format.space_before is None
        assert paragraph.paragraph_format.space_after == (0 if bilingual else None)
    english_tables = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Table Caption EN"]
    assert len(english_tables) == 2
    assert all(paragraph.paragraph_format.space_before == 0
               and paragraph.paragraph_format.space_after is None for paragraph in english_tables)
    assert sum(instruction.startswith("SEQ ") for instruction in instructions) == (10 if native else 0)
    assert "PMT_DOCX_CAPTION:" not in document.element.xml


@pytest.mark.parametrize("postprocess,native", [(True, False), (True, True), (False, False)])
def test_custom_fields_preserve_instructions_results_and_typography(
    tmp_path: Path, rust_executable: Path, postprocess: bool, native: bool,
) -> None:
    """Catch lost field switches, malformed XML, and formatting loss in body and page stories."""
    source = tmp_path / "input.md"
    text = (ROOT / "tests/snapshot_cases/word_fields/word_fields.md").read_text(encoding="utf-8")
    source.write_text(text.replace("  mathtype: false", f"  mathtype: false\n  docxNativeCrossref: {str(native).lower()}"), encoding="utf-8")
    output = tmp_path / "output.docx"
    result = build_insertions(rust_executable, source, output, postprocess=postprocess)
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    instructions = [node.text for node in document.element.iter(qn("w:instrText"))]
    assert instructions == [
        "NUMPAGES", "STYLEREF 1", r"SEQ Example \* ARABIC", 'DATE \\@ "yyyy-MM-dd"',
        r"SEQ 自定义 \* ARABIC", 'IF 1 < 2 "A & B" "C > D"', "NUMPAGES", "PAGE", "NUMPAGES",
        "STYLEREF 1",
    ]
    total = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Total pages:"))
    assert total.text == "Total pages: 1."
    assert next(run for run in total.runs if run.text == "1").bold
    chinese = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Chinese instruction:"))
    assert next(run for run in chinese.runs if run.text == "一").italic
    assert "XML characters: A & B." in [paragraph.text for paragraph in document.paragraphs]
    assert "No initial result: ." in [paragraph.text for paragraph in document.paragraphs]
    assert "Identical instruction and result: PAGE." in [paragraph.text for paragraph in document.paragraphs]
    assert sum(paragraph.text == "Current heading: Heading pending." for paragraph in document.paragraphs) == 2
    spaces = next(node for node in document.element.iter(qn("w:t")) if node.text == " 1 ")
    assert spaces.get(qn("xml:space")) == "preserve"
    first, second = document.sections
    for section in (first, second):
        assert [node.text for node in section.header._element.iter(qn("w:instrText"))] == ["NUMPAGES", "STYLEREF 1"]
        assert [node.text for node in section.footer._element.iter(qn("w:instrText"))] == (
            ["PAGE", "NUMPAGES"] if section is first else ["PAGE"]
        )
    with ZipFile(output) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                assert b"PMT_DOCX_FIELD:" not in archive.read(name)


@pytest.mark.parametrize("postprocess,native", [(True, False), (True, True), (False, False)])
def test_section_page_variants_preserve_visible_stories_and_layout(
    tmp_path: Path, rust_executable: Path, postprocess: bool, native: bool,
) -> None:
    """Catch global-toggle blank pages, lost variant content, and broken inheritance or edge distances."""
    source = tmp_path / "input.md"
    text = (ROOT / "tests/snapshot_cases/docx_footer/docx_footer.md").read_text(encoding="utf-8")
    source.write_text(text.replace("  mathtype: false", f"  mathtype: false\n  docxNativeCrossref: {str(native).lower()}"), encoding="utf-8")
    output = tmp_path / "output.docx"
    result = build_insertions(rust_executable, source, output, postprocess=postprocess)
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    initial, alternate, triple, inherited, copied, uniform, cleared = document.sections
    assert document.settings.odd_and_even_pages_header_footer
    assert initial.header.paragraphs[0].text == initial.even_page_header.paragraphs[0].text == "Review copy - 1 -"
    assert initial.footer.paragraphs[0].text == initial.even_page_footer.paragraphs[0].text == "Confidential - 1 -"
    assert alternate.header.paragraphs[0].text == "Odd A"
    assert alternate.even_page_header.paragraphs[0].text == "Even B"
    assert alternate.footer.paragraphs[0].text == alternate.even_page_footer.paragraphs[0].text == "Confidential - 1 -"
    assert alternate.header.paragraphs[0].runs[0].bold
    assert alternate.even_page_header.paragraphs[0].runs[0].italic
    assert not initial.different_first_page_header_footer
    for section in (triple, inherited, copied):
        assert section.different_first_page_header_footer
        assert section.header.paragraphs[0].text.startswith("Default header")
        assert section.even_page_header.paragraphs[0].text.startswith("Even header")
        assert section.first_page_header.paragraphs[0].text == "First header"
        assert section.first_page_header.tables[0].cell(1, 0).text == "Table"
        assert section.footer.paragraphs[0].text == "Default - 1 -"
        assert section.even_page_footer.paragraphs[0].text == "Even - 1 -"
        assert section.first_page_footer.paragraphs[0].text == "First - I -"
        assert section.header_distance.twips == 1417
        assert section.footer_distance.twips == 1080
        for story, target in [(section.header, "default"), (section.even_page_header, "even")]:
            assert any(rel.target_ref == f"https://example.org/{target}" for rel in story.part.rels.values())
    for story_name in ("header", "even_page_header", "first_page_header", "footer", "even_page_footer", "first_page_footer"):
        assert getattr(inherited, story_name).is_linked_to_previous
        assert not getattr(copied, story_name).is_linked_to_previous
        assert getattr(copied, story_name).part.partname != getattr(triple, story_name).part.partname
    assert alternate.header_distance.twips == 1701
    assert alternate.footer_distance.twips == 1304
    assert uniform.header_distance.twips == cleared.header_distance.twips == 0
    assert uniform.footer_distance.twips == cleared.footer_distance.twips == 360
    assert not uniform.different_first_page_header_footer
    assert not cleared.different_first_page_header_footer
    assert uniform.header.paragraphs[0].text == uniform.even_page_header.paragraphs[0].text == "Unified header"
    assert uniform.footer.paragraphs[0].text == uniform.even_page_footer.paragraphs[0].text == "Unified 1"
    assert cleared.header.paragraphs[0].text == "Unified header"
    assert cleared.even_page_header.paragraphs[0].text == ""
    assert cleared.footer.paragraphs[0].text == cleared.even_page_footer.paragraphs[0].text == "Unified 1"
    with ZipFile(output) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                assert b"PMT_DOCX_" not in archive.read(name)


def test_partial_section_variants_preserve_other_stories(tmp_path: Path, rust_executable: Path) -> None:
    """Preserve the untouched footer on first-page activation and override shared fallbacks per variant."""
    source = tmp_path / "input.md"
    source.write_text('''---
papperSettings:
  mathtype: false
  docxHeader: {content: Original}
  docxFooter: {content: Original footer}
---
Before.

[]{t="sectionBreak" headerDefault="Interim header"}

Uniform before any even variant.

[]{t="sectionBreak" headerFirst="First only" footerDefault="New default footer"}

First-page change.

[]{t="sectionBreak" header="Shared" headerEven="Even override" footer="Common" footerFirst="First override"}

Mixed fallbacks.
''', encoding="utf-8")
    output = tmp_path / "output.docx"
    result = build_insertions(rust_executable, source, output)
    assert result.returncode == 0, result.stdout + result.stderr
    _, interim, first, mixed = Document(output).sections
    assert interim.header.paragraphs[0].text == interim.even_page_header.paragraphs[0].text == "Interim header"
    assert first.first_page_header.paragraphs[0].text == "First only"
    assert first.header.paragraphs[0].text == "Interim header"
    assert first.even_page_header.paragraphs[0].text == "Interim header"
    assert first.first_page_footer.paragraphs[0].text == "New default footer"
    assert first.footer.paragraphs[0].text == "New default footer"
    assert first.even_page_footer.paragraphs[0].text == "New default footer"
    assert mixed.header.paragraphs[0].text == mixed.first_page_header.paragraphs[0].text == "Shared"
    assert mixed.even_page_header.paragraphs[0].text == "Even override"
    assert mixed.footer.paragraphs[0].text == mixed.even_page_footer.paragraphs[0].text == "Common"
    assert mixed.first_page_footer.paragraphs[0].text == "First override"


@pytest.mark.parametrize("enabled", [False, True])
def test_section_variants_preserve_active_reference_stories(
    tmp_path: Path, rust_executable: Path, enabled: bool,
) -> None:
    """Keep active template stories, but do not reveal dormant stories when a mode is newly enabled."""
    document = Document()
    document.settings.odd_and_even_pages_header_footer = enabled
    section = document.sections[0]
    section.different_first_page_header_footer = enabled
    section.header.paragraphs[0].text = "Reference normal header"
    section.even_page_header.paragraphs[0].text = "Reference even header"
    section.first_page_header.paragraphs[0].text = "Reference first header"
    section.footer.paragraphs[0].text = "Reference normal footer"
    section.even_page_footer.paragraphs[0].text = "Reference even footer"
    section.first_page_footer.paragraphs[0].text = "Reference first footer"
    reference = tmp_path / "reference.docx"
    document.save(reference)
    source = tmp_path / "input.md"
    source.write_text('''---
papperSettings:
  docxHeader: null
  docxFooter: null
---
Original section.

[]{t="sectionBreak" headerEven="New even header" headerFirst="New first header"}

Changed section.
''', encoding="utf-8")
    output = tmp_path / "output.docx"
    result = build_insertions(rust_executable, source, output, reference=reference)
    assert result.returncode == 0, result.stdout + result.stderr
    initial, changed = Document(output).sections
    assert initial.header.paragraphs[0].text == "Reference normal header"
    assert initial.even_page_header.paragraphs[0].text == (
        "Reference even header" if enabled else "Reference normal header"
    )
    assert changed.header.paragraphs[0].text == "Reference normal header"
    assert changed.even_page_header.paragraphs[0].text == "New even header"
    assert changed.first_page_header.paragraphs[0].text == "New first header"
    assert changed.footer.paragraphs[0].text == "Reference normal footer"
    assert changed.even_page_footer.paragraphs[0].text == (
        "Reference even footer" if enabled else "Reference normal footer"
    )
    assert changed.first_page_footer.paragraphs[0].text == (
        "Reference first footer" if enabled else "Reference normal footer"
    )


@pytest.mark.parametrize("markdown,message", [
    ('[]{t="toc" depth="10"}', "depth must be between"),
    ('[]{t="listofFigures" includeSubfigures="yes"}', "must be true or false"),
    ('[]{t="listofFigures" lang="cn"}', "Unknown DOCX listofFigures attribute: lang"),
    ('[]{t="listofTables" lang="en"}', "Unknown DOCX listofTables attribute: lang"),
    ('[]{t="listofFigures" caption-en="yes"}', "caption-en must be true or false"),
    ('[]{t="listofTables" caption-en="EN"}', "caption-en must be true or false"),
    ('[]{t="sectionBreak" pageStart="-1"}', "pageStart must be an integer"),
    ('[]{t="sectionBreak" header="Text" headerLink="true"}', "conflicts with headerLink"),
    ('[]{t="sectionBreak" headerEven="Text" headerLink="true"}', "conflicts with headerLink"),
    ('[]{t="sectionBreak" footerFirst="Text" footerLink="true"}', "conflicts with footerLink"),
    ('[]{t="sectionBreak" headerFirst="Text" differentFirstPage="false"}', "conflicts with differentFirstPage"),
    ('[]{t="sectionBreak" differentFirstPage="yes"}', "must be true or false"),
    ('[]{t="sectionBreak" headerDistance="-1cm"}', "must be a nonnegative Word length"),
    ('[]{t="sectionBreak" footerDistance="far"}', "Invalid sectionBreak footerDistance"),
    ('[]{t="sectionBreak" footerDistance="999999999999in"}', "must be a nonnegative Word length"),
    ('[]{t="sectionBreak" headerAlignment="left"}', "requires header content"),
    ('Text []{t="toc"}', "must occupy its own paragraph"),
    ('> []{t="sectionBreak"}', "top-level body paragraph"),
    ('1. []{t="sectionBreak"}', "top-level body paragraph"),
    ('[Text]{t="pageBreak"}', "spans must be empty"),
    ('[]{t="date" format="unsupported"}', "Unsupported DOCX date format"),
    ('[]{t="field"}', "field instruction must not be empty"),
    ('[]{t="field" instruction=""}', "field instruction must not be empty"),
    ('[]{t="field" instruction="&#32;"}', "field instruction must not be empty"),
    ('[]{t="field" instruction="PAGE&#10;NUMPAGES"}', "single line without control characters"),
    ('[]{t="unknown"}', "Unsupported DOCX directive"),
])
def test_invalid_insertions_fail_without_replacing_existing_output(
    tmp_path: Path, rust_executable: Path, markdown: str, message: str,
) -> None:
    """Catch invalid layout requests before publishing over an existing output file."""
    source = tmp_path / "input.md"
    source.write_text(markdown, encoding="utf-8")
    output = tmp_path / "output.docx"
    output.write_bytes(b"existing output")
    result = build_insertions(rust_executable, source, output)
    assert result.returncode != 0
    assert message in result.stdout + result.stderr
    assert output.read_bytes() == b"existing output"


def test_native_insertions_have_no_effect_on_html(tmp_path: Path, rust_executable: Path) -> None:
    """A DOCX-only directive must not create directories or page-story content in HTML."""
    source = tmp_path / "input.md"
    source.write_text('Before []{t="date"}. []{t="field" instruction="NUMPAGES" result="DOCX field"}\n\n[]{t="toc" title="DOCX contents"}\n\n'
                      '[]{t="sectionBreak" header="DOCX header"}\n\nAfter.\n', encoding="utf-8")
    output = tmp_path / "output.html"
    result = build_insertions(rust_executable, source, output, target="html")
    assert result.returncode == 0, result.stdout + result.stderr
    text = html.fromstring(output.read_text(encoding="utf-8")).text_content()
    assert "Before" in text and "After." in text
    assert "DOCX contents" not in text and "DOCX header" not in text and "DOCX field" not in text
