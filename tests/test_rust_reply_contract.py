"""Verify native reviewer replies using actual document content, styles, and PDF geometry."""

from __future__ import annotations

import json
import os
from pathlib import Path
import shutil
import subprocess
from zipfile import ZipFile

import pytest


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def reply_executable(rust_executable: Path) -> Path:
    """Share the current session's owned Rust CLI with reply contracts."""
    return rust_executable


class ReplyProject:
    """Run the native product CLI against an isolated set of reviewer documents."""

    def __init__(self, directory: Path, executable: Path) -> None:
        """Create a project with local references and stable settings for real parity checks."""
        self.directory = directory
        self.executable = executable
        self.environment = {**os.environ, "PAPPER_HOME": str(directory / "state"), "PAPPER_RESOURCE_ROOT": str(ROOT)}
        (directory / "style.yml").write_text("mathtype: false\nlang: en-US\ncsl: numeric.csl\nbibliography: references.bib\nnumberSections: true\nsecPrefix: Section\nfigPrefix: Figure\ntblPrefix: Table\neqnPrefix: Equation\n", encoding="utf-8")
        shutil.copy2(ROOT / "pandoc/csl/elsevier-vancouver.csl", directory / "numeric.csl")
        (directory / "references.bib").write_text("@article{alpha, author={Alice Smith}, title={Alpha}, journal={J}, year={2020}}\n@article{beta, author={Bob Jones}, title={Beta}, journal={J}, year={2021}}\n", encoding="utf-8")
        (directory / "diagram.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="20" height="10"><rect width="20" height="10" fill="blue"/></svg>', encoding="utf-8")
        (directory / "manuscript.md").write_text("---\ntitle: Manuscript\n---\n\n# Introduction {#sec:intro}\n\nCitations [@alpha; @beta].\n\n![Overview](diagram.svg){#fig:overview}\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n: Metrics. {#tbl:metrics}\n\n$$ z=1 $$ {#eq:prior}\n\n$$ x+y $$ {#eq:sum}\n\n# Detail {#sec:intro.detail}\n", encoding="utf-8")

    def run(self, *arguments: str, success: bool = True) -> subprocess.CompletedProcess[str]:
        """Capture the complete user command and fail with actionable build diagnostics."""
        result = subprocess.run([str(self.executable), "build-reply", *arguments], cwd=self.directory, env=self.environment, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=90)
        if success:
            assert result.returncode == 0, result.stdout + result.stderr
        else:
            assert result.returncode != 0, result.stdout + result.stderr
        return result


@pytest.fixture
def reply_project(tmp_path: Path, reply_executable: Path) -> ReplyProject:
    """Provide independently writable reply assets and caches to each contract."""
    return ReplyProject(tmp_path, reply_executable)


def test_native_reply_txt_references_clusters_and_copied_captions(reply_project: ReplyProject) -> None:
    """Preserve original numbering, cluster punctuation and unresolved bibliography keys."""
    project = reply_project
    text = """We revised Figure @fig:overview and [@sec:intro.detail].

::: {custom-style="Reply to Reviewers"}
Prior work [@alpha; @beta] and bare @alpha; unresolved [@alpha; @missing].<br>
1\\. Updated explanation.

![Overview](diagram.svg) {#fig:overview width=80%}

> : Metrics. {#tbl:metrics}

$$ x+y $$ {#eq:sum}
:::
"""
    (project.directory / "reply.md").write_text(text, encoding="utf-8")
    project.run("reply.md", "-o", "native.txt")
    expected = (project.directory / "native.txt").read_text(encoding="utf-8")
    assert "Figure 1" in expected and "missing?" in expected
    assert "[Image: Figure 1 Overview]" in expected and "1. Updated" in expected


@pytest.mark.parametrize("page_margins", [False, True])
def test_native_reply_docx_styles_equations_and_no_author_footnote(reply_project: ReplyProject, page_margins: bool) -> None:
    """Preserve reply styles, configured margins, and the original numbered OMML equation."""
    project = reply_project
    if page_margins:
        with (project.directory / "style.yml").open("a", encoding="utf-8") as handle:
            handle.write("docxPageMargins:\n  left: 2cm\n  right: 2cm\n")
    (project.directory / "reply.md").write_text("---\ntitle: Reviewer reply\nauthor:\n- name: Reply Author\n---\n\nReviewer comment.\n\n::: {custom-style=\"Reply to Reviewers\"}\nWe revised Equation @eq:sum and Table @tbl:metrics.\n\n$$ x+y $$ {#eq:sum}\n:::\n", encoding="utf-8")
    project.run("reply.md", "-o", "native.docx")
    with ZipFile(project.directory / "native.docx") as archive:
        document = archive.read("word/document.xml")
        assert b"footnoteReference" not in document
        assert b"oMath" in document and b"center" in document and b"right" in document
        assert b"Equation 2" in document and b"(2)" in document
    from docx import Document
    output = Document(project.directory / "native.docx")
    if page_margins:
        assert output.sections[0].left_margin.cm == pytest.approx(2, abs=0.001)
        assert output.sections[0].right_margin.cm == pytest.approx(2, abs=0.001)
    # Numbered reply equations must use the same available width as the actual page.
    from docx.oxml.ns import qn
    section = output.sections[0]
    text_width = section.page_width.twips - section.left_margin.twips - section.right_margin.twips
    equations = [paragraph for paragraph in output.paragraphs if "(2)" in paragraph.text]
    assert equations
    tabs = equations[0]._p.findall(f".//{qn('w:tab')}")
    positions = {tab.get(qn("w:val")): int(tab.get(qn("w:pos")))
                 for tab in tabs if tab.get(qn("w:pos")) is not None}
    assert positions["center"] == round(text_width / 2)
    assert positions["right"] == text_width
    replies = [paragraph for paragraph in output.paragraphs if "We revised" in paragraph.text]
    assert replies and replies[0].style.name == "Reply to Reviewers"
    assert replies[0].style.font.color.rgb is not None


@pytest.mark.parametrize("body,expected_rows", [
    pytest.param("| Group | !<! | !<! |\n", [["Group", "Group", "Group"]], id="horizontal"),
    pytest.param(
        "| Group | Second | Third |\n| !^! | Fourth | Fifth |\n| !^! | Sixth | Seventh |\n",
        [["Group", "Second", "Third"], ["Group", "Fourth", "Fifth"], ["Group", "Sixth", "Seventh"]],
        id="vertical",
    ),
    pytest.param(
        "| Group | !<! | Third |\n| !^! | !<! | Fifth |\n",
        [["Group", "Group", "Third"], ["Group", "Group", "Fifth"]],
        id="combined",
    ),
])
def test_native_reply_docx_merged_cells(
    reply_project: ReplyProject, body: str, expected_rows: list[list[str]],
) -> None:
    """Verify real reply cell spans and preserved neighbors, including chained markers."""
    from docx import Document

    project = reply_project
    (project.directory / "reply.md").write_text(
        '::: {custom-style="Reply to Reviewers"}\n\n'
        "| A | B | C |\n|---|---|---|\n" + body + "\n:::\n",
        encoding="utf-8",
    )
    project.run("reply.md", "-o", "native.docx")
    output = Document(project.directory / "native.docx")
    assert len(output.tables) == 1
    table = output.tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ["A", "B", "C"], *expected_rows,
    ]
    # Repeated text must come from one merged Word cell, rather than copied values.
    anchor = table.cell(1, 0)._tc
    for row_index, values in enumerate(expected_rows, start=1):
        for column_index, value in enumerate(values):
            assert (table.cell(row_index, column_index)._tc is anchor) == (value == "Group")


def test_native_reply_and_manuscript_share_docx_defaults(reply_project: ReplyProject) -> None:
    """Catch configuration drift by applying an authored output filter to both builds."""
    from docx import Document
    from native_support import native_pandoc_executable
    import yaml

    engine = native_pandoc_executable()
    if engine is None:
        pytest.skip("Shared defaults integration requires the native Pandoc engine")
    project = reply_project
    resources = project.directory / "resources"
    shutil.copytree(ROOT / "pandoc", resources / "pandoc")
    # Reuse the real worker without copying its large binary into this isolated resource root.
    record = resources / ".pmt/pandoc-worker/current.json"
    record.parent.mkdir(parents=True)
    record.write_text(json.dumps({"executable": str(engine)}), encoding="utf-8")
    project.environment["PAPPER_RESOURCE_ROOT"] = str(resources)
    filter_path = resources / "pandoc/filters/shared/output_annotation.lua"
    filter_path.write_text(
        '-- Annotate rendered prose to verify shared defaults through actual output.\n'
        '-- Prefix one paragraph while retaining its resolved references and styles.\n'
        'function Para(paragraph)\n'
        '  paragraph.content:insert(1, pandoc.Space())\n'
        '  paragraph.content:insert(1, pandoc.Str("Shared-render:"))\n'
        '  return paragraph\n'
        'end\n',
        encoding="utf-8",
    )
    defaults_path = resources / "pandoc/pandoc-docx.yml"
    defaults = yaml.safe_load(defaults_path.read_text(encoding="utf-8"))
    defaults["filters"].append("${.}/filters/shared/output_annotation.lua")
    defaults_path.write_text(yaml.safe_dump(defaults, sort_keys=False), encoding="utf-8")
    markdown = project.directory / "reply.md"
    markdown.write_text("Shared prose reaches both DOCX pipelines.\n", encoding="utf-8")
    project.run("reply.md", "-o", "reply.docx")
    built = subprocess.run(
        [str(project.executable), "build", "docx", "-m", "reply.md", "-o", "manuscript.docx", "--no-mathtype"],
        cwd=project.directory, env=project.environment, capture_output=True,
        text=True, encoding="utf-8", errors="replace", timeout=90,
    )
    assert built.returncode == 0, built.stdout + built.stderr
    for name in ["reply.docx", "manuscript.docx"]:
        paragraphs = [paragraph.text for paragraph in Document(project.directory / name).paragraphs]
        assert "Shared-render: Shared prose reaches both DOCX pipelines." in paragraphs


def write_numbered_pdf(path: Path, producer: str, layout: bool) -> None:
    """Generate actual PDF geometry whose stream order differs from its margin layout."""
    import pymupdf

    with pymupdf.open() as document:
        document.set_metadata({"producer": producer})
        page = document.new_page()
        for number, text in [(101, "The distinctive alpha experiment is accurate."), (102, "A second result extends the analysis."), (103, "The repeated marker remains repeated.")]:
            y = 70 + (number - 101) * 20
            if layout:
                page.insert_text((30, y), str(number))
                page.insert_text((75, y), text)
            else:
                page.insert_text((75, y), text)
                page.insert_text((75, y + 10), str(number))
        document.save(path)


@pytest.mark.parametrize("producer,layout", [("Microsoft Word", True), ("LibreOffice", True), ("Independent PDF tool", False)])
def test_native_reply_real_pdf_line_geometry_and_regex(reply_project: ReplyProject, producer: str, layout: bool) -> None:
    """Resolve unique regexes from real PDFs while retaining invalid and ambiguous placeholders."""
    project = reply_project
    write_numbered_pdf(project.directory / "source.pdf", producer, layout)
    (project.directory / "reply.md").write_text("Changes (Line `distinctive ALPHA`) and spanning (Line `accurate.*A second`).\nInvalid (Line `[`) and ambiguous (Line `repeated`).\n", encoding="utf-8")
    project.run("reply.md", "--manuscript-line-source", "source.pdf", "-o", "native.txt")
    expected = (project.directory / "native.txt").read_text(encoding="utf-8")
    assert "Changes (Line 101)" in expected and "spanning (Line 101)" in expected
    assert "(Line `[`)" in expected and "(Line `repeated`)" in expected


def test_native_pdf_extract_cjk_metadata(reply_executable: Path, tmp_path: Path) -> None:
    """Decode predefined Chinese CMaps and Unicode metadata from an unencrypted PDF."""
    import pymupdf

    source = tmp_path / "中文稿件.pdf"
    expected = "中文测试：论文第十六行，方法改进。"
    with pymupdf.open() as document:
        document.set_metadata({"producer": "Microsoft Word", "title": "中文论文标题"})
        page = document.new_page()
        # Predefined Chinese CMaps must decode without an embedded ToUnicode map.
        page.insert_text((75, 70), expected, fontname="china-s")
        document.save(source)
    result = subprocess.run([str(reply_executable), "__pdf_extract", str(source)], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    extracted = json.loads(result.stdout)
    assert extracted["metadata"]["Title"] == "中文论文标题"
    assert extracted["metadata"]["Producer"] == "Microsoft Word"
    assert len(extracted["pages"]) == 1
    assert extracted["pages"][0]["text"].strip() == expected


def test_native_reply_bad_pdf_and_missing_input_preserve_output(reply_project: ReplyProject) -> None:
    """Failures in the real isolated PDF engine never replace an existing journal submission."""
    project = reply_project
    output = project.directory / "submission.txt"
    output.write_text("Keep this submitted version\n", encoding="utf-8")
    (project.directory / "reply.md").write_text("Changes (Line `actual text`).\n", encoding="utf-8")
    (project.directory / "bad.pdf").write_bytes(b"this is not a PDF")
    result = project.run("reply.md", "--manuscript-line-source", "bad.pdf", "-o", str(output), success=False)
    assert "PDF" in result.stdout + result.stderr
    assert output.read_text(encoding="utf-8") == "Keep this submitted version\n"
    project.run("missing.md", "-o", str(output), success=False)
    assert output.read_text(encoding="utf-8") == "Keep this submitted version\n"


def test_native_reply_nested_companions_and_default_output(reply_project: ReplyProject) -> None:
    """Find companions beside a nested reply and derive the established output basename."""
    project = reply_project
    folder = project.directory / "reviewer"
    folder.mkdir()
    (folder / "manuscript.md").write_text("# First\n\n# Revised {#sec:local}\n", encoding="utf-8")
    (folder / "reply.md").write_text("We revised @sec:local.\n", encoding="utf-8")
    project.run("reviewer/reply.md")
    path = project.directory / "output/docx/reply.docx"
    assert path.is_file()
    with ZipFile(path) as archive:
        assert b"Section 2" in archive.read("word/document.xml")


def test_native_reply_template_pdf_line_anchors(reply_project: ReplyProject) -> None:
    """Resolve anchors in the retained multi-page paper fixture with real native extraction."""
    project = reply_project
    lines = [(16, "coherent narrative that guides readers from the general context"), (45, "caption attributes above are applied to the DOCX table"), (68, r"procedure is useful\. In this template, pseudocode"), (75, "synthetic trend chart in Figure 1"), (140, "authors declare no conflict of interest")]
    (project.directory / "reply.md").write_text("\n".join(f"Anchor {number}: (Line `{pattern}`)." for number, pattern in lines) + "\n", encoding="utf-8")
    project.run("reply.md", "--manuscript-line-source", str(ROOT / "tests/fixtures/template-manuscript.pdf"), "-o", "native.txt")
    expected = (project.directory / "native.txt").read_text(encoding="utf-8")
    for number, _ in lines:
        assert f"Anchor {number}: (Line {number})." in expected


def test_native_reply_word_line_source_rebuild_and_timestamp_cache(reply_project: ReplyProject) -> None:
    """Reuse layout-identical source packages and invalidate PDF layout after a body edit."""
    if os.name == "nt":
        import winreg

        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application"):
                pass
        except FileNotFoundError:
            pytest.skip("Microsoft Word is required for the Windows native line-source integration")
    elif not shutil.which("soffice"):
        pytest.skip("LibreOffice is required for the native line-source integration")
    project = reply_project
    source = project.directory / "line-source.md"
    source.write_text("---\ntitle: Layout source\n---\n\nThe distinctive native line-source experiment is accurate.\n", encoding="utf-8")
    with (project.directory / "style.yml").open("a", encoding="utf-8") as handle:
        handle.write("docxShowLineNumbers: true\n")
    reply = project.directory / "reply.md"
    reply.write_text("Changes (Line `distinctive native line-source experiment`).\n", encoding="utf-8")
    project.run("reply.md", "--manuscript-line-source", str(source), "-o", "first.txt")
    first = (project.directory / "first.txt").read_text(encoding="utf-8")
    assert "`" not in first and "(Line " in first
    pdfs = list((project.directory / "state").glob("projects/*/cache/reply/line-source/pdf/*.pdf"))
    assert len(pdfs) == 1
    original_pdf = pdfs[0].read_bytes()
    original_time = pdfs[0].stat().st_mtime_ns
    project.run("reply.md", "--manuscript-line-source", str(source), "-o", "second.txt")
    assert (project.directory / "second.txt").read_text(encoding="utf-8") == first
    assert pdfs[0].read_bytes() == original_pdf and pdfs[0].stat().st_mtime_ns == original_time
    source.write_text("---\ntitle: Layout source\n---\n\nA new leading paragraph changes subsequent line positions.\n\nThe distinctive native line-source experiment is accurate.\n", encoding="utf-8")
    project.run("reply.md", "--manuscript-line-source", str(source), "-o", "third.txt")
    third = (project.directory / "third.txt").read_text(encoding="utf-8")
    assert "`" not in third and third != first
    assert len(list((project.directory / "state").glob("projects/*/cache/reply/line-source/pdf/*.pdf"))) == 2
    assert pdfs[0].read_bytes() == original_pdf


def test_native_reply_mathtype_retains_original_manuscript_number(reply_project: ReplyProject) -> None:
    """Publish an actual embedded MathType object with the original equation number and no markers."""
    project = reply_project
    with (project.directory / "style.yml").open("a", encoding="utf-8") as handle:
        handle.write("mathtype: true\nmathtypeConversionMethod: rust\nmathtypeSvgBackend: typst\nmathtypeTypstMathFont: XITS Math\n")
    (project.directory / "reply.md").write_text("We revised Equation @eq:sum.\n\n$$ x+y $$ {#eq:sum}\n", encoding="utf-8")
    project.run("reply.md", "-o", "native.docx")
    with ZipFile(project.directory / "native.docx") as archive:
        document = archive.read("word/document.xml")
        assert b"OLEObject" in document and b"MTLATEX:" not in document
        assert b"Equation 2" in document and b"(2)" in document
        embedded = [name for name in archive.namelist() if name.startswith("word/embeddings/")]
        assert len(embedded) == 1
        assert archive.read(embedded[0]).startswith(bytes.fromhex("d0cf11e0a1b11ae1"))


def test_native_reply_literal_attribute_example_without_manuscript(reply_project: ReplyProject) -> None:
    """Literal syntax examples do not require a manuscript or alter escaped non-list prose."""
    project = reply_project
    (project.directory / "reply.md").write_text("We explained the literal attribute `{#fig:literal}`.\n\n1\\.example remains literal.\n", encoding="utf-8")
    project.run("reply.md", "--reply-manuscript", "absent.md", "-o", "native.txt")
    assert "`{#fig:literal}`" in (project.directory / "native.txt").read_text(encoding="utf-8")
    assert "1\\.example" in (project.directory / "native.txt").read_text(encoding="utf-8")


@pytest.mark.parametrize("target", ["html", "docx"])
def test_build_detects_reply_header_and_preserves_selected_manuscript_numbers(
    reply_project: ReplyProject, target: str,
) -> None:
    """A nested reply uses its YAML manuscript and explicit style, including copied equation numbers."""
    from docx import Document
    from lxml import html

    project = reply_project
    folder = project.directory / "responses"
    folder.mkdir()
    manuscript = folder / "selected paper.md"
    manuscript.write_text("$$ a=1 $$ {#eq:first}\n\n$$ b=2 $$ {#eq:second}\n\n$$ x+y $$ {#eq:sum}\n", encoding="utf-8")
    (folder / "style.yml").write_text("reply:\n  pandocMetadata:\n    title: Wrong local title\n", encoding="utf-8")
    (project.directory / "selected.yml").write_text(
        "mathtype: true\nreply:\n  pandocMetadata:\n    title: Selected reply title\n", encoding="utf-8",
    )
    reply = folder / "answer.md"
    reply.write_bytes(("\ufeff---\r\nreply: selected paper.md\r\n---\r\n\r\n"
                       '::: {custom-style="Reply to Reviewers"}\r\n'
                       "We revised Equation @eq:sum.\r\n\r\n$$ x+y $$ {#eq:sum}\r\n:::\r\n").encode("utf-8"))
    output = project.directory / f"answer.{target}"
    command = [str(project.executable), "build", target, "responses/answer.md", "-o", str(output),
               "--style-file", "selected.yml"]
    if target == "docx":
        command.append("--no-mathtype")
    result = subprocess.run(command, cwd=project.directory, env=project.environment, capture_output=True,
                            text=True, encoding="utf-8", timeout=90)
    assert result.returncode == 0, result.stdout + result.stderr
    if target == "docx":
        paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
        assert "Selected reply title" in paragraphs and "Wrong local title" not in paragraphs
        assert "We revised Equation 3." in paragraphs
        assert any("(3)" in paragraph for paragraph in paragraphs)
        with ZipFile(output) as archive:
            assert b"OLEObject" not in archive.read("word/document.xml")
    else:
        rendered = html.fromstring(output.read_text(encoding="utf-8"))
        assert "Selected reply title" in rendered.text_content() and "Wrong local title" not in rendered.text_content()
        assert "We revised Equation 3." in rendered.text_content()
        assert "\\tag{3}" in rendered.xpath('string(//span[contains(@class,"math")])')


@pytest.mark.parametrize("target", ["html", "docx"])
@pytest.mark.parametrize("selector", ["null", "false", "[]", '""', "missing.md"])
def test_build_invalid_reply_header_preserves_existing_output(
    reply_project: ReplyProject, target: str, selector: str,
) -> None:
    """Invalid reply selectors fail before publishing a replacement submission."""
    project = reply_project
    (project.directory / "reply.md").write_text(f"---\nreply: {selector}\n---\n\nResponse.\n", encoding="utf-8")
    output = project.directory / f"submission.{target}"
    output.write_bytes(b"Keep submitted version")
    result = subprocess.run([str(project.executable), "build", target, "reply.md", "-o", str(output)],
                            cwd=project.directory, env=project.environment, capture_output=True,
                            text=True, encoding="utf-8", timeout=90)
    assert result.returncode != 0
    assert "reply" in result.stderr.lower()
    assert output.read_bytes() == b"Keep submitted version"
