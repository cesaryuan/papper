"""Build DOCX fixtures and compare Word and DOCXIDE against independent page baselines.

Run ``uv run --no-sync pytest tests/test_docx_visual.py --visual`` on Windows with
Microsoft Word and the pinned DOCXIDE installation. Use --docx-backend to select
one renderer; both run by default. Baselines are explicit (--visual-update) and separate from ZIP/XML
content snapshots. Each export opens a unique read-only copy, closes only that copy,
and verifies existing user documents stay open; it never quits the Word application.
PyMuPDF rasterizes print-quality PDFs at 144 DPI, so pagination, fonts, formulas,
tables, revision text and image layout are tested independently of XML serialization.
"""

from __future__ import annotations

import json
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest
import pymupdf
from PIL import Image, ImageChops
from docx import Document
from docx.shared import RGBColor
from lxml import etree

from docx_visual_support import VISUAL_PLATFORM, DocxVisualSession, DocxideVisualSession, WordVisualSession
from native_support import papper_command
from test_build_snapshots import CASE_ROOT, CASES, build_case, copy_case


def freeze_visual_dates(source: Path, destination: Path) -> None:
    """Lock DATE caches only in the visual copy so field updates cannot depend on today's date."""
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    dates = 0
    with ZipFile(source) as original, ZipFile(destination, 'w') as frozen:
        for member in original.infolist():
            data = original.read(member.filename)
            if member.filename.startswith('word/') and member.filename.endswith('.xml'):
                root = etree.fromstring(data)
                for paragraph in root.iter(namespace + 'p'):
                    begin, in_date, result = None, False, False
                    for run in paragraph.iter(namespace + 'r'):
                        field = run.find(namespace + 'fldChar')
                        if field is not None:
                            kind = field.get(namespace + 'fldCharType')
                            if kind == 'begin':
                                begin = field
                            elif kind == 'separate':
                                result = True
                            elif kind == 'end':
                                begin, in_date, result = None, False, False
                        instruction = run.find(namespace + 'instrText')
                        if instruction is not None and (instruction.text or '').strip().startswith('DATE '):
                            assert begin is not None, 'DATE field is missing its beginning'
                            begin.set(namespace + 'fldLock', 'true')
                            begin.set(namespace + 'dirty', 'false')
                            in_date = True
                            dates += 1
                        if in_date and result:
                            for value in run.iter(namespace + 't'):
                                value.text = '2000-01-02'
                data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            frozen.writestr(member, data)
    assert dates > 0, 'Native insertion visual fixture lost its DATE field'


@pytest.fixture(scope="session")
def docx_visual_session(
    docx_backend: str, request: pytest.FixtureRequest, tmp_path_factory: pytest.TempPathFactory,
) -> DocxVisualSession:
    """Initialize only the selected renderer; DOCXIDE-only runs never attach to Word."""
    renderer = WordVisualSession if docx_backend == "word" else DocxideVisualSession
    return renderer(tmp_path_factory.mktemp(f"{docx_backend}-visual"), request.config.getoption("--visual-update"))


@pytest.fixture(scope="module")
def docx_visual_inputs() -> dict[str, Path]:
    """Keep the exact same generated DOCX bytes available across backend parameter groups."""
    return {}


@pytest.fixture(scope="module")
def docx_visual_input(
    case_name: str, tmp_path_factory: pytest.TempPathFactory, docx_visual_inputs: dict[str, Path],
) -> Path:
    """Build an immutable DOCX input shared by the two renderer checks for this case."""
    if case_name in docx_visual_inputs:
        return docx_visual_inputs[case_name]
    tmp_path = tmp_path_factory.mktemp(f"docx-{case_name}")
    project = tmp_path / case_name
    if case_name == "reply":
        copy_case(CASE_ROOT / "reply", project)
        output = project / "rendered.docx"
        result = subprocess.run(
            [*papper_command(), "build", "docx", str(project / "reply.md"),
             "--manuscript-line-source", str(project / "manuscript.pdf"), "-o", str(output)],
            cwd=project, capture_output=True, text=True, encoding="utf-8", timeout=60,
        )
        assert result.returncode == 0, result.stdout + result.stderr
    else:
        case_dir, markdown = CASES[case_name]
        copy_case(case_dir, project)
        if case_name in {"bilingual_captions", "docx_directives"}:
            # Relative sibling image references must resolve before Word pagination.
            copy_case(CASES["crossrefs"][0], tmp_path / "crossrefs")
        output = project / "rendered.docx"
        style_file = project / "style.yml" if case_name in {"native_crossrefs", "equation_attributes_no_mathtype"} else None
        build_case(project, markdown, "docx", output, style_file=style_file)
    update_fields = case_name in {"docx_directives", "word_fields"}
    if update_fields:
        visual_copy = project / "visual-fixed-date.docx"
        freeze_visual_dates(output, visual_copy)
        output = visual_copy
    docx_visual_inputs[case_name] = output
    return output


@pytest.mark.visual
@pytest.mark.parametrize("case_name", [*CASES, "reply"], scope="module")
def test_docx_render_matches_visual_snapshot(
    case_name: str, docx_visual_input: Path, docx_visual_session: DocxVisualSession,
) -> None:
    """Protect each renderer's pages while retaining Word's stronger field/layout assertions."""
    output = docx_visual_input
    update_fields = case_name in {"docx_directives", "word_fields"}
    baseline = CASE_ROOT / case_name / "snapshots-visual" / VISUAL_PLATFORM / f"docx-{docx_visual_session.backend}"
    docx_visual_session.check(output, baseline, update_fields=update_fields)
    # Word-specific semantics remain authoritative while DOCXIDE has the field and
    # layout limitations recorded in upstream issue #142. Its own pixels still gate changes.
    if docx_visual_session.backend != "word":
        return
    if case_name == "docx_footer":
        # Verify Word chooses the right live story and computes its page number;
        # XML snapshots alone cannot catch the wrong story being rendered on a page.
        with pymupdf.open(output.parent / "word-render/export-2/rendered.pdf") as rendered:
            expected = [
                ("Review copy - 1 -", "Confidential - 1 -"),
                ("Review copy - 2 -", "Confidential - 2 -"),
                ("Odd A", "Confidential - 3 -"),
                ("Even B", "Confidential - 4 -"),
                ("Odd A", "Confidential - 5 -"),
                ("First header", "First - VI -"),
                ("Default header", "Default - 7 -"),
                ("Even header", "Even - 8 -"),
                ("First header", "First - IX -"),
                ("First header", "First - X -"),
                ("Unified header", "Unified 11"),
                ("Unified header", "Unified 12"),
                ("Unified header", "Unified 13"),
                ("", "Unified 14"),
            ]
            for page, (header, footer) in zip(rendered, expected, strict=True):
                text = page.get_text()
                assert header in text and footer in text
            assert "Table" in rendered[5].get_text() and "Cell" in rendered[5].get_text()
            assert "Unified header" not in rendered[13].get_text()
    if case_name == "paragraph_format":
        # Word can let inherited character indents override direct physical
        # indents; XML property checks alone cannot detect the wrong placement.
        margin = Document(output).sections[0].left_margin.pt
        with pymupdf.open(output.parent / "word-render/export-2/rendered.pdf") as rendered:
            spans = [span for block in rendered[0].get_text("dict")["blocks"]
                     for line in block.get("lines", []) for span in line["spans"]]
            physical = next(span for span in spans if span["text"].startswith("Chinese Word size"))
            inherited = next(span for span in spans if span["text"].startswith("Only the font size"))
            character = next(span for span in spans if span["text"].startswith("One-character local"))
            hanging = next(span for span in spans if span["text"].startswith("Local physical hanging"))
            assert abs(physical["bbox"][0] - margin - 12) < 0.2
            assert abs(inherited["bbox"][0] - margin - 18) < 0.2
            assert abs(character["bbox"][0] - margin - 10) < 0.2
            assert abs(hanging["bbox"][0] - margin - 12) < 0.2
            # Word must actually draw the borders; schema-valid XML alone cannot prove this.
            drawings = [drawing for page in rendered for drawing in page.get_drawings()]
            assert len(drawings) >= 20, "Word lost the paragraph border drawings"
            assert any(drawing.get("color") and drawing["color"][2] > 0.4
                       and drawing["color"][0] < 0.1 for drawing in drawings), "Word lost the authored blue borders"
    if case_name == "docx_directives":
        # Check real Word output as well as pixels: a stable placeholder image
        # must not become the baseline when directory updating silently fails.
        with pymupdf.open(output.parent / "word-render/export-2/rendered.pdf") as rendered:
            text = '\n'.join(page.get_text() for page in rendered)
            assert 'Update field in Word.' not in text
            assert '2000-01-02' in text
            for caption in ('Figure 1 Primary figure', 'Figure 2 Combined panels',
                            '(a) First panel', '(b) Second panel', 'Table 1 Measurements'):
                assert text.count(caption) >= 2, f'Directory lost caption: {caption}'
    if case_name == "word_fields":
        # Word must calculate the authored fields, rather than just display caches.
        with pymupdf.open(output.parent / "word-render/export-2/rendered.pdf") as rendered:
            text = '\n'.join(page.get_text() for page in rendered)
            assert f'Total pages: {len(rendered)}.' in text
            assert f'No initial result: {len(rendered)}.' in text
            assert 'Identical instruction and result: 1.' in text
            assert 'XML characters: A & B.' in text
            # STYLEREF must follow each page's heading in both body and header.
            for page, heading in zip(rendered, ('1 Authored instructions', '2 Next section'), strict=True):
                page_text = page.get_text()
                assert f'Chapter: {heading}' in page_text
                assert f'Current heading: {heading}.' in page_text
                assert 'Heading pending' not in page_text


@pytest.mark.visual
def test_docx_visual_reports_color_and_pagination_changes(
    tmp_path: Path, docx_visual_session: DocxVisualSession, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Verify both real renderers reject changed text color and additional pages."""
    source = tmp_path / "layout.docx"
    document = Document()
    run = document.add_paragraph().add_run("DOCX visual regression contract")
    document.save(source)
    images, geometry = docx_visual_session.capture(source, tmp_path / "original")
    baseline = tmp_path / "baseline"
    baseline.mkdir()
    for index, image in enumerate(images, 1):
        (baseline / f"page-{index:03d}.png").write_bytes(image)
    (baseline / "pages.json").write_text(json.dumps({"pages": geometry}), encoding="utf-8")
    # Even baseline-generation runs must independently prove visible changes fail.
    monkeypatch.setattr(docx_visual_session, "update", False)
    run.font.color.rgb = RGBColor(255, 0, 0)
    document.save(source)
    with pytest.raises(AssertionError, match="Visual mismatch"):
        docx_visual_session.check(source, baseline)
    with Image.open(baseline / "page-001.png") as original, Image.open(baseline / "page-001-diff.png") as difference:
        assert ImageChops.difference(original, difference).getbbox(), "Changed text color produced no visible diff"

    document.add_page_break()
    document.add_paragraph("Additional DOCX page")
    document.save(source)
    with pytest.raises(AssertionError, match="DOCX page count/geometry changed"):
        docx_visual_session.check(source, baseline)
    with pymupdf.open(tmp_path / f"{docx_visual_session.backend}-render/export-2/rendered.pdf") as actual:
        assert actual.page_count > len(geometry), "Failure PDF lost the additional page"
