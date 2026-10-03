"""Verify footer page visibility using the native DOCX package editor."""

import json
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

from test_rust_docx_contract import rust_postprocessor


def footer_page_instructions(doc: Document) -> list[str]:
    """Read PAGE fields from every section's primary footer."""
    return [instruction.text or "" for section in doc.sections
            for paragraph in section.footer.paragraphs
            for instruction in paragraph._p.iter(qn("w:instrText"))]


def test_native_page_numbers_preserve_footer_and_obey_explicit_visibility(tmp_path: Path, rust_postprocessor: Path) -> None:
    """Toggle PAGE fields in linked sections without losing unrelated authored footer text."""
    doc = Document()
    doc.styles.add_style("page number", WD_STYLE_TYPE.CHARACTER)
    doc.sections[0].footer.paragraphs[0].text = "Confidential manuscript"
    doc.add_section()
    source = tmp_path / "source.docx"
    doc.save(source)
    metadata = tmp_path / "metadata.json"
    for index, visible in enumerate((None, True, None, False)):
        metadata.write_text(json.dumps({"pmt_settings": {"values": {"docxShowPageNumbers": visible},
                             "provided": [], "pandoc_metadata": {}, "reply": None},
                             "pandoc_metadata": {}, "has_yaml_header": False}), encoding="utf-8")
        output = tmp_path / f"output-{index}.docx"
        result = subprocess.run([str(rust_postprocessor), str(source), str(output), str(metadata)],
                                capture_output=True, text=True, encoding="utf-8", timeout=30)
        assert result.returncode == 0, result.stdout + result.stderr
        doc = Document(output)
        assert footer_page_instructions(doc) == ([" PAGE ", " PAGE "] if index in (1, 2) else [])
        assert all("Confidential manuscript" in "\n".join(p.text for p in section.footer.paragraphs)
                   for section in doc.sections)
        assert doc.settings.element.find(qn("w:updateFields")) is None
        source = output
