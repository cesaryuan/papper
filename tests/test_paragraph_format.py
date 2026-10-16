"""Verify local paragraph formatting, scope boundaries and failed-build data preservation.

Content snapshots omit HTML CSS; these output assertions protect the formatting
itself and the public CLI failure contract for invalid local typography.
"""

from __future__ import annotations

import os
from pathlib import Path
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree, html
import pytest

from native_support import ROOT


def build(executable: Path, source: Path, output: Path, *, postprocess: bool = True) -> subprocess.CompletedProcess:
    """Build real output through the public CLI, including mandatory finalization mode."""
    return subprocess.run(
        [str(executable), "build", output.suffix[1:], str(source), "-o", str(output),
         *(["--no-mathtype"] if output.suffix == ".docx" else [])],
        cwd=output.parent, capture_output=True, text=True, encoding="utf-8", timeout=90,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT),
             "PMT_ENABLE_DOCX_POSTPROCESS": str(postprocess).lower()},
    )


@pytest.mark.parametrize("target,postprocess", [("html", True), ("docx", True), ("docx", False)])
def test_local_formats_preserve_scopes_and_inline_emphasis(
    tmp_path: Path, rust_executable: Path, target: str, postprocess: bool,
) -> None:
    """Catch dropped formatting, outer-scope leaks and lost explicit inline styles in actual output."""
    source = ROOT / "tests/snapshot_cases/paragraph_format/paragraph_format.md"
    output = tmp_path / f"output.{target}"
    result = build(rust_executable, source, output, postprocess=postprocess)
    assert result.returncode == 0, result.stdout + result.stderr
    if target == "html":
        document = html.fromstring(output.read_text(encoding="utf-8"))
        paragraphs = {" ".join(node.text_content().split()): node for node in document.xpath("//p")}
        outer = next(node for text, node in paragraphs.items() if text.startswith("Outer paragraph with"))
        inner = next(node for text, node in paragraphs.items() if text.startswith("Inner paragraph"))
        assert "font-size: 10pt;" in outer.get("style", "")
        assert "font-weight: normal;" in outer.get("style", "")
        assert "margin-bottom: 6pt;" in outer.get("style", "")
        assert "line-height: 18pt;" in inner.get("style", "")
        assert "text-align: right;" in inner.get("style", "")
        assert "text-indent: -6pt;" in inner.get("style", "")
        assert "margin-top: 3pt;" in inner.get("style", "")
        assert "margin-bottom: 0pt;" in inner.get("style", "")
        assert "text-align: center;" in paragraphs["Outer paragraph after the nested block."].get("style", "")
        assert [" ".join(node.text_content().split()) for node in outer.xpath(".//strong")] == ["bold words"]
        assert [" ".join(node.text_content().split()) for node in outer.xpath(".//em")] == ["emphasis"]
        assert outer.xpath(".//span[@data-custom-style='Revision Char']")
        assert "data-papper-paragraph-format" not in output.read_text(encoding="utf-8")
        assert document.xpath("//div[@id='local-format'][@data-somekey='preserved']")
        assert "font-size: 9pt;" in paragraphs["Chinese Word size and proportional line spacing."].get("style", "")
        assert "line-height: 1.5;" in paragraphs["Chinese Word size and proportional line spacing."].get("style", "")
        assert "font-size: 8pt;" in paragraphs["Explicitly formatted note paragraph."].get("style", "")
        note = next(node for text, node in paragraphs.items() if text.startswith("A note keeps its own formatting."))
        assert "font-size: 10pt" not in note.get("style", "")
        for node in document.xpath("//h2 | //td/p | //figcaption/p"):
            assert "font-size: 10pt" not in node.get("style", "")
        assert "font-size: 10pt" not in paragraphs["Control before the formatted block."].get("style", "")
        return
    document = Document(output)
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    outer = next(paragraph for text, paragraph in paragraphs.items() if text.startswith("Outer paragraph with"))
    inner = next(paragraph for text, paragraph in paragraphs.items() if text.startswith("Inner paragraph"))
    assert outer.paragraph_format.space_before.pt == 3
    assert outer.paragraph_format.space_after.pt == 6
    assert inner.paragraph_format.space_before.pt == 3
    assert inner.paragraph_format.space_after.pt == 0
    assert inner.paragraph_format.line_spacing.pt == 18
    assert inner.paragraph_format.first_line_indent.pt == -6
    assert inner._p.pPr.jc.val == 2  # Word right alignment
    assert outer._p.pPr.ind.get(qn("w:firstLineChars")) == "0"
    assert inner._p.pPr.ind.get(qn("w:firstLineChars")) == "0"
    assert next(run for run in outer.runs if run.text == "bold words").bold is True
    assert next(run for run in outer.runs if run.text == "emphasis").italic is True
    revision = next(run for run in outer.runs if run.text == "revision words")
    assert revision.style.name == "Revision Char"
    assert revision.font.color.rgb is None  # Character style retains its own red color
    assert next(run for run in outer.runs if run.text.startswith("Outer paragraph")).font.size.pt == 10
    assert paragraphs["Control before the formatted block."].paragraph_format.space_after is None
    assert all(run.font.size is None for run in document.tables[0].cell(1, 0).paragraphs[0].runs)
    assert "PMT_PARAGRAPH_FORMAT:" not in document.element.xml


@pytest.mark.parametrize("target", ["html", "docx"])
def test_inner_override_does_not_hide_invalid_outer_format(
    tmp_path: Path, rust_executable: Path, target: str,
) -> None:
    """An invalid outer declaration must fail even when only an overriding inner block has text."""
    source = tmp_path / "nested.md"
    source.write_text('::::: {fontSize="0pt"}\n\n::: {fontSize="10pt"}\nParagraph.\n:::\n\n:::::\n', encoding="utf-8")
    result = build(rust_executable, source, tmp_path / f"nested.{target}")
    assert result.returncode != 0
    assert "fontSize" in result.stderr


@pytest.mark.parametrize("postprocess", [True, False])
def test_local_format_applies_to_page_stories_and_explicit_note_blocks(
    tmp_path: Path, rust_executable: Path, postprocess: bool,
) -> None:
    """Prevent formatting disappearing when Word moves paragraphs into a header or note part."""
    source = tmp_path / "stories.md"
    source.write_text('''---
papperSettings:
  docxHeader:
    alignment: center
    content: |
      ::: {fontSize="9pt" alignment="right" paragraphSpacing.after="4pt"}
      Header with **bold** text.
      :::
---
::: {fontSize="10pt"}
Body with a note.[^note]
:::

[^note]: Ordinary note text.

    ::: {fontSize="8pt"}
    Local note text.
    :::
''', encoding="utf-8")
    output = tmp_path / "stories.docx"
    result = build(rust_executable, source, output, postprocess=postprocess)
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    header = next(paragraph for paragraph in document.sections[0].header.paragraphs if paragraph.text)
    assert header.paragraph_format.space_after.pt == 4
    assert header._p.pPr.jc.val == 2
    assert next(run for run in header.runs if run.text.startswith("Header")).font.size.pt == 9
    assert next(run for run in header.runs if run.text == "bold").bold is True
    with ZipFile(output) as archive:
        notes = etree.fromstring(archive.read("word/footnotes.xml"))
        paragraphs = {"".join(node.itertext()): node for node in notes.iter(qn("w:p"))}
        local = paragraphs["Local note text."]
        ordinary = next(node for text, node in paragraphs.items() if "Ordinary note text." in text)
        assert local.find('.//' + qn("w:sz")).get(qn("w:val")) == "16"
        assert ordinary.find('.//' + qn("w:sz")) is None
        assert b"PMT_PARAGRAPH_FORMAT:" not in archive.read("word/footnotes.xml")


@pytest.mark.parametrize("target", ["html", "docx"])
@pytest.mark.parametrize("attributes,error", [
    ('fontSize="0pt"', "fontSize"),
    ('lineSpacing="0"', "lineSpacing"),
    ('bold="maybe"', "bold"),
    ('firstLineIndentChars="-1"', "firstLineIndentChars"),
    ('firstLineIndentChars="0" indentation.hanging="6pt"', "indent"),
    ('paragraphSpacing.aftr="6pt"', "aftr"),
    ('fontFamily="Arial" fontFamily.chinese="SimSun"', "cannot combine"),
])
def test_invalid_local_format_preserves_previous_output(
    tmp_path: Path, rust_executable: Path, target: str, attributes: str, error: str,
) -> None:
    """Reject malformed values and conflicting formats without replacing the user's output."""
    source = tmp_path / "invalid.md"
    source.write_text(f"::: {{{attributes}}}\nParagraph.\n:::\n", encoding="utf-8")
    output = tmp_path / f"output.{target}"
    output.write_bytes(b"previous output")
    result = build(rust_executable, source, output)
    assert result.returncode != 0
    assert error.lower() in (result.stdout + result.stderr).lower()
    assert output.read_bytes() == b"previous output"
