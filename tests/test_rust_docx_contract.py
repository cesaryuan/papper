"""Compare actual native DOCX builds to the existing complete OPC snapshots.

The fixtures cover user-visible document text, numbering, layout, author
footnotes, revision attributes, style metadata, and image resources. The Rust
command runs directly; Python only canonicalizes artifacts for verification.
"""

from __future__ import annotations

import os
import json
import shutil
import subprocess
from pathlib import Path

import pytest

from snapshot_utils import assert_snapshot, canonical_docx
from test_build_snapshots import CASE_ROOT, CASES, ROOT, copy_case


@pytest.fixture(scope="module")
def rust_postprocessor() -> Path:
    """Build the native standalone package editor for crafted-document contracts."""
    result = subprocess.run(["cargo", "build", "--offline", "-p", "papper-document", "--example", "postprocess_docx"],
                            cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=180)
    assert result.returncode == 0, result.stdout + result.stderr
    return ROOT / "target/debug/examples" / ("postprocess_docx.exe" if os.name == "nt" else "postprocess_docx")


@pytest.fixture(scope="module")
def rust_mathtype_converter() -> Path:
    """Build the package converter for native OLE, layout, and cache behavior checks."""
    result = subprocess.run(["cargo", "build", "--offline", "-p", "papper-document", "--example", "convert_mathtype_docx"],
                            cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=180)
    assert result.returncode == 0, result.stdout + result.stderr
    return ROOT / "target/debug/examples" / ("convert_mathtype_docx.exe" if os.name == "nt" else "convert_mathtype_docx")


@pytest.mark.parametrize("case_name", CASES)
def test_native_docx_matches_existing_snapshot(case_name: str, tmp_path: Path, rust_executable: Path) -> None:
    """Require full package parity for native Word equations and retained Lua filters."""
    case_dir, markdown = CASES[case_name]
    if case_name == "chinese_crossrefs":
        copied = tmp_path / case_name
        copy_case(case_dir, copied)
        case_dir = copied
    output = tmp_path / f"{case_name}.docx"
    command = [str(rust_executable), "build", "docx", "-m", str(case_dir / markdown),
               "-o", str(output), "--no-mathtype"]
    if case_name == "native_crossrefs":
        command.extend(["--style-file", str(case_dir / "style.yml")])
    result = subprocess.run(command, cwd=tmp_path, env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT)},
                            capture_output=True, text=True, encoding="utf-8", timeout=45)
    assert result.returncode == 0, result.stdout + result.stderr
    actual = canonical_docx(output, repository_root=ROOT, project_dir=tmp_path,
                            normalize_native_crossrefs=case_name == "native_crossrefs")
    assert_snapshot(actual, CASE_ROOT / case_name / "snapshots-content" / "docx.snap", update=False)


@pytest.mark.parametrize("postprocess", ["false", "true"])
def test_docx_subfigure_tables_use_reference_style_without_restyling_adjacent_tables(
    tmp_path: Path, rust_executable: Path, postprocess: str,
) -> None:
    """Keep layout styles through DOCX builds and preserve tables beside image captions."""
    from lxml import etree
    from zipfile import ZipFile

    image = (ROOT / "template/examples/images/subfigure-a-example.png").as_posix()
    source = tmp_path / "subfigures.md"
    source.write_text(f"""---
subfigGrid: true
---
<div id="fig:layout">
![Left panel.]({image}){{#fig:left width=49%}}
![Right panel.]({image}){{#fig:right width=49%}}

Grouped panels.
</div>

| Regular cell | Value |
|---|---|
| Ordinary data | 1 |

: {{custom-style="TableNoBorder"}}

`<w:p><w:pPr><w:pStyle w:val="ImageCaption"/></w:pPr><w:r><w:t>Ordinary caption</w:t></w:r></w:p>`{{=openxml}}
""", encoding="utf-8")
    output = tmp_path / "subfigures.docx"
    result = subprocess.run(
        [str(rust_executable), "build", "docx", "-m", str(source), "-o", str(output), "--no-mathtype"],
        cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PMT_ENABLE_DOCX_POSTPROCESS": postprocess},
        capture_output=True, text=True, encoding="utf-8", timeout=45,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    with ZipFile(output) as archive:
        document = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    layouts = document.xpath("//w:tbl[w:tblPr/w:tblStyle/@w:val='TableSubfigure']", namespaces=ns)
    assert len(layouts) == 1
    assert len(layouts[0].xpath(".//w:drawing", namespaces=ns)) == 2
    # Direct zero margins would override future changes to the reference style.
    assert layouts[0].find("w:tblPr/w:tblCellMar", ns) is None
    regular = document.xpath("//w:tbl[.//w:t='Regular cell']", namespaces=ns)
    assert len(regular) == 1
    assert regular[0].find("w:tblPr/w:tblStyle", ns).get(f"{{{ns['w']}}}val") == "TableNoBorder"
    assert regular[0].find("w:tblPr/w:tblCellMar", ns) is None


@pytest.mark.parametrize("has_table_text", [True, False])
@pytest.mark.parametrize("source_style", ["Compact", "First Paragraph", "Body Text"])
def test_native_table_text_preserves_authored_style_or_keeps_source_style(
    tmp_path: Path, rust_postprocessor: Path, has_table_text: bool, source_style: str,
) -> None:
    """Honor reference style inheritance without creating missing table paragraph styles."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt

    document = Document()
    if has_table_text:
        base = document.styles.add_style("Authored Table Base", WD_STYLE_TYPE.PARAGRAPH)
        base.font.size = Pt(13)
        style = document.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = base
    if source_style not in document.styles:
        document.styles.add_style(source_style, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Body contents", style=source_style)
    paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph.text = "Table contents"
    paragraph.style = source_style
    # Authored cell styles must survive normalization of Pandoc's default styles.
    document.tables[0].cell(0, 0).add_paragraph("Authored cell contents", style="Caption")
    source = tmp_path / "source.docx"
    document.save(source)
    metadata = tmp_path / "metadata.json"
    metadata.write_text(json.dumps({"pmt_settings": {"values": {}, "provided": [], "pandoc_metadata": {}, "reply": None},
                                   "pandoc_metadata": {}, "has_yaml_header": False}), encoding="utf-8")
    output = tmp_path / "output.docx"
    result = subprocess.run([str(rust_postprocessor), str(source), str(output), str(metadata)],
                            capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    reopened = Document(output)
    paragraph = reopened.tables[0].cell(0, 0).paragraphs[0]
    assert paragraph.text == "Table contents"
    assert paragraph.style.name == ("Table Text" if has_table_text else source_style)
    assert reopened.paragraphs[0].style.name == source_style
    assert reopened.tables[0].cell(0, 0).paragraphs[1].style.name == "Caption"
    if has_table_text:
        assert paragraph.style.base_style.name == "Authored Table Base"
        assert paragraph.style.base_style.font.size.pt == 13
    else:
        assert "Table Text" not in reopened.styles


def test_docx_consecutive_grid_tables_keep_table_text_in_multi_paragraph_cells(
    tmp_path: Path, rust_executable: Path,
) -> None:
    """Keep both tables uniformly styled without losing paragraphs inside grid cells."""
    from docx import Document

    source = tmp_path / "grid-tables.md"
    # A blank cell line makes Pandoc use First Paragraph/Body Text throughout
    # the second table; crafted single-cell DOCX tests cannot cover this reader path.
    source.write_text("""Paragraph before the tables.

+-------+-------+
| A     | B     |
+=======+=======+
| 1     | 2     |
+-------+-------+

: First table

+-------+-------+
| C     | D     |
+=======+=======+
| 3     | 4     |
+-------+-------+
| 5     | 6     |
|       |       |
|       | 7     |
+-------+-------+

: Second table

Paragraph after the tables.
""", encoding="utf-8")
    output = tmp_path / "grid-tables.docx"
    result = subprocess.run(
        [str(rust_executable), "build", "docx", "-m", str(source), "-o", str(output), "--no-mathtype"],
        cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PMT_ENABLE_DOCX_POSTPROCESS": "true"},
        capture_output=True, text=True, encoding="utf-8", timeout=45,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(output)
    assert len(document.tables) == 2
    expected = [
        [[["A"], ["B"]], [["1"], ["2"]]],
        [[["C"], ["D"]], [["3"], ["4"]], [["5"], ["6", "7"]]],
    ]
    for table, expected_rows in zip(document.tables, expected):
        assert [
            [[paragraph.text for paragraph in cell.paragraphs] for cell in row.cells]
            for row in table.rows
        ] == expected_rows
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    assert paragraph.style.name == "Table Text", paragraph.text
    for text in ["Paragraph before the tables.", "Paragraph after the tables."]:
        paragraph = next(paragraph for paragraph in document.paragraphs if paragraph.text == text)
        assert paragraph.style.name != "Table Text"


def test_native_package_editor_preserves_unknown_parts_and_cancels_inherited_equation_tabs(
    tmp_path: Path, rust_postprocessor: Path,
) -> None:
    """Preserve embedded data and style precedence in a document with custom equation tabs."""
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    from lxml import etree
    from zipfile import ZipFile, ZIP_DEFLATED

    document = Document()
    document.sections[0].page_width = Inches(8.5)
    document.sections[0].left_margin = Inches(1)
    document.sections[0].right_margin = Inches(1.5)
    document.styles["Body Text"].paragraph_format.tab_stops.add_tab_stop(Inches(1))
    equation = document.add_paragraph()
    equation.add_run("\t")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = "x=1"
    math_run.append(math_text)
    math.append(math_run)
    equation._p.append(math)
    equation.add_run("\t(1)")
    equation.paragraph_format.tab_stops.add_tab_stop(Inches(2), WD_TAB_ALIGNMENT.CENTER)
    equation.paragraph_format.space_before = Pt(3)
    equation.paragraph_format.space_after = Pt(18)
    equation.paragraph_format.line_spacing = 2
    source = tmp_path / "source.docx"
    document.save(source)
    payload = bytes(range(256)) * 3
    with ZipFile(source, "a", ZIP_DEFLATED) as archive:
        archive.writestr("customXml/preserved-user-data.bin", payload)
    metadata = tmp_path / "metadata.json"
    metadata.write_text(json.dumps({"pmt_settings": {"values": {}, "provided": [], "pandoc_metadata": {}, "reply": None},
                                   "pandoc_metadata": {}, "has_yaml_header": False}), encoding="utf-8")
    target = tmp_path / "target.docx"
    result = subprocess.run([str(rust_postprocessor), str(source), str(target), str(metadata)],
                            capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    with ZipFile(target) as archive:
        assert archive.read("customXml/preserved-user-data.bin") == payload
    reopened = Document(target)
    paragraph = reopened.paragraphs[0]
    assert paragraph.style.name == "Para Equation"
    assert paragraph.style.base_style.name == "Body Text"
    assert paragraph.paragraph_format.space_before.pt == 3
    assert paragraph._p.pPr.find(qn("w:tabs")) is None
    assert dict(paragraph._p.pPr.find(qn("w:spacing")).attrib) == {qn("w:before"): "60"}
    assert [(tab.position.inches, tab.alignment) for tab in paragraph.style.paragraph_format.tab_stops] == [
        (1, WD_TAB_ALIGNMENT.CLEAR), (3, WD_TAB_ALIGNMENT.CENTER), (6, WD_TAB_ALIGNMENT.RIGHT),
    ]
    assert reopened.styles["Body Text"].paragraph_format.tab_stops[0].position.inches == 1
    # Canonical XML compares the formula itself without depending on inherited
    # namespace declaration order chosen by the package serializer.
    assert etree.tostring(paragraph._p.find(qn("m:oMath")), method="c14n", exclusive=True) == etree.tostring(
        math, method="c14n", exclusive=True,
    )


def test_native_mathtype_preserves_failed_formulas_layout_and_invalidates_corrupt_cache(
    tmp_path: Path, rust_mathtype_converter: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Keep unsupported OMML and attachments while matching real OLE/WMF and recovering bad caches."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from lxml import etree
    from zipfile import ZipFile, ZIP_DEFLATED

    document = Document()
    formulas = [("inline", "x_1"), ("display", r"\frac{"), ("display", "y+2"), ("inline", "x_1")]
    for style, latex in formulas:
        paragraph = document.add_paragraph()
        marker = paragraph.add_run(f"MTLATEX:{style}:{latex}")
        marker.font.hidden = True
        marker.font.size = Pt(10.5)
        math = OxmlElement("m:oMath")
        math_run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = latex
        math_run.append(text)
        math.append(math_run)
        paragraph._p.append(math)
    source = tmp_path / "marked.docx"
    document.save(source)
    payload = b"Unrelated custom data must survive native equation conversion\x00\xff"
    with ZipFile(source, "a", ZIP_DEFLATED) as archive:
        archive.writestr("customXml/attachment.bin", payload)
    metadata = tmp_path / "metadata.json"
    metadata.write_text(json.dumps({"pmt_settings": {"values": {"mathtypeConversionMethod": "rust", "mathtypeSvgBackend": "typst", "mathtypeTypstMathFont": "XITS Math"},
                                                   "provided": [], "pandoc_metadata": {}, "reply": None},
                                    "pandoc_metadata": {}, "has_yaml_header": False}), encoding="utf-8")
    target = tmp_path / "native.docx"
    # An authored resource tree without converter DLLs must still produce real
    # equations: both Rust equation libraries are linked into the executable.
    native_resources = tmp_path / "native-resources"
    (native_resources / "pandoc").mkdir(parents=True)
    (native_resources / "mathtype").mkdir()
    shutil.copy2(ROOT / "pandoc/pandoc-html.yml", native_resources / "pandoc/pandoc-html.yml")
    shutil.copy2(ROOT / "src/pandoc_manuscript/mathtype/Times+Symbol 12.eqp",
                 native_resources / "mathtype/Times+Symbol 12.eqp")
    environment = {**os.environ, "PAPPER_RESOURCE_ROOT": str(native_resources), "PAPPER_HOME": str(tmp_path / "home")}

    def convert() -> dict:
        """Build through the native runner and return its artifact-level conversion report."""
        result = subprocess.run([str(rust_mathtype_converter), str(source), str(target), str(metadata), str(tmp_path)],
                                env=environment, capture_output=True, text=True, encoding="utf-8", timeout=90)
        assert result.returncode == 0, result.stdout + result.stderr
        return json.loads(result.stdout.strip())

    first = convert()
    assert first == {"total": 4, "converted": 3, "failures": 1, "cache_hits": 1}
    original_parts = {}
    with ZipFile(target) as archive:
        assert archive.read("customXml/attachment.bin") == payload
        xml = etree.fromstring(archive.read("word/document.xml"))
        assert b"MTLATEX:" not in archive.read("word/document.xml")
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
              "o": "urn:schemas-microsoft-com:office:office"}
        assert xml.xpath("count(.//o:OLEObject)", namespaces=ns) == 3
        assert xml.xpath(".//m:oMath//m:t/text()", namespaces=ns) == [r"\frac{"]
        for index in (0, 2, 3):
            ole = archive.read(f"word/embeddings/mathtype_formula_{index + 1}.bin")
            wmf = archive.read(f"word/media/mathtype_formula_{index + 1}.wmf")
            assert ole.startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
            assert wmf.startswith(bytes.fromhex("d7cdc69a"))
            original_parts[index] = (ole, wmf)
        assert original_parts[0] == original_parts[3]
        # A display formula must not inherit inline baseline positioning.
        display = xml.xpath(".//w:r[o:OLEObject or w:object/o:OLEObject[@ShapeID='_x0000_i3003']]", namespaces=ns)
        assert display and not display[0].xpath("w:rPr/w:position", namespaces=ns)
        inline = xml.xpath(".//w:r[w:object/o:OLEObject[@ShapeID='_x0000_i3001']]/w:rPr/w:position/@w:val", namespaces=ns)
        metadata_path = next(preview.parent / "metadata.json" for preview in (tmp_path / "home").rglob("preview.wmf")
                             if preview.read_bytes() == original_parts[0][1])
        baseline = json.loads(metadata_path.read_bytes())["mathtype"]["baseline_from_bottom_pt"]
        assert inline == [str(-round(baseline * 2))]
    assert convert()["cache_hits"] == 3
    caches = list((tmp_path / "home").rglob("equation.ole.bin"))
    assert len(caches) == 2
    for cached in caches:
        cached.write_bytes(b"broken cache must never enter a document")
    recovered = convert()
    assert recovered == first
    with ZipFile(target) as archive:
        assert archive.read("word/embeddings/mathtype_formula_1.bin") == original_parts[0][0]
        assert archive.read("word/media/mathtype_formula_1.wmf") == original_parts[0][1]


def test_markdown_mathtype_build_preserves_control_space_and_supports_mspace(
    tmp_path: Path, rust_executable: Path,
) -> None:
    """Convert the reported equation shapes through Lua markers, native OLE, and Typst previews."""
    from lxml import etree
    from zipfile import ZipFile
    import io
    import olefile

    source = tmp_path / "spacing.md"
    source.write_text(r"""---
mathtypeConversionMethod: rust
mathtypeSvgBackend: typst
---

$$\left\{ \begin{aligned}
 & U_n(\omega) = H_n(\omega) F_e(\omega) \\
 & U_m(\omega) = H_m(\omega) F_e(\omega)
\end{aligned} \right.\ $$ {#eq:control-space}

$$\mathcal{L}_{adv} = \sum_{i=1}^{n_s}\mspace{2mu}\gamma(y_i) + \sum_{j=1}^{n_t}\mspace{2mu}\eta(x_j)$$ {#eq:adversarial}

$$\mathcal{L}_{total} = \lambda_d\sum_{m=1}^{N_m}\mspace{2mu}\left(L_{adv}+L_{cov}\right) + \lambda_e\sum_{m=1}^{N_m}\mspace{2mu}L_{ent}$$ {#eq:total}
""", encoding="utf-8")
    target = tmp_path / "spacing.docx"
    result = subprocess.run(
        [str(rust_executable), "build", "docx", "-m", str(source), "-o", str(target)],
        cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "home")},
        capture_output=True, text=True, encoding="utf-8", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    with ZipFile(target) as archive:
        document = archive.read("word/document.xml")
        root = etree.fromstring(document)
        ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
              "o": "urn:schemas-microsoft-com:office:office"}
        assert root.xpath("count(.//o:OLEObject)", namespaces=ns) == 3
        assert not root.xpath(".//m:oMath", namespaces=ns)
        assert b"MTLATEX:" not in document
        for index in range(1, 4):
            wmf = archive.read(f"word/media/mathtype_formula_{index}.wmf")
            assert wmf.startswith(bytes.fromhex("d7cdc69a"))
            with olefile.OleFileIO(io.BytesIO(archive.read(f"word/embeddings/mathtype_formula_{index}.bin"))) as ole:
                native = ole.openstream("Equation Native").read()
                # Preview compatibility must preserve the recoverable equation source.
                if index == 1:
                    assert b"\\right.\\ " in native
                else:
                    assert b"\\mspace{2mu}" in native


@pytest.mark.parametrize("part", ["body", "header", "footer"])
def test_native_docx_preserves_angle_bracket_notation(
    tmp_path: Path, rust_postprocessor: Path, part: str,
) -> None:
    """Preserve literal scientific notation without rejecting it as unrendered HTML."""
    from docx import Document

    document = Document()
    paragraph = (
        document.add_paragraph()
        if part == "body"
        else getattr(document.sections[0], part).paragraphs[0]
    )
    # Word can split the reported <k> notation across runs; validation must
    # accept the joined text and preserve every character in the saved result.
    for text in ["式中，<", "k", ">为网络的平均度；<degree>、<K> 和 <k:avg> 为其他记号"]:
        paragraph.add_run(text)
    expected = paragraph.text
    source = tmp_path / "notation.docx"
    document.save(source)
    metadata = tmp_path / "metadata.json"
    metadata.write_text(json.dumps({"pmt_settings": {"values": {}, "provided": [], "pandoc_metadata": {}, "reply": None},
                                   "pandoc_metadata": {}, "has_yaml_header": False}), encoding="utf-8")
    output = tmp_path / "output.docx"
    result = subprocess.run([str(rust_postprocessor), str(source), str(output), str(metadata)],
                            capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    reopened = Document(output)
    actual = (
        reopened.paragraphs[0]
        if part == "body"
        else getattr(reopened.sections[0], part).paragraphs[0]
    )
    assert actual.text == expected


def test_native_docx_rejects_unrendered_reference_syntax_without_replacing_existing_output(
    tmp_path: Path, rust_postprocessor: Path,
) -> None:
    """Reject broken visible references in body and footer while retaining an existing result."""
    from docx import Document

    metadata = tmp_path / "metadata.json"
    metadata.write_text(json.dumps({"pmt_settings": {"values": {}, "provided": [], "pandoc_metadata": {}, "reply": None},
                                   "pandoc_metadata": {}, "has_yaml_header": False}), encoding="utf-8")
    output = tmp_path / "previous.docx"
    original = b"previous successful output"
    for text in [
        "Equation (Equation 6)", "Table 3 Table 3", "[@citation]", "::: {.note}", "Figure 2.1 Figure 2.1",
        "<div>", "</span>", '<img src="figure.png">', "<BR/>", "<!-- unrendered comment -->",
        "式中，<k>为网络的平均度 <span>残留标签</span>",
    ]:
        document = Document()
        if text.startswith("Figure"):
            document.add_paragraph("A valid body")
            document.sections[0].footer.paragraphs[0].text = text
        else:
            document.add_paragraph(text)
        source = tmp_path / "broken.docx"
        document.save(source)
        output.write_bytes(original)
        result = subprocess.run([str(rust_postprocessor), str(source), str(output), str(metadata)],
                                capture_output=True, text=True, encoding="utf-8", timeout=30)
        assert result.returncode != 0, text
        assert "unrendered Pandoc syntax" in result.stderr
        assert output.read_bytes() == original
    # Formula TeX is an intermediate hidden marker, not final visible text; it
    # can legally contain strings that resemble broken prose references.
    document = Document()
    paragraph = document.add_paragraph("A valid formula follows")
    hidden = paragraph.add_run(r"MTLATEX:inline:\text{Figure Figure}")
    hidden.font.hidden = True
    source = tmp_path / "marker.docx"
    document.save(source)
    result = subprocess.run([str(rust_postprocessor), str(source), str(output), str(metadata)],
                            capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    assert Document(output).paragraphs[0].text.startswith("A valid formula follows")
