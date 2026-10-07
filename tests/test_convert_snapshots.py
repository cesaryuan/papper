"""Snapshot DOCX imports as Markdown and a normalized media filename tree.

Each Convert case keeps its input DOCX beside ``snapshots-content/``. Tests run
the public Rust Convert command against the checked-in DOCX, without rebuilding
it. Temporary output and golden PNG/WMF files use tiny white images with black
borders; their paths, not bytes, are the contract here. Refresh explicitly with
``--snapshot-update`` and review the diff.
"""

from __future__ import annotations

import os
import struct
import subprocess
import zlib
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree

from native_support import ROOT
from snapshot_utils import assert_snapshot


CASE_ROOT = Path(__file__).with_name("snapshot_cases_convert")


@pytest.mark.parametrize(("text", "expected_language"), [
    pytest.param("中文正文转换语言中a", "zh-CN", id="exactly-90-percent"),
    pytest.param("中" * 91 + "a" * 9, "zh-CN", id="above-90-percent"),
    pytest.param("中" * 89 + "a" * 11, None, id="below-90-percent"),
    pytest.param("𠀀" * 9 + "a", "zh-CN", id="supplementary-han"),
    pytest.param("中文， 正文！\t转换：语言；中。 a", "zh-CN", id="ignore-punctuation-and-space"),
    pytest.param("中" * 8 + "a1", None, id="count-digits"),
    pytest.param("An English manuscript", None, id="english"),
    pytest.param("", None, id="empty"),
    pytest.param("，。！？", None, id="punctuation-only"),
])
def test_convert_detects_chinese_language_from_document_text(
    text: str, expected_language: str | None, tmp_path: Path, rust_executable: Path,
) -> None:
    """Check actual imports at the language threshold, preserving text and portable output."""
    import yaml
    from docx import Document

    source = tmp_path / "language.docx"
    document = Document()
    document.add_paragraph(text)
    document.save(source)
    output = tmp_path / "converted"
    result = subprocess.run(
        [str(rust_executable), "convert", str(source), "-o", str(output)], cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "home")},
        capture_output=True, text=True, encoding="utf-8", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    markdown = (output / "language.md").read_bytes().decode("utf-8")
    assert "\r" not in markdown
    if expected_language:
        assert markdown.startswith("---\n")
        header, body = markdown[4:].split("\n---\n", 1)
        assert yaml.safe_load(header)["lang"] == expected_language
    else:
        assert not markdown.startswith("---\n")
        body = markdown
    # Pandoc may escape punctuation or normalize spaces; authored letters and
    # numbers must still survive the new standalone output mode in order.
    assert "".join(c for c in body if c.isalnum()) == "".join(c for c in text if c.isalnum())


@pytest.mark.parametrize("note_type", ["footnote", "endnote"])
def test_convert_includes_notes_in_language_detection(
    note_type: str, tmp_path: Path, rust_executable: Path,
) -> None:
    """Include actual note text in the language ratio and retain it in Markdown."""
    import yaml
    from docx import Document
    from docx.oxml import OxmlElement

    source = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("a").add_run()._r.append(OxmlElement(f"w:{note_type}Reference"))
    reference = document.paragraphs[0]._p[-1][-1]
    reference.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id", "1")
    document.save(source)
    with ZipFile(source) as archive:
        parts = {entry.filename: archive.read(entry) for entry in archive.infolist()}
    word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    note_text = "中文正文转换语言中"
    parts[f"word/{note_type}s.xml"] = (
        f'<w:{note_type}s xmlns:w="{word_ns}"><w:{note_type} w:id="1">'
        f'<w:p><w:r><w:t>{note_text}</w:t></w:r></w:p></w:{note_type}></w:{note_type}s>'
    ).encode("utf-8")
    relationships = etree.fromstring(parts["word/_rels/document.xml.rels"])
    etree.SubElement(relationships, f"{{{relationships.nsmap[None]}}}Relationship", {
        "Id": "rIdPapperNote", "Target": f"{note_type}s.xml",
        "Type": f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{note_type}s",
    })
    parts["word/_rels/document.xml.rels"] = etree.tostring(relationships)
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(content_types, f"{{{content_types.nsmap[None]}}}Override", {
        "PartName": f"/word/{note_type}s.xml",
        "ContentType": f"application/vnd.openxmlformats-officedocument.wordprocessingml.{note_type}s+xml",
    })
    parts["[Content_Types].xml"] = etree.tostring(content_types)
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    output = tmp_path / "converted"
    result = subprocess.run(
        [str(rust_executable), "convert", str(source), "-o", str(output)], cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "home")},
        capture_output=True, text=True, encoding="utf-8", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    markdown = (output / "notes.md").read_text(encoding="utf-8")
    assert markdown.startswith("---\n")
    header, body = markdown[4:].split("\n---\n", 1)
    assert yaml.safe_load(header)["lang"] == "zh-CN"
    assert "[^1]: " + note_text in body


def png_chunk(kind: bytes, data: bytes) -> bytes:
    """Encode a PNG chunk with its length and CRC, without ancillary metadata."""
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))


def media_placeholder(extension: str) -> bytes:
    """Create small valid 16×16 white/black PNG and vector WMF placeholders."""
    if extension == ".png":
        # One-bit grayscale avoids palettes and RGB overhead. Zero scanline
        # filters plus maximum DEFLATE compression keep the simple border tiny.
        pixels = b"\x00\x00\x00" + b"\x00\x7f\xfe" * 14 + b"\x00\x00\x00"
        header = struct.pack(">IIBBBBB", 16, 16, 1, 0, 0, 0, 0)
        return (b"\x89PNG\r\n\x1a\n" + png_chunk(b"IHDR", header)
                + png_chunk(b"IDAT", zlib.compress(pixels, level=9)) + png_chunk(b"IEND", b""))
    if extension == ".wmf":
        # A placeable WMF scales to 16 pixels at 96 dpi. Draw one rectangle
        # using a white solid brush and a black pen; no raster payload is needed.
        operations = [
            (0x0103, (8,)),                         # Anisotropic mapping
            (0x020B, (0, 0)),                       # Window origin
            (0x020C, (160, 160)),                   # Window extent
            (0x02FA, (0, 10, 0, 0, 0)),             # Solid black pen
            (0x012D, (0,)),                         # Select pen
            (0x02FC, (0, 0xFFFF, 0x00FF, 0)),       # Solid white brush
            (0x012D, (1,)),                         # Select brush
            (0x041B, (155, 155, 5, 5)),             # Inset rectangle
            (0x0000, ()),                           # End metafile
        ]
        records = [struct.pack("<IH", 3 + len(parameters), operation)
                   + struct.pack(f"<{len(parameters)}H", *parameters)
                   for operation, parameters in operations]
        placeable = struct.pack("<IHhhhhHI", 0x9AC6CDD7, 0, 0, 0, 160, 160, 960, 0)
        checksum = 0
        for word in struct.unpack("<10H", placeable):
            checksum ^= word
        body = b"".join(records)
        header = struct.pack("<HHHIHIH", 1, 9, 0x0300, 9 + len(body) // 2, 2,
                             max(len(record) // 2 for record in records), 0)
        return placeable + struct.pack("<H", checksum) + header + body
    # Other extracted media still use filename-only placeholders until a
    # format-specific stand-in is needed, rather than copying document content.
    return b""


def output_tree(directory: Path) -> list[str]:
    """Represent the complete published tree, including unexpected or empty directories."""
    return sorted(path.relative_to(directory).as_posix() + ("/" if path.is_dir() else "")
                  for path in directory.rglob("*"))


def normalize_output_media(directory: Path) -> None:
    """Replace extracted media in the test output before snapshot comparison/copy."""
    # Normalize actual output too: previously only golden files were replaced,
    # so inspecting pytest's converted/media directory still displayed originals.
    for path in (directory / "media").rglob("*"):
        if path.is_file():
            path.write_bytes(media_placeholder(path.suffix.lower()))


@pytest.mark.parametrize("case_name", ["comprehensive"])
def test_convert_output_matches_snapshot(
    case_name: str, tmp_path: Path, rust_executable: Path, snapshot_update: bool,
) -> None:
    """Detect changed imported syntax, missing pictures, and leaked formula previews."""
    source = CASE_ROOT / case_name / f"{case_name}.docx"
    output = tmp_path / "converted"
    result = subprocess.run(
        [str(rust_executable), "convert", str(source), "-o", str(output)],
        cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "home")},
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    normalize_output_media(output)
    expected = CASE_ROOT / case_name / "snapshots-content"
    markdown_name = f"{case_name}.md"
    # read_bytes preserves line endings: portable LF is part of Convert output.
    markdown = (output / markdown_name).read_bytes().decode("utf-8")
    # A Word save can merge fallback previews with genuine equations. Updating
    # snapshots must not silently accept a case that stopped covering recovery.
    assert "[@eq:_RefConvertTwoCell]" in markdown and "[@eq:_RefConvertTextLabel]" in markdown, (
        "Comprehensive fixture lost equation recovery; check shared Word previews before refreshing snapshots"
    )
    # These authored pictures have no OLE; the snapshot must exercise WMF MTEF
    # recovery rather than silently accepting leaked formula preview images.
    wmf_section = markdown.split("# 仅有WMF的公式\n", 1)[1].split("[^1]:", 1)[0]
    assert wmf_section.count("$") == 14 and "![](" not in wmf_section
    assert "![](media/undecodable.wmf)" in markdown
    assert "![](media/non-equation.wmf)" in markdown
    assert markdown.count("![](media/shared-preview.wmf)") == 2
    if snapshot_update:
        # Delete only this case's previous golden output so removed media names
        # cannot silently remain in the expected directory after a refresh.
        assert expected.resolve().parent == (CASE_ROOT / case_name).resolve(), (
            "Snapshot update must remain inside its case root"
        )
        for path in sorted(expected.rglob("*"), key=lambda value: len(value.parts), reverse=True):
            if path.is_dir():
                path.rmdir()
            else:
                path.unlink()
        for path in output.rglob("*"):
            target = expected / path.relative_to(output)
            if path.is_dir():
                target.mkdir(parents=True, exist_ok=True)
            elif path.relative_to(output).parts[0] == "media":
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(path.read_bytes())
    assert_snapshot(markdown, expected / markdown_name, update=snapshot_update)
    assert output_tree(output) == output_tree(expected), "Convert output file/directory tree differs from the snapshot"


@pytest.mark.parametrize("ole_state", ["corrupt", "missing", "absent"])
def test_convert_recovers_wmf_when_ole_is_unusable(
    ole_state: str, tmp_path: Path, rust_executable: Path,
) -> None:
    """Recover the same visible formula from a preview after OLE corruption or loss."""
    source = CASE_ROOT / "comprehensive/comprehensive.docx"
    with ZipFile(source) as archive:
        parts = {entry.filename: archive.read(entry) for entry in archive.infolist()}
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "o": "urn:schemas-microsoft-com:office:office",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    document = etree.fromstring(parts["word/document.xml"])
    body = document.find("w:body", ns)
    paragraph = next(p for p in body.findall("w:p", ns) if p.find(".//o:OLEObject", ns) is not None)
    ole = paragraph.find(".//o:OLEObject", ns)
    relations = etree.fromstring(parts["word/_rels/document.xml.rels"])
    target = next(rel.get("Target") for rel in relations if rel.get("Id") == ole.get(f"{{{ns['r']}}}id"))
    if ole_state == "corrupt":
        parts[f"word/{target}"] = b"damaged OLE"
    elif ole_state == "missing":
        del parts[f"word/{target}"]
    else:
        ole.getparent().remove(ole)
    # Isolate this occurrence so another identical equation cannot hide a failure.
    for child in list(body):
        if child is not paragraph and child.tag != f"{{{ns['w']}}}sectPr":
            body.remove(child)
    parts["word/document.xml"] = etree.tostring(document, encoding="UTF-8", xml_declaration=True)
    damaged = tmp_path / "fallback.docx"
    with ZipFile(damaged, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    output = tmp_path / "converted"
    result = subprocess.run(
        [str(rust_executable), "convert", str(damaged), "-o", str(output)], cwd=tmp_path,
        env={**os.environ, "PAPPER_RESOURCE_ROOT": str(ROOT), "PAPPER_HOME": str(tmp_path / "home")},
        capture_output=True, text=True, encoding="utf-8", timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    markdown = (output / "fallback.md").read_text(encoding="utf-8")
    assert markdown.strip() == r"$$\sqrt{{b^2}-4ac}$$"
    assert not list(output.rglob("*.wmf")), "Recovered previews must not leak into published media"
